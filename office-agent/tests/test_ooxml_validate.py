from pathlib import Path

from lxml import etree

from coscribe.tools._ooxml_validate import assert_ooxml_valid, ooxml_errors

_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"

_NSMAP = {"p": _P_NS, "a": _A_NS, "r": _R_NS}


def _bare_slide() -> object:
    """A minimal, schema-correct <p:sld> -- same shape python-pptx's own
    `add_slide` produces (cSld/spTree/nvGrpSpPr with the mandatory
    id/name/cNvGrpSpPr/nvPr, then an empty grpSpPr)."""
    sld = etree.Element(f"{{{_P_NS}}}sld", nsmap=_NSMAP)
    c_sld = etree.SubElement(sld, f"{{{_P_NS}}}cSld")
    sp_tree = etree.SubElement(c_sld, f"{{{_P_NS}}}spTree")
    nv_grp = etree.SubElement(sp_tree, f"{{{_P_NS}}}nvGrpSpPr")
    etree.SubElement(nv_grp, f"{{{_P_NS}}}cNvPr", id="1", name="")
    etree.SubElement(nv_grp, f"{{{_P_NS}}}cNvGrpSpPr")
    etree.SubElement(nv_grp, f"{{{_P_NS}}}nvPr")
    etree.SubElement(sp_tree, f"{{{_P_NS}}}grpSpPr")
    return sld


def test_valid_real_pptx_slide_passes(tmp_path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Hello"
    prs.save(tmp_path / "deck.pptx")

    reopened = Presentation(tmp_path / "deck.pptx")
    assert ooxml_errors(reopened.slides[0].element) == []
    assert_ooxml_valid(reopened.slides[0].element, "test")  # must not raise


def test_theme_root_validates_against_the_same_schema() -> None:
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    prs = Presentation()
    theme_part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    theme_root = etree.fromstring(theme_part.blob)

    assert ooxml_errors(theme_root) == []


def test_wrong_child_order_is_rejected() -> None:
    # ECMA-376's CT_Slide requires <p:transition> after <p:cSld>, never
    # before it -- a real defect this validator exists to catch.
    sld = _bare_slide()
    transition = etree.Element(f"{{{_P_NS}}}transition")
    sld.insert(0, transition)

    errors = ooxml_errors(sld)
    assert errors != []
    assert any("transition" in error for error in errors)


def test_assert_ooxml_valid_raises_with_context_and_leaves_no_side_effect() -> None:
    sld = _bare_slide()
    sld.insert(0, etree.Element(f"{{{_P_NS}}}transition"))

    try:
        assert_ooxml_valid(sld, "my_tool")
        raise AssertionError("expected RuntimeError")
    except RuntimeError as exc:
        assert "my_tool" in str(exc)
        assert "not modified" in str(exc)


def test_extension_attribute_without_mc_ignorable_is_rejected() -> None:
    # The false-positive this module exists to avoid also has a real
    # negative case: an extension attribute with NO matching mc:Ignorable
    # declaration anywhere is a genuine defect, not something to wave
    # through just because it's in a foreign namespace.
    sld = _bare_slide()
    transition = etree.SubElement(
        sld, f"{{{_P_NS}}}transition", nsmap={"p14": _P14_NS}
    )
    transition.set(f"{{{_P14_NS}}}dur", "800")
    etree.SubElement(transition, f"{{{_P_NS}}}fade")

    errors = ooxml_errors(sld)
    assert errors != []


def test_extension_attribute_with_mc_ignorable_is_accepted() -> None:
    # Same document as above, but with the mc:Ignorable declaration real
    # PowerPoint-authored files carry alongside p14: content -- this is
    # exactly what _set_slide_transition/_mark_mce_ignorable produce.
    sld = _bare_slide()
    transition = etree.SubElement(
        sld, f"{{{_P_NS}}}transition", nsmap={"p14": _P14_NS}
    )
    transition.set(f"{{{_P14_NS}}}dur", "800")
    etree.SubElement(transition, f"{{{_P_NS}}}fade")
    sld.set(f"{{{_MC_NS}}}Ignorable", "p14")

    assert ooxml_errors(sld) == []
    # The real element passed in must be untouched -- validation works on
    # a stripped copy, never mutates the caller's own tree.
    assert sld.get(f"{{{_MC_NS}}}Ignorable") == "p14"
    assert transition.get(f"{{{_P14_NS}}}dur") == "800"


def test_unrelated_ignorable_namespace_does_not_hide_a_real_defect() -> None:
    # mc:Ignorable="p14" must only excuse p14: content -- a genuine
    # out-of-order core element right next to it should still fail.
    sld = _bare_slide()
    sld.set(f"{{{_MC_NS}}}Ignorable", "p14")
    bogus = etree.Element(f"{{{_P_NS}}}transition")
    sld.insert(0, bogus)

    errors = ooxml_errors(sld)
    assert errors != []
