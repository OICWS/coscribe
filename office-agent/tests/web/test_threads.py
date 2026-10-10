# ruff: noqa: E402
"""Web tests: threads, folders, uploads, activity and history replay."""

import asyncio
import json
import shutil
import tempfile
import threading
import time
from pathlib import Path
from typing import Any

import pytest

pytest.importorskip("langgraph", reason="needs the langgraph_spike extra installed")

from docx import Document
from langchain_core.messages import (
    AIMessage,
    BaseMessage,
    HumanMessage,
    ToolMessage,
)
from langchain_core.outputs import ChatGeneration, ChatResult
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from coscribe.conversation.thread_meta import ThreadMetaStore
from coscribe.tools.presentations import PresentationToolkit

from .helpers import (
    FakeToolCallingChatModel,
    _client_lg,
    _receive_until,
    _run_turn,
    _stub_script_env,
    _tool_call,
    _wait_for_run_status,
)


def test_threads_can_be_grouped_archived_and_their_groups_renamed_and_deleted_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/g1") as ws:
            ws.receive_json()
            ws.receive_json()
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")

        moved = client.post("/api/threads/g1/meta", json={"group": " Reports "}).json()
        again = client.post("/api/threads/g1/meta", json={"group": "reports"}).json()
        archived = client.post("/api/threads/g1/meta", json={"archived": True}).json()
        with client.websocket_connect("/ws/g1") as ws:
            ws.receive_json()
            ws.receive_json()
            # No more scripted replies: the message only has to reach the turn.
            ws.send_json({"type": "user_message", "text": "again"})
            for _ in range(40):
                if not ThreadMetaStore(tmp_path / "state").get("g1")["archived"]:
                    break
                time.sleep(0.05)
        revived = ThreadMetaStore(tmp_path / "state").get("g1")["archived"]
        client.post("/api/threads/g1/meta", json={"archived": True})
        listed = {t["thread_id"]: t for t in client.get("/api/threads").json()}
        renamed = client.post("/api/thread-groups/Reports/rename", json={"name": "Monthly"}).json()
        after_rename = client.get("/api/threads").json()[0]
        client.post("/api/threads/g1/meta", json={"group": "Other"})
        clash = client.post("/api/thread-groups/Other/rename", json={"name": "monthly"})
        blank = client.post("/api/threads/g1/meta", json={"group": "   "})
        client.post("/api/threads/g1/meta", json={"group": "Monthly"})
        deleted = client.delete("/api/thread-groups/Monthly").json()
        after_delete = client.get("/api/threads").json()[0]
        missing = client.delete("/api/thread-groups/Nope")
        client.delete("/api/threads/g1")
        leftover = (tmp_path / "state" / "g1.meta.json").exists()

    assert moved["group"] == "Reports" and moved["groups"] == ["Reports"]
    assert again["group"] == "Reports" and again["groups"] == ["Reports"]
    assert archived["archived"] is True
    assert revived is False
    assert listed["g1"]["group"] == "Reports" and listed["g1"]["archived"] is True
    assert renamed["groups"] == ["Monthly"] and after_rename["group"] == "Monthly"
    assert clash.status_code == 400
    assert blank.status_code == 200 and blank.json()["group"] is None
    assert deleted["groups"] == ["Other"] and after_delete["group"] is None
    assert missing.status_code == 404
    assert not leftover


def _libreoffice_actually_works() -> bool:
    """Same probe technique as test_presentations_tool.py's identical
    helper (not imported from there -- each test file in this project
    stays self-contained): `shutil.which` alone can't tell a genuinely
    broken soffice install from a working one, so this actually tries a
    trivial conversion in a throwaway temp dir."""
    probe_dir = Path(tempfile.mkdtemp(prefix="coscribe_lo_probe_"))
    try:
        result = PresentationToolkit(probe_dir).write_pptx(
            path="probe.pptx", content="# Probe\n- hi"
        )
        return result["qa_skipped_reason"] is None
    finally:
        shutil.rmtree(probe_dir, ignore_errors=True)


def _write_test_deck(workspace: Path, name: str = "deck.pptx") -> None:
    """A one-slide deck with a single plain (unfilled) textbox -- enough
    for the approval-preview tests below to have something real to
    diff a fill-color edit against."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text_frame.text = "hello"
    (workspace).mkdir(parents=True, exist_ok=True)
    prs.save(str(workspace / name))


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed/functional in this environment",
)
def test_approval_required_carries_a_before_after_preview_for_a_pptx_edit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """web/session.py's _build_document_edit_preview dry-runs the exact same
    edit_pptx_shape call against a throwaway copy of the deck and renders
    both states, so the approval card can show what the edit will
    actually produce instead of a raw JSON arguments dump (see
    ROADMAP.md's approval-preview entry). End-to-end through a real WS
    connection and a real LibreOffice/poppler conversion -- not mocked."""
    _write_test_deck(tmp_path / "workspace")
    call = _tool_call(
        "call_1",
        "edit_pptx_shape",
        {"path": "deck.pptx", "slide": 1, "shape_index": 0, "fill_color": "38BDF8"},
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_pptx_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "recolor that shape"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["tool_name"] == "edit_pptx_shape"
            before_name = approval["before_preview"]
            after_name = approval["after_preview"]
            assert before_name is not None
            assert after_name is not None
            assert before_name != after_name

            # The dry run must never touch the real file -- still
            # un-filled at this point, approval hasn't happened yet.
            real_deck = Presentation(str(tmp_path / "workspace" / "deck.pptx"))
            assert real_deck.slides[0].shapes[0].fill.type is None or str(
                real_deck.slides[0].shapes[0].fill.type
            ).startswith("BACKGROUND")

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")

    previews_dir = tmp_path / "state" / "previews"
    assert (previews_dir / before_name).is_file()
    assert (previews_dir / after_name).is_file()


def test_approval_preview_is_absent_for_a_non_pptx_tool(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """write_file has nothing to do with PresentationToolkit at all --
    _build_document_edit_preview must fall through to (None, None) instantly,
    not attempt (and fail at) treating its arguments as a pptx edit. No
    LibreOffice needed for this one, so it isn't gated behind the
    real_libreoffice marker."""
    call = _tool_call("call_1", "write_file", {"path": "note.txt", "content": "hi"})
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_no_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "write hi to note.txt"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["before_preview"] is None
            assert approval["after_preview"] is None

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": False})
            _receive_until(ws, "agent_message")


def test_approval_preview_is_absent_when_the_target_file_does_not_exist_yet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """edit_pptx_shape on a path that doesn't exist yet has no "before"
    to show -- _build_document_edit_preview must recognize that up front
    (real_path.is_file() is False) and return (None, None), not attempt a
    dry run that would just fail. No LibreOffice needed: this returns
    before ever shelling out to soffice."""
    call = _tool_call(
        "call_1",
        "edit_pptx_shape",
        {"path": "missing.pptx", "slide": 1, "shape_index": 0, "fill_color": "38BDF8"},
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_missing_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "recolor that shape"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["before_preview"] is None
            assert approval["after_preview"] is None

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": False})
            _receive_until(ws, "agent_message")


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed/functional in this environment",
)
def test_approval_preview_covers_write_docx_overwriting_an_existing_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """_build_document_edit_preview's dispatch table covers docx too,
    dispatched via DocumentToolkit -- write_docx regenerates the *whole*
    document from markdown-like `content`, so calling it against an
    already-existing file (this scenario: revising a one-paragraph status
    report) is exactly the "before/after" case worth previewing. A fresh
    write_docx with no existing target still correctly gets no preview
    (nothing to diff against) -- covered separately below, not repeated
    here."""
    workspace = tmp_path / "workspace"
    workspace.mkdir(parents=True, exist_ok=True)
    original = Document()
    original.add_paragraph("Status: on track.")
    original.save(str(workspace / "status.docx"))

    call = _tool_call(
        "call_1",
        "write_docx",
        {
            "path": "status.docx",
            "content": "Status: **delayed** -- see risks below.",
            "overwrite": True,
        },
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_docx_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "update the status report"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["tool_name"] == "write_docx"
            before_name = approval["before_preview"]
            after_name = approval["after_preview"]
            assert before_name is not None
            assert after_name is not None
            assert before_name != after_name

            # Untouched until approved.
            real_doc = Document(str(workspace / "status.docx"))
            assert real_doc.paragraphs[0].text == "Status: on track."

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")

    previews_dir = tmp_path / "state" / "previews"
    assert (previews_dir / before_name).is_file()
    assert (previews_dir / after_name).is_file()
    updated_doc = Document(str(workspace / "status.docx"))
    assert "delayed" in updated_doc.paragraphs[0].text


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed/functional in this environment",
)
def test_approval_preview_covers_format_xlsx_cells(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """_build_document_edit_preview's dispatch table covers xlsx too, via
    SpreadsheetToolkit -- format_xlsx_cells never creates a file (only
    READ_check's the target), so real_path.is_file() being required is
    what makes this safe to dry-run unconditionally. Scenario: bolding
    and red-coloring a budget overage cell in an existing sheet."""
    workspace = tmp_path / "workspace"
    workspace.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet["A1"] = "Category"
    sheet["B1"] = "Spent"
    sheet["A2"] = "Marketing"
    sheet["B2"] = 15200
    workbook.save(str(workspace / "budget.xlsx"))

    call = _tool_call(
        "call_1",
        "format_xlsx_cells",
        {
            "path": "budget.xlsx",
            "sheet_name": "Budget",
            "cell_range": "B2",
            "font_color": "FF0000",
            "bold": True,
        },
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_xlsx_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "flag the marketing overage in red bold"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["tool_name"] == "format_xlsx_cells"
            before_name = approval["before_preview"]
            after_name = approval["after_preview"]
            assert before_name is not None
            assert after_name is not None
            assert before_name != after_name

            real_workbook = load_workbook(str(workspace / "budget.xlsx"))
            assert real_workbook["Budget"]["B2"].font.bold is not True

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")

    previews_dir = tmp_path / "state" / "previews"
    assert (previews_dir / before_name).is_file()
    assert (previews_dir / after_name).is_file()
    updated_workbook = load_workbook(str(workspace / "budget.xlsx"))
    assert updated_workbook["Budget"]["B2"].font.bold is True


def test_approval_preview_is_absent_for_a_brand_new_write_docx(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """write_docx creating a file that doesn't exist yet has no "before"
    -- same "nothing to diff against" reasoning fill_pptx_template gets,
    now confirmed for docx's own create-vs-overwrite ambiguity (write_docx
    is *both* a create and an overwrite tool depending on whether the
    target already exists). No LibreOffice needed -- this returns before
    ever shelling out to soffice."""
    call = _tool_call(
        "call_1",
        "write_docx",
        {"path": "new_report.docx", "content": "Hello.", "overwrite": True},
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_new_docx") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "write a new report"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["before_preview"] is None
            assert approval["after_preview"] is None

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")


def test_approval_preview_leaves_the_real_file_untouched_on_denial(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Denying an edit_pptx_shape call must leave the real file exactly as
    it was -- true by construction for the real *execution* (a denied
    HumanInTheLoopMiddleware decision never runs the tool at all), but
    this is specifically about the *preview* mechanism: confirms the dry
    run itself (which runs regardless of what the human eventually
    decides, since the preview is built before the approval prompt is
    even sent) never wrote through to the real path. No LibreOffice
    needed for the assertion that matters here (the file bytes are
    unchanged); the preview images themselves are covered elsewhere."""
    _write_test_deck(tmp_path / "workspace")
    original_bytes = (tmp_path / "workspace" / "deck.pptx").read_bytes()
    call = _tool_call(
        "call_1",
        "edit_pptx_shape",
        {"path": "deck.pptx", "slide": 1, "shape_index": 0, "fill_color": "38BDF8"},
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="denied, ok")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_deny_preview") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "recolor that shape"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": False})
            _receive_until(ws, "agent_message")

    assert (tmp_path / "workspace" / "deck.pptx").read_bytes() == original_bytes


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed/functional in this environment",
)
def test_approval_preview_resolves_a_file_in_an_extra_writable_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """_build_document_edit_preview builds its own WorkspaceScope from
    self.settings.extra_readable_dirs/extra_writable_dirs, the same
    settings the real session's own tools use -- not just the plain
    workspace root. Confirms it actually resolves and renders a file
    that only exists via extra_writable_dirs (a directory outside the
    main workspace), the "file system is the core of this feature" case
    worth nailing explicitly rather than trusting by inspection alone --
    a resolve() bug here would silently degrade to (None, None) instead
    of erroring, so only a real render (non-None previews) proves it
    actually reached the file."""
    extra_dir = tmp_path / "shared_drive"
    _write_test_deck(extra_dir, name="external.pptx")
    call = _tool_call(
        "call_1",
        "edit_pptx_shape",
        {
            "path": str(extra_dir / "external.pptx"),
            "slide": 1,
            "shape_index": 0,
            "fill_color": "38BDF8",
        },
    )
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model, extra_writable_dirs=[extra_dir]) as client:
        with client.websocket_connect("/ws/t_extra_writable") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "recolor that shape"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            assert approval["before_preview"] is not None
            assert approval["after_preview"] is not None

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": False})
            _receive_until(ws, "agent_message")


def test_clear_wipes_conversation_history(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi there!")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_clear") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")

            ws.send_json({"type": "user_message", "text": "/clear"})
            cleared = ws.receive_json()
            assert cleared == {"type": "cleared"}

            # Behavioral proof the history is really gone, not just that the
            # event fired -- /compact's own "nothing much to compact" guard
            # (< 4 messages) only trips on a genuinely empty thread.
            ws.send_json({"type": "user_message", "text": "/compact"})
            error = ws.receive_json()

    assert error["type"] == "error"
    assert "Nothing much to compact" in error["message"]


def test_stop_mid_tool_closes_the_orphaned_tool_call_instead_of_breaking_the_thread(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Stop hard-cancels a turn while a tool is still running: the
    checkpoint is left with an AIMessage whose tool_call never got a
    ToolMessage. The next message must reach the model with that call
    closed off (OpenAI-compatible providers 400 on the whole history
    otherwise), and reconnecting must not silently re-run the tool."""
    from langchain_core.messages import ToolCall

    started = threading.Event()
    release = threading.Event()
    calls: list[str] = []

    def slow_search_pdf(self: Any, query: str, **kwargs: Any) -> list[dict[str, object]]:
        calls.append(query)
        started.set()
        release.wait(timeout=10)
        return []

    monkeypatch.setattr("coscribe.tools.documents.DocumentToolkit.search_pdf", slow_search_pdf)
    call = ToolCall(name="search_pdf", args={"query": "Error code"}, id="call_1")
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="hi")]
    )
    try:
        with _client_lg(tmp_path, monkeypatch, fake_model) as client:
            with client.websocket_connect("/ws/t_stop_mid_tool") as ws:
                ws.receive_json()  # state
                ws.receive_json()  # history
                ws.send_json({"type": "user_message", "text": "search the manual"})
                assert started.wait(timeout=5), "tool never started"
                ws.send_json({"type": "stop"})
                stopped = _receive_until(ws, "agent_message")[-1]
                assert stopped["text"] == "[stopped]"
                release.set()

            with client.websocket_connect("/ws/t_stop_mid_tool") as ws:
                ws.receive_json()  # state
                ws.receive_json()  # history
                ws.send_json({"type": "user_message", "text": "just say hi"})
                messages = _receive_until(ws, "tasks_changed")
    finally:
        release.set()

    assert calls == ["Error code"]
    assert not any(m["type"] == "error" for m in messages)
    assert [m for m in messages if m["type"] == "agent_message"][-1]["text"] == "hi"
    history = fake_model.received[-1]
    tool_call_index = next(
        i for i, m in enumerate(history) if isinstance(m, AIMessage) and m.tool_calls
    )
    closing = history[tool_call_index + 1]
    assert isinstance(closing, ToolMessage)
    assert closing.tool_call_id == "call_1"
    assert isinstance(history[tool_call_index + 2], HumanMessage)


def test_close_orphaned_tool_calls_heals_a_thread_already_corrupted_mid_history() -> None:
    """A thread broken before the fix has a newer HumanMessage *after*
    the orphaned tool call -- the placeholder has to be inserted directly
    after its AIMessage, not appended at the end."""
    from langchain.agents import create_agent
    from langchain_core.messages import ToolCall
    from langgraph.checkpoint.memory import InMemorySaver

    from coscribe.conversation.session import _close_orphaned_tool_calls

    def lookup(q: str) -> str:
        """Look something up."""
        return "found"

    calls = [
        ToolCall(name="lookup", args={"q": "a"}, id="done"),
        ToolCall(name="lookup", args={"q": "b"}, id="orphan"),
    ]
    agent = create_agent(
        FakeToolCallingChatModel(responses=[]), [lookup], checkpointer=InMemorySaver()
    )
    config = {"configurable": {"thread_id": "t"}}

    async def scenario() -> list[BaseMessage]:
        await agent.aupdate_state(
            config,
            {
                "messages": [
                    HumanMessage(content="go"),
                    AIMessage(content="", tool_calls=calls),
                    ToolMessage(content="found", tool_call_id="done", name="lookup"),
                    HumanMessage(content="second"),
                ]
            },
            as_node="tools",
        )
        assert await _close_orphaned_tool_calls(agent, config) is True
        assert await _close_orphaned_tool_calls(agent, config) is False
        state = await agent.aget_state(config)
        assert state.next == ()
        return list(state.values["messages"])

    messages = asyncio.run(scenario())
    assert [type(m).__name__ for m in messages] == [
        "HumanMessage",
        "AIMessage",
        "ToolMessage",
        "ToolMessage",
        "HumanMessage",
    ]
    assert [m.tool_call_id for m in messages[2:4]] == ["done", "orphan"]


def test_compact_collapses_history_and_survives_a_later_turn(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(content="hi there!"),
            AIMessage(content="nice to hear"),
            AIMessage(content="a short summary of the chat"),
            AIMessage(content="still here"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t4e") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")
            ws.send_json({"type": "user_message", "text": "how are you"})
            _receive_until(ws, "tasks_changed")

            ws.send_json({"type": "user_message", "text": "/compact"})
            compacted = ws.receive_json()
            assert compacted["type"] == "compacted"
            assert compacted["before"] == 4
            assert compacted["after"] == 1

            ws.send_json({"type": "user_message", "text": "are you still there?"})
            messages = _receive_until(ws, "tasks_changed")

    agent_message = next(m for m in messages if m["type"] == "agent_message")
    assert agent_message["text"] == "still here"


def test_load_older_messages_reveals_pre_compact_history(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Real, live-reported bug: after /compact, the chat log only ever
    showed the synthetic summary note -- no way to scroll up and see the
    real conversation, even though it was never actually deleted (see
    ChatSessionLG.load_older_messages' own docstring for the checkpoint
    mechanics). This is the fix's own regression test: the pre-compact
    turns come back verbatim on a "load_older_messages" request."""
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(content="hi there!"),
            AIMessage(content="nice to hear"),
            AIMessage(content="a short summary of the chat"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_older1") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")
            ws.send_json({"type": "user_message", "text": "how are you"})
            _receive_until(ws, "tasks_changed")

            ws.send_json({"type": "user_message", "text": "/compact"})
            compacted = ws.receive_json()
            assert compacted["type"] == "compacted"

            ws.send_json({"type": "load_older_messages"})
            older = ws.receive_json()
            # has_more is a "try again" hint, not a real lookahead (see
            # load_older_messages' own docstring) -- true here even
            # though this genuinely is the last epoch, confirmed instead
            # by a *second* call actually coming back empty below.
            ws.send_json({"type": "load_older_messages"})
            second_call = ws.receive_json()

    assert older["type"] == "older_messages"
    assert older["has_more"] is True
    texts = [(e["kind"], e.get("text")) for e in older["entries"]]
    assert texts == [
        ("user", "hello"),
        ("agent", "hi there!"),
        ("user", "how are you"),
        ("agent", "nice to hear"),
    ]
    assert second_call == {"type": "older_messages", "entries": [], "has_more": False}


def test_load_older_messages_is_empty_when_the_thread_was_never_compacted(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi there!")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_older2") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")

            ws.send_json({"type": "load_older_messages"})
            older = ws.receive_json()

    assert older == {"type": "older_messages", "entries": [], "has_more": False}


def test_edit_message_truncates_history_and_regenerates_from_the_edit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(content="hi there!"),
            AIMessage(content="nice to hear"),
            AIMessage(content="edited reply"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_edit1") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")
            ws.send_json({"type": "user_message", "text": "how are you"})
            _receive_until(ws, "tasks_changed")

            # index 1 -- the latest user turn -- edited. That turn onward
            # ("how are you", "nice to hear") is discarded, and the edited
            # text runs as a fresh turn; the earlier turn is untouched.
            ws.send_json({"type": "edit_message", "index": 1, "text": "how are you, edited"})
            messages = _receive_until(ws, "tasks_changed")

        # Reconnecting re-reads straight from the checkpointer -- proves
        # the truncation is real, not just a client-side illusion.
        with client.websocket_connect("/ws/t_edit1") as ws:
            ws.receive_json()  # state
            history = ws.receive_json()

    agent_message = next(m for m in messages if m["type"] == "agent_message")
    assert agent_message["text"] == "edited reply"
    assert history["entries"] == [
        {"kind": "user", "text": "hello"},
        {"kind": "agent", "text": "hi there!"},
        {"kind": "user", "text": "how are you, edited"},
        {"kind": "agent", "text": "edited reply"},
    ]


def test_task_list_endpoint_reflects_task_create(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    call = _tool_call("call_1", "task_create", {"content": "write the report"})
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(content="", tool_calls=[call]),
            AIMessage(content="tracked it"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t5") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "please plan this out"})
            _receive_until(ws, "tasks_changed")

        response = client.get("/api/threads/t5/tasks")

    assert response.status_code == 200
    tasks = response.json()
    assert len(tasks) == 1
    assert tasks[0]["content"] == "write the report"
    assert tasks[0]["status"] == "pending"


def test_get_tools_lists_builtin_tools_with_metadata(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/tools")

    assert response.status_code == 200
    tools = {tool["name"]: tool for tool in response.json()["tools"]}
    assert tools["read_docx"]["category"] == "documents"
    assert tools["write_xlsx"]["requires_approval"] is True
    assert tools["task_create"]["category"] == "tasks"
    assert tools["web_search"]["category"] == "web"
    assert tools["web_search"]["requires_approval"] is False
    assert tools["add_xlsx_chart"]["category"] == "documents"
    assert tools["add_xlsx_chart"]["requires_approval"] is True
    assert tools["recalc_xlsx"]["category"] == "documents"
    assert tools["recalc_xlsx"]["requires_approval"] is True
    assert tools["format_xlsx_cells"]["category"] == "documents"
    assert tools["format_xlsx_cells"]["requires_approval"] is True
    assert tools["add_pptx_chart"]["category"] == "documents"
    assert tools["add_pptx_chart"]["requires_approval"] is True
    assert tools["add_pptx_image"]["category"] == "documents"
    assert tools["add_pptx_image"]["requires_approval"] is True
    assert tools["set_pptx_notes"]["category"] == "documents"
    assert tools["set_pptx_notes"]["requires_approval"] is True
    assert tools["set_pptx_transition"]["category"] == "documents"
    assert tools["set_pptx_transition"]["requires_approval"] is True
    assert tools["add_pptx_animation"]["category"] == "documents"
    assert tools["add_pptx_animation"]["requires_approval"] is True
    assert tools["add_pptx_hyperlink"]["category"] == "documents"
    assert tools["add_pptx_hyperlink"]["requires_approval"] is True
    assert tools["edit_pptx_theme_colors"]["category"] == "documents"
    assert tools["edit_pptx_theme_colors"]["requires_approval"] is True
    assert tools["run_python_script"]["category"] == "scripts"
    assert tools["run_python_script"]["risk_category"] == "EXEC"
    assert tools["run_python_script"]["requires_approval"] is True
    assert tools["run_node_script"]["category"] == "scripts"
    assert tools["run_node_script"]["risk_category"] == "EXEC"
    assert tools["run_node_script"]["requires_approval"] is True
    assert tools["set_pptx_background_image"]["category"] == "documents"
    assert tools["set_pptx_background_image"]["requires_approval"] is True
    assert tools["edit_pptx_text"]["category"] == "documents"
    assert tools["edit_pptx_text"]["requires_approval"] is True
    assert tools["delete_pptx_slide"]["category"] == "documents"
    assert tools["delete_pptx_slide"]["requires_approval"] is True
    assert tools["duplicate_pptx_slide"]["category"] == "documents"
    assert tools["duplicate_pptx_slide"]["requires_approval"] is True
    assert tools["reorder_pptx_slide"]["category"] == "documents"
    assert tools["reorder_pptx_slide"]["requires_approval"] is True
    assert tools["edit_file"]["category"] == "filesystem"
    assert tools["edit_file"]["requires_approval"] is True
    assert tools["edit_file_batch"]["category"] == "filesystem"
    assert tools["edit_file_batch"]["requires_approval"] is True
    assert tools["get_file_info"]["category"] == "filesystem"
    assert tools["get_file_info"]["requires_approval"] is False
    assert tools["delete_file"]["category"] == "filesystem"
    assert tools["delete_file"]["requires_approval"] is True
    assert tools["move_file"]["category"] == "filesystem"
    assert tools["move_file"]["requires_approval"] is True
    assert tools["copy_file"]["category"] == "filesystem"
    assert tools["copy_file"]["requires_approval"] is True
    assert tools["list_pptx_shapes"]["category"] == "documents"
    assert tools["list_pptx_shapes"]["requires_approval"] is False
    assert tools["edit_pptx_shape"]["category"] == "documents"
    assert tools["edit_pptx_shape"]["requires_approval"] is True
    assert tools["delete_pptx_shape"]["category"] == "documents"
    assert tools["delete_pptx_shape"]["requires_approval"] is True
    assert tools["replace_pptx_image"]["category"] == "documents"
    assert tools["replace_pptx_image"]["requires_approval"] is True
    assert tools["list_pptx_icons"]["category"] == "documents"
    assert tools["list_pptx_icons"]["requires_approval"] is False
    assert tools["add_pptx_icon"]["category"] == "documents"
    assert tools["add_pptx_icon"]["requires_approval"] is True
    assert tools["recolor_pptx_icon"]["category"] == "documents"
    assert tools["recolor_pptx_icon"]["requires_approval"] is True
    assert tools["edit_pptx_table_cell"]["category"] == "documents"
    assert tools["edit_pptx_table_cell"]["requires_approval"] is True
    assert tools["merge_pptx_table_cells"]["category"] == "documents"
    assert tools["merge_pptx_table_cells"]["requires_approval"] is True
    assert tools["search_images"]["category"] == "web"
    assert tools["search_images"]["requires_approval"] is False
    assert tools["download_image"]["category"] == "web"
    assert tools["download_image"]["requires_approval"] is True
    assert tools["add_pptx_scrim"]["category"] == "documents"
    assert tools["add_pptx_scrim"]["requires_approval"] is True
    assert tools["sleep_until"]["category"] == "selfwake"
    assert tools["sleep_until"]["requires_approval"] is True
    assert tools["list_wakes"]["category"] == "selfwake"
    assert tools["list_wakes"]["requires_approval"] is False


def test_upload_writes_file_to_workspace_and_returns_relative_path(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.post(
            "/api/upload",
            files={"file": ("report.pdf", b"%PDF-1.4 fake content", "application/pdf")},
        )

    assert response.status_code == 200
    body = response.json()
    assert body["path"] == "report.pdf"
    assert body["bytes_written"] == len(b"%PDF-1.4 fake content")
    assert (tmp_path / "workspace" / "report.pdf").read_bytes() == b"%PDF-1.4 fake content"


def test_an_upload_and_a_slide_overlay_use_the_conversations_own_folder_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Real bug: with a folder chosen for a conversation, an attached file
    landed in the app's default workspace -- the path the model was given
    pointed nowhere in its folder -- and the click-a-shape overlay couldn't
    find a deck that lived there."""
    from pptx import Presentation

    chosen = tmp_path / "project-folder"
    chosen.mkdir()
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5])
    prs.save(str(chosen / "deck.pptx"))
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect(f"/ws/t_own_folder?workspace={chosen}") as ws:
            ws.receive_json()
            ws.receive_json()
            uploaded = client.post(
                "/api/upload",
                data={"thread_id": "t_own_folder"},
                files={"file": ("brief.txt", b"hello", "text/plain")},
            ).json()
            shapes = client.get(
                "/api/pptx-shapes",
                params={"path": "deck.pptx", "slide": 1, "thread_id": "t_own_folder"},
            )
            without_thread = client.get(
                "/api/pptx-shapes", params={"path": "deck.pptx", "slide": 1}
            )

    assert uploaded["path"] == "brief.txt"
    assert (chosen / "brief.txt").read_bytes() == b"hello"
    assert not (tmp_path / "workspace" / "brief.txt").exists()
    assert shapes.status_code == 200 and shapes.json()["slide_width_in"] > 0
    assert without_thread.status_code == 400


def test_upload_auto_renames_on_name_collision(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        first = client.post("/api/upload", files={"file": ("notes.txt", b"first", "text/plain")})
        second = client.post("/api/upload", files={"file": ("notes.txt", b"second", "text/plain")})

    assert first.json()["path"] == "notes.txt"
    assert second.json()["path"] == "notes (1).txt"
    workspace = tmp_path / "workspace"
    assert (workspace / "notes.txt").read_bytes() == b"first"
    assert (workspace / "notes (1).txt").read_bytes() == b"second"


def test_upload_sanitizes_path_traversal_attempt_in_filename(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.post(
            "/api/upload", files={"file": ("../../evil.txt", b"data", "text/plain")}
        )

    assert response.status_code == 200
    assert response.json()["path"] == "evil.txt"
    assert (tmp_path / "workspace" / "evil.txt").read_bytes() == b"data"
    assert not (tmp_path / "evil.txt").exists()


def test_upload_rejects_oversized_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("coscribe.web.routes.files.MAX_UPLOAD_BYTES", 10)
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.post(
            "/api/upload", files={"file": ("big.txt", b"0123456789 -- too big", "text/plain")}
        )

    assert response.status_code == 413
    assert not (tmp_path / "workspace" / "big.txt").exists()


def test_get_preview_serves_an_existing_preview_png(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    previews_dir = tmp_path / "state" / "previews"
    previews_dir.mkdir(parents=True)
    name = "0123456789abcdef0123456789abcdef.png"
    (previews_dir / name).write_bytes(b"\x89PNG\r\n\x1a\n fake png bytes")

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get(f"/api/previews/{name}")

    assert response.status_code == 200
    assert response.headers["content-type"] == "image/png"
    assert response.content == b"\x89PNG\r\n\x1a\n fake png bytes"


def test_get_preview_404s_for_unknown_name(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/previews/0123456789abcdef0123456789abcdef.png")

    assert response.status_code == 404


def test_get_preview_rejects_path_traversal_attempts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Proves the name-shape check, not just a missing-file 404 -- a name
    # that doesn't match render_thumbnail's uuid4().hex pattern is rejected
    # before any filesystem lookup, so it can never escape state_dir/previews/.
    secret = tmp_path / "state" / "secret.txt"
    secret.parent.mkdir(parents=True)
    secret.write_text("do not serve me")

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/previews/..%2Fsecret.txt")

    assert response.status_code == 404


def test_get_pptx_shapes_returns_shapes_and_slide_dimensions(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Backs the click-a-shape-in-the-preview feature (ChatLog.tsx's
    PptxShapeOverlay) -- a plain UI-facing REST read, not a tool call."""
    _write_test_deck(tmp_path / "workspace")

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/pptx-shapes", params={"path": "deck.pptx", "slide": 1})

    assert response.status_code == 200
    body = response.json()
    assert body["slide"] == 1
    assert body["shape_count"] == 1
    assert body["slide_width_in"] > 0
    assert body["slide_height_in"] > 0
    [shape] = body["shapes"]
    assert shape["index"] == 0
    assert shape["left_in"] == pytest.approx(1.0)
    assert shape["top_in"] == pytest.approx(1.0)
    assert shape["width_in"] == pytest.approx(2.0)
    assert shape["height_in"] == pytest.approx(1.0)
    assert shape["text_preview"] == "hello"


def test_get_pptx_shapes_400s_for_a_missing_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/pptx-shapes", params={"path": "nope.pptx", "slide": 1})

    assert response.status_code == 400


def test_get_pptx_shapes_400s_for_an_out_of_range_slide(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _write_test_deck(tmp_path / "workspace")

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/pptx-shapes", params={"path": "deck.pptx", "slide": 5})

    assert response.status_code == 400


def test_get_pptx_shapes_400s_for_a_path_outside_the_workspace(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    outside = tmp_path / "outside.pptx"
    Presentation().save(str(outside))

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/pptx-shapes", params={"path": str(outside), "slide": 1})

    assert response.status_code == 400


def test_get_commands_includes_plan_accept_edits_and_compact(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/commands")

    assert response.status_code == 200
    names = {command["name"] for command in response.json()}
    assert {"init", "plan", "accept-edits", "compact"} <= names


def test_unknown_slash_command_is_still_rejected_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_slash_unknown") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "/nope do something"})
            error = ws.receive_json()

    assert error["type"] == "error"
    assert "Unknown command: /nope" in error["message"]
    assert fake_model.i == 0


def test_an_attached_pdf_can_be_previewed_page_by_page_and_nothing_else_can(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import io

    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    for number in (1, 2):
        pdf.drawString(100, 700, f"page {number}")
        pdf.showPage()
    pdf.save()

    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        uploaded = client.post(
            "/api/upload", files={"file": ("cv.pdf", buffer.getvalue(), "application/pdf")}
        ).json()
        text = client.post("/api/upload", files={"file": ("a.txt", b"hi", "text/plain")}).json()
        info = client.get("/api/attachment/pdf", params={"path": uploaded["path"]})
        page = client.get(
            "/api/attachment/pdf/page", params={"path": uploaded["path"], "page": 2, "width": 300}
        )
        past_end = client.get(
            "/api/attachment/pdf/page", params={"path": uploaded["path"], "page": 3}
        )
        not_pdf = client.get("/api/attachment/pdf", params={"path": text["path"]})
        outside = client.get("/api/attachment/pdf", params={"path": "../../etc/passwd"})

    assert info.json() == {"pages": 2}
    assert page.headers["content-type"] == "image/png" and page.content.startswith(b"\x89PNG")
    assert past_end.status_code == 404
    assert not_pdf.status_code == 404
    assert outside.status_code in (400, 404)


def test_workspace_query_param_resolves_workspace_on_connect(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    chosen = tmp_path / "project-a"
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect(f"/ws/t_ws_qs?workspace={chosen}") as ws:
            state = ws.receive_json()
            ws.receive_json()  # history

    assert state["workspace_root"] == str(chosen)
    assert state["workspace_explicit"] is True
    assert state["folders"] == [str(chosen)]
    sidecar = tmp_path / "state" / "t_ws_qs.workspace"
    assert json.loads(sidecar.read_text(encoding="utf-8")) == [str(chosen)]


def test_new_thread_without_workspace_falls_back_to_settings_workspace_root(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_ws_default") as ws:
            state = ws.receive_json()
            ws.receive_json()  # history

    assert state["workspace_root"] == str(tmp_path / "workspace")
    assert state["workspace_explicit"] is False
    assert not (tmp_path / "state" / "t_ws_default.workspace").exists()


def test_workspace_choice_persists_across_reconnect_without_the_query_param(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    chosen = tmp_path / "project-b"
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as first_process:
        with first_process.websocket_connect(f"/ws/t_ws_persist?workspace={chosen}") as ws:
            ws.receive_json()
            ws.receive_json()  # history

    # Fresh create_app_lg() -- empty in-memory sessions dict, same tmp_path
    # on disk -- so a reconnect with no ?workspace= this time can only pick
    # the choice back up from the sidecar, not in-process state.
    with _client_lg(tmp_path, monkeypatch, fake_model) as second_process:
        with second_process.websocket_connect("/ws/t_ws_persist") as ws:
            state = ws.receive_json()
            ws.receive_json()  # history

    assert state["workspace_root"] == str(chosen)


def test_set_folders_changes_a_conversations_folders_any_time(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    first = tmp_path / "project-first"
    second = tmp_path / "project-second"
    first.mkdir()
    second.mkdir()
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_ws_folders") as ws:
            ws.receive_json()  # state (falls back to settings.workspace_root)
            ws.receive_json()  # history

            ws.send_json({"type": "set_folders", "folders": [str(first)]})
            one = ws.receive_json()
            ws.send_json({"type": "set_folders", "folders": [str(first), str(second)]})
            two = ws.receive_json()
            ws.send_json({"type": "set_folders", "folders": [str(second)]})
            swapped = ws.receive_json()

    assert one["workspace_root"] == str(first)
    assert one["workspace_explicit"] is True
    assert two["folders"] == [str(first), str(second)]
    assert swapped["workspace_root"] == str(second)
    assert swapped["folders"] == [str(second)]
    sidecar = tmp_path / "state" / "t_ws_folders.workspace"
    assert json.loads(sidecar.read_text(encoding="utf-8")) == [str(second)]


def test_removing_every_folder_returns_to_the_default_workspace(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    chosen = tmp_path / "project-gone"
    chosen.mkdir()
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect(f"/ws/t_ws_clear?workspace={chosen}") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "set_folders", "folders": []})
            cleared = ws.receive_json()

    assert cleared["folders"] == []
    assert cleared["workspace_explicit"] is False
    assert cleared["workspace_root"] == str(tmp_path / "workspace")
    assert not (tmp_path / "state" / "t_ws_clear.workspace").exists()


def test_a_folder_that_does_not_exist_is_refused_and_nothing_changes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    kept = tmp_path / "project-kept"
    kept.mkdir()
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect(f"/ws/t_ws_bad?workspace={kept}") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            missing = tmp_path / "missing"
            ws.send_json({"type": "set_folders", "folders": [str(kept), str(missing)]})
            refused = ws.receive_json()

    assert refused["type"] == "error"
    assert "isn't an existing folder" in refused["message"]
    sidecar = tmp_path / "state" / "t_ws_bad.workspace"
    assert json.loads(sidecar.read_text(encoding="utf-8")) == [str(kept)]


def test_a_single_path_sidecar_from_before_multiple_folders_still_loads(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    chosen = tmp_path / "project-legacy"
    (tmp_path / "state").mkdir(parents=True, exist_ok=True)
    (tmp_path / "state" / "t_ws_legacy.workspace").write_text(str(chosen), encoding="utf-8")
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_ws_legacy") as ws:
            state = ws.receive_json()
            ws.receive_json()  # history

    assert state["folders"] == [str(chosen)]


def test_the_agent_can_write_into_a_conversations_second_folder(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    main = tmp_path / "project-main"
    extra = tmp_path / "project-extra"
    main.mkdir()
    extra.mkdir()
    call = _tool_call("call_1", "write_file", {"path": str(extra / "note.txt"), "content": "hi"})
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_ws_extra") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "set_folders", "folders": [str(main), str(extra)]})
            ws.receive_json()  # state
            ws.send_json({"type": "user_message", "text": "write hi into the second folder"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")

    assert (extra / "note.txt").read_text() == "hi"


def test_delete_thread_removes_workspace_sidecar_too(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    chosen = tmp_path / "project-d"
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect(f"/ws/t_ws_delete?workspace={chosen}") as ws:
            ws.receive_json()
            ws.receive_json()  # history
        sidecar = tmp_path / "state" / "t_ws_delete.workspace"
        assert sidecar.is_file()

        client.delete("/api/threads/t_ws_delete")

        assert not sidecar.is_file()


def test_write_file_lands_in_the_threads_own_workspace_not_the_global_default(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """End-to-end proof that a per-thread workspace_root actually reaches
    the tool layer (coordinator.py's build_coordinator_agent -- not just
    that the WS "state" event reports the right string, which the other
    workspace tests above already cover)."""
    custom_workspace = tmp_path / "custom-project"
    custom_workspace.mkdir()
    call = _tool_call("call_1", "write_file", {"path": "note.txt", "content": "hi"})
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(content="", tool_calls=[call]),
            AIMessage(content="done"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect(f"/ws/t_ws_write?workspace={custom_workspace}") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "write hi to note.txt"})

            approval = ws.receive_json()
            assert approval["type"] == "approval_required"

            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")

    assert (custom_workspace / "note.txt").read_text() == "hi"
    assert not (tmp_path / "workspace" / "note.txt").exists()


def test_get_threads_endpoint_includes_workspace_root(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    chosen = tmp_path / "project-e"
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="hi"), AIMessage(content="hi again")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect(f"/ws/t_ws_list_a?workspace={chosen}") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")

        with client.websocket_connect("/ws/t_ws_list_b") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello again"})
            _receive_until(ws, "tasks_changed")

        threads = {t["thread_id"]: t for t in client.get("/api/threads").json()}

    assert threads["t_ws_list_a"]["workspace_root"] == str(chosen)
    assert threads["t_ws_list_b"]["workspace_root"] == str(tmp_path / "workspace")


def test_get_threads_endpoint_lists_saved_threads_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Regression test for a real, live-reported bug: list_threads used to
    glob settings.state_dir for *.json files -- the old runtime's storage
    shape, not this one's. runtime_lg persists conversation history in the
    shared AsyncSqliteSaver checkpointer, so that glob always came back
    empty and the session-switcher UI never showed any history, even
    though every thread's real state was sitting right there in the
    checkpoint database."""
    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_list1") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")

        response = client.get("/api/threads")

    assert response.status_code == 200
    threads = {t["thread_id"]: t for t in response.json()}
    assert "t_list1" in threads
    entry = threads["t_list1"]
    assert entry["preview"] == "hello"
    assert entry["message_count"] == 2  # the human "hello" + the AI's "hi" reply
    assert entry["updated_at"]  # a real ISO timestamp, not asserting the exact value


def test_get_threads_endpoint_strips_mode_note_and_sorts_by_recency_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Covers strip_mode_note (a live turn's mode-note prefix, e.g.
    "[normal mode: ...] ", shouldn't leak into the preview text a user
    never actually typed) and the most-recently-updated-first sort."""
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="hi"), AIMessage(content="hi again")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_older") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "first thread"})
            _receive_until(ws, "tasks_changed")

        with client.websocket_connect("/ws/t_newer") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "second thread"})
            _receive_until(ws, "tasks_changed")

        threads = client.get("/api/threads").json()

    assert [t["thread_id"] for t in threads] == ["t_newer", "t_older"]
    newer = threads[0]
    assert not newer["preview"].startswith("[")
    assert "second thread" in newer["preview"]


def test_delete_thread_endpoint_removes_checkpoints_and_tasks_sidecar_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from coscribe.tools.tasks import TaskToolkit

    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_del1") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")

        TaskToolkit("t_del1", tmp_path / "state").create("write the report")
        assert (tmp_path / "state" / "t_del1.tasks.json").is_file()
        thread_ids = {t["thread_id"] for t in client.get("/api/threads").json()}
        assert "t_del1" in thread_ids

        response = client.delete("/api/threads/t_del1")

        assert response.status_code == 200
        assert response.json() == {"deleted": "t_del1"}
        assert not (tmp_path / "state" / "t_del1.tasks.json").exists()
        thread_ids = {t["thread_id"] for t in client.get("/api/threads").json()}
        assert "t_del1" not in thread_ids


def test_delete_thread_endpoint_unknown_thread_404s_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.delete("/api/threads/nope")

    assert response.status_code == 404


def test_rename_thread_endpoint_overrides_preview_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_rename1") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "original first message"})
            _receive_until(ws, "tasks_changed")

        before = next(t for t in client.get("/api/threads").json() if t["thread_id"] == "t_rename1")
        assert "original first message" in before["preview"]

        response = client.post(
            "/api/threads/t_rename1/rename", json={"title": "Quarterly report draft"}
        )

        assert response.status_code == 200
        assert response.json() == {"thread_id": "t_rename1", "title": "Quarterly report draft"}
        after = next(t for t in client.get("/api/threads").json() if t["thread_id"] == "t_rename1")
        assert after["preview"] == "Quarterly report draft"


def test_rename_thread_endpoint_unknown_thread_404s_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.post("/api/threads/nope/rename", json={"title": "New title"})

    assert response.status_code == 404


def test_rename_thread_endpoint_rejects_blank_title_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_rename2") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")

        response = client.post("/api/threads/t_rename2/rename", json={"title": "   "})

    assert response.status_code == 400


def test_delete_thread_endpoint_removes_title_sidecar_too(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_del_title") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")

        client.post("/api/threads/t_del_title/rename", json={"title": "renamed"})
        assert (tmp_path / "state" / "t_del_title.title").is_file()

        response = client.delete("/api/threads/t_del_title")

        assert response.status_code == 200
        assert not (tmp_path / "state" / "t_del_title.title").exists()


def test_browse_dirs_lists_subdirectories_only_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Regression test for a real, live-reported bug: app.py never
    defined /api/browse-dirs at all (unlike web/app.py), even though the
    shared app.js frontend's Settings-panel folder picker calls it
    unconditionally -- clicking "select folder" 404'd and the picker
    modal opened with nothing in it, with no visible error."""
    (tmp_path / "Downloads").mkdir()
    (tmp_path / "Documents").mkdir()
    (tmp_path / "not_a_dir.txt").write_text("x", encoding="utf-8")
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/browse-dirs", params={"path": str(tmp_path)})

    assert response.status_code == 200
    body = response.json()
    resolved = tmp_path.resolve()
    assert body["path"] == str(resolved)
    assert body["directories"] == [
        {"name": "Documents", "path": str(resolved / "Documents")},
        {"name": "Downloads", "path": str(resolved / "Downloads")},
        # create_app_lg() auto-creates settings.skills_dir at startup --
        # the skills_by_name dict built for GET /api/skills/select_skills
        # validation calls load_skills(settings.skills_dir), which
        # creates it if missing.
        {"name": "skills", "path": str(resolved / "skills")},
        # The lifespan also opens the AsyncSqliteSaver checkpointer under
        # settings.state_dir (tmp_path/state), creating that directory too
        # -- app.py's equivalent test has no state_dir-creating lifespan
        # step, so this entry is specific to runtime_lg.
        {"name": "state", "path": str(resolved / "state")},
    ]
    assert body["parent"] == str(resolved.parent)


def test_browse_dirs_defaults_to_home_when_no_path_given_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/browse-dirs")

    assert response.status_code == 200
    assert response.json()["path"] == str(Path.home().resolve())


def test_browse_dirs_nonexistent_path_returns_error_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/browse-dirs", params={"path": str(tmp_path / "does-not-exist")})

    assert response.status_code == 200
    assert "error" in response.json()


def test_browse_dirs_file_path_returns_error_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    file_path = tmp_path / "a.txt"
    file_path.write_text("x", encoding="utf-8")
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        response = client.get("/api/browse-dirs", params={"path": str(file_path)})

    assert response.status_code == 200
    assert "error" in response.json()


def test_history_is_empty_for_a_brand_new_thread_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_hist_new") as ws:
            ws.receive_json()  # state
            history = ws.receive_json()

    assert history == {"type": "history", "entries": [], "has_older": False}


def test_history_reports_has_older_true_after_reconnecting_to_a_compacted_thread(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Real, live-reported bug: the frontend's "load earlier messages"
    control used to be offered unconditionally on *every* thread,
    including a brand-new one with nothing to page through -- this
    `has_older` flag is what the frontend now gates that control on.
    Regression-tests the O(1) proxy itself: a thread that's been
    /compact'd, then reconnected to (a fresh connection, so send_history
    -- not any state left over from the compacting connection -- is what
    has to get this right), reports has_older=True."""
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(content="hi there!"),
            AIMessage(content="nice to hear"),
            AIMessage(content="a short summary of the chat"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_hist_has_older") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history -- empty, brand new thread
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")
            ws.send_json({"type": "user_message", "text": "how are you"})
            _receive_until(ws, "tasks_changed")
            ws.send_json({"type": "user_message", "text": "/compact"})
            assert ws.receive_json()["type"] == "compacted"

        with client.websocket_connect("/ws/t_hist_has_older") as ws:
            ws.receive_json()  # state
            history = ws.receive_json()

    assert history["has_older"] is True


def test_reconnect_replays_conversation_history_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Regression test for a real, live-reported bug: the chat log was
    only ever built from *live* events during the current WS connection,
    so switching to (or reconnecting to) an existing thread always showed
    a blank pane even though the model still remembered the whole
    conversation -- see serialize_history_for_ws_lg's docstring."""
    call = _tool_call("call_1", "read_file", {"path": "note.txt"})
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(content="", tool_calls=[call]),
            AIMessage(content="the file says hello"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_hist_reconnect") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history -- empty, brand new thread
            ws.send_json({"type": "user_message", "text": "what does note.txt say?"})
            _receive_until(ws, "tasks_changed")

        # Reconnect -- a fresh WS connection to the same thread, exactly
        # what switching sessions or reloading the page does.
        with client.websocket_connect("/ws/t_hist_reconnect") as ws:
            ws.receive_json()  # state
            history = ws.receive_json()

    assert history == {
        "type": "history",
        "entries": [
            {"kind": "user", "text": "what does note.txt say?"},
            {
                "kind": "tool",
                "tool_name": "read_file",
                "arguments": {"path": "note.txt"},
                "result": "File does not exist: note.txt",
                "is_error": True,
            },
            {"kind": "agent", "text": "the file says hello"},
        ],
        "has_older": False,
    }


def test_history_omits_a_call_still_pending_approval_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A tool call awaiting approval has no ToolMessage result yet --
    serialize_history_for_ws_lg must skip it rather than show a
    phantom step with no outcome; resume_after_reconnect already
    redelivers the live approval_required event for it separately."""
    call = _tool_call("call_1", "write_file", {"path": "note.txt", "content": "hi"})
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_hist_pending") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "write hi to note.txt"})
            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            # Disconnect without ever answering -- the call stays pending.

        with client.websocket_connect("/ws/t_hist_pending") as ws:
            ws.receive_json()  # state
            history = ws.receive_json()

    assert history == {
        "type": "history",
        "entries": [{"kind": "user", "text": "write hi to note.txt"}],
        "has_older": False,
    }


def test_history_replay_includes_an_approved_calls_real_result_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Regression test for a real, live-reported bug: reconnecting to (or
    reloading) an existing thread showed every past tool call as click-to-
    expand, but expanding one showed nothing -- ToolCallRow's own `result
    !== undefined` guard always failed on a replayed item, because
    serialize_history_for_ws_lg computed each call's real result (into
    results_by_id, to decide whether to include the entry at all) and then
    silently dropped it instead of putting it on the emitted entry. Uses
    an *approved* call specifically (not the plain read_file case the
    sibling reconnect test above already covers), since that's the
    real-world shape reported live -- an approval-gated run_command/
    write-file call whose result vanished on reload."""
    call = _tool_call("call_1", "write_file", {"path": "note.txt", "content": "hi"})
    fake_model = FakeToolCallingChatModel(
        responses=[AIMessage(content="", tool_calls=[call]), AIMessage(content="done")]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/t_hist_approved_result") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "write hi to note.txt"})
            approval = ws.receive_json()
            assert approval["type"] == "approval_required"
            ws.send_json({"type": "approval_response", "id": approval["id"], "approved": True})
            _receive_until(ws, "tasks_changed")

        with client.websocket_connect("/ws/t_hist_approved_result") as ws:
            ws.receive_json()  # state
            history = ws.receive_json()

    tool_entry = next(e for e in history["entries"] if e["kind"] == "tool")
    assert tool_entry["result"] == {
        "path": "note.txt",
        "bytes_written": 2,
        "lines_added": 1,
        "lines_removed": 0,
    }
    assert tool_entry["is_error"] is False


def test_thread_activity_lists_outputs_references_tools_and_progress_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _stub_script_env(monkeypatch)
    (tmp_path / "workspace").mkdir()
    (tmp_path / "workspace" / "input.txt").write_text("data", encoding="utf-8")
    script = "open('chart.png', 'w').write('png')"
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(
                content="",
                tool_calls=[
                    _tool_call("c1", "task_create", {"content": "Read the input"}),
                    _tool_call("c2", "read_file", {"path": "input.txt"}),
                    _tool_call("c3", "list_files", {"path": "."}),
                ],
            ),
            AIMessage(
                content="",
                tool_calls=[
                    _tool_call("c4", "write_file", {"path": "report.md", "content": "# R"}),
                    _tool_call("c5", "write_file", {"path": "../outside.md", "content": "x"}),
                ],
            ),
            # Its own turn: a script's files_written is a before/after
            # snapshot of the workspace, so a write running alongside it
            # would be counted as the script's too.
            AIMessage(
                content="",
                tool_calls=[
                    _tool_call("c6", "run_python_script", {"script": script, "description": "d"}),
                ],
            ),
            AIMessage(content="done"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        _run_turn(client, "t_activity", "make a report")
        activity = client.get("/api/threads/t_activity/activity").json()

    assert [t["content"] for t in activity["tasks"]] == ["Read the input"]
    # Newest first; the write outside the workspace failed, so it isn't one.
    assert [o["path"] for o in activity["outputs"]] == ["chart.png", "report.md"]
    assert all(o["exists"] and o["openable"] for o in activity["outputs"])
    assert [r["path"] for r in activity["references"]] == ["input.txt"]
    assert {t["name"]: t["count"] for t in activity["tools"]} == {
        "read_file": 1,
        "list_files": 1,
        "write_file": 1,
        "run_python_script": 1,
    }
    assert activity["connectors"] == [] and activity["skills"] == []


def test_thread_activity_marks_files_created_edited_and_read_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    (workspace / "notes.txt").write_text("keep", encoding="utf-8")
    (workspace / "plan.md").write_text("old line", encoding="utf-8")
    (workspace / "draft.md").write_text("old", encoding="utf-8")
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(
                content="",
                tool_calls=[
                    _tool_call("c1", "read_file", {"path": "notes.txt"}),
                    _tool_call("c2", "read_file", {"path": "plan.md"}),
                ],
            ),
            AIMessage(
                content="",
                tool_calls=[
                    _tool_call(
                        "c3", "edit_file", {"path": "plan.md", "old_text": "old", "new_text": "new"}
                    ),
                    _tool_call("c4", "write_file", {"path": "draft.md", "content": "replaced"}),
                    _tool_call("c5", "write_file", {"path": "summary.md", "content": "# S"}),
                ],
            ),
            AIMessage(
                content="",
                tool_calls=[
                    _tool_call("c6", "write_file", {"path": "summary.md", "content": "# S2"})
                ],
            ),
            AIMessage(content="done"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        _run_turn(client, "t_activity_actions", "tidy up")
        activity = client.get("/api/threads/t_activity_actions/activity").json()

    # summary.md was made here, so it stays "created" after being rewritten;
    # draft.md existed (its old line was removed) and plan.md was read first.
    assert {o["path"]: o["action"] for o in activity["outputs"]} == {
        "summary.md": "created",
        "draft.md": "edited",
        "plan.md": "edited",
    }
    assert [(r["path"], r["action"]) for r in activity["references"]] == [("notes.txt", "read")]


def test_thread_activity_leaves_scratch_folders_out_of_the_outputs_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A run that checks its work leaves renders and test files behind;
    those aren't what the person asked for."""
    (tmp_path / "workspace").mkdir()
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(
                content="",
                tool_calls=[
                    _tool_call("c1", "write_file", {"path": "report.md", "content": "# R"}),
                    _tool_call("c2", "write_file", {"path": "qa/page-1.md", "content": "x"}),
                    _tool_call("c3", "write_file", {"path": ".cache/x.md", "content": "x"}),
                    _tool_call("c4", "write_file", {"path": "out/qa.md", "content": "x"}),
                ],
            ),
            AIMessage(content="done"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        _run_turn(client, "t_activity_scratch", "write")
        activity = client.get("/api/threads/t_activity_scratch/activity").json()

    assert sorted(o["path"] for o in activity["outputs"]) == ["out/qa.md", "report.md"]


def test_thread_activity_leaves_out_files_the_assistant_made_and_removed_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    (tmp_path / "workspace").mkdir()
    fake_model = FakeToolCallingChatModel(
        responses=[
            AIMessage(
                content="",
                tool_calls=[
                    _tool_call("c1", "write_file", {"path": "report.md", "content": "# R"}),
                    _tool_call("c2", "write_file", {"path": "check.png", "content": "x"}),
                ],
            ),
            AIMessage(
                content="", tool_calls=[_tool_call("c3", "delete_file", {"path": "check.png"})]
            ),
            AIMessage(content="done"),
        ]
    )
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        _run_turn(client, "t_activity_removed", "write")
        activity = client.get("/api/threads/t_activity_removed/activity").json()

    assert [o["path"] for o in activity["outputs"]] == ["report.md"]


def test_opening_a_thread_file_only_hands_documents_to_the_os_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    opened: list[tuple[str, bool]] = []
    monkeypatch.setattr(
        "coscribe.web.routes.threads.open_in_os",
        lambda path, *, reveal: opened.append((path.name, reveal)),
    )
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    (workspace / "report.xlsx").write_bytes(b"x")
    (workspace / "run.bat").write_text("echo hi", encoding="utf-8")
    (tmp_path / "secret.txt").write_text("s", encoding="utf-8")
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        url = "/api/threads/t_files/files/open"
        doc = client.post(url, json={"path": "report.xlsx"})
        script = client.post(url, json={"path": "run.bat"})
        revealed = client.post(url, json={"path": "run.bat", "reveal": True})
        outside = client.post(url, json={"path": "../secret.txt"})
        download = client.get("/api/threads/t_files/files/download", params={"path": "report.xlsx"})
        missing = client.get("/api/threads/t_files/files/download", params={"path": "nope.xlsx"})

    assert doc.status_code == 200 and script.status_code == 400
    assert revealed.status_code == 200 and outside.status_code == 404
    assert opened == [("report.xlsx", False), ("run.bat", True)]
    assert download.status_code == 200 and download.content == b"x"
    assert "report.xlsx" in download.headers["content-disposition"]
    assert missing.status_code == 404


def test_a_saved_run_waits_for_a_folder_it_hasnt_been_allowed_then_carries_on_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    (tmp_path / "workspace").mkdir()
    elsewhere = tmp_path / "elsewhere"
    elsewhere.mkdir()
    workflow = {
        "inputs": [{"name": "dest"}],
        "steps": [
            {"id": "save", "kind": "tool", "title": "Save the report", "tool": "write_file",
             "args": {"path": "{{dest}}", "content": "report"}},
        ],
    }  # fmt: skip
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        created = client.post(
            "/api/scheduled-tasks",
            json={"name": "Report", "kind": "manual", "at": "", "workflow": workflow},
        )
        task = created.json()
        target = str(elsewhere / "out.txt")
        started = client.post(
            f"/api/scheduled-tasks/{task['trigger_id']}/run", json={"inputs": {"dest": target}}
        )
        run = started.json()["run"]
        parked = _wait_for_run_status(client, task["trigger_id"], run["run_id"])
        declined_early = (elsewhere / "out.txt").exists()

        answered = client.post(
            f"/api/scheduled-tasks/{task['trigger_id']}/runs/{run['run_id']}/answer",
            json={"approved": True},
        )
        finished = _wait_for_run_status(client, task["trigger_id"], run["run_id"])
        saved = client.get("/api/scheduled-tasks").json()
        permissions = client.post(
            "/api/workflows/permissions",
            json={"workflow": workflow, "trigger_id": task["trigger_id"]},
        ).json()

    waiting = parked["steps"][-1]
    assert parked["status"] == "needs_approval"
    assert waiting["status"] == "waiting"
    assert waiting["permission"]["kind"] == "folder"
    assert waiting["permission"]["target"] == str(elsewhere)
    assert not declined_early
    assert answered.status_code == 200
    assert finished["status"] == "completed", finished
    assert (elsewhere / "out.txt").read_text(encoding="utf-8") == "report"
    granted = next(t for t in saved if t["trigger_id"] == task["trigger_id"])
    assert granted["permissions"] == {"folders": [str(elsewhere)]}
    assert permissions["permissions"]["folders"] == [str(elsewhere)]


class _TitlingModel(FakeToolCallingChatModel):
    """Answers the naming request on its own, so that it neither takes a
    scripted reply nor depends on whether it comes before or after the
    turn's own call."""

    def _generate(
        self,
        messages: list[BaseMessage],
        stop: list[str] | None = None,
        run_manager: Any = None,
        **kwargs: Any,
    ) -> ChatResult:
        if "Name this conversation" in str(messages[0].content):
            self.received.append(list(messages))
            return ChatResult(
                generations=[ChatGeneration(message=AIMessage(content='"Q3 销售汇总"'))]
            )
        return super()._generate(messages, stop, run_manager, **kwargs)


def test_a_conversation_is_named_from_its_first_message_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = _TitlingModel(responses=[AIMessage(content="Here's the summary.")])
    with _client_lg(tmp_path, monkeypatch, fake_model, auto_title_threads=True) as client:
        with client.websocket_connect("/ws/t_named") as ws:
            ws.receive_json()  # state
            ws.receive_json()  # history
            ws.send_json({"type": "user_message", "text": "帮我汇总一下第三季度的销售数据"})
            titled = _receive_until(ws, "thread_titled")[-1]
            _receive_until(ws, "tasks_changed")
        threads = client.get("/api/threads").json()

    assert titled == {"type": "thread_titled", "title": "Q3 销售汇总"}
    [thread] = [t for t in threads if t["thread_id"] == "t_named"]
    assert thread["preview"] == "Q3 销售汇总"
    [title_request] = [
        str(m[-1].content) for m in fake_model.received if "Name this" in str(m[0].content)
    ]
    # Asked for before there is a reply, so a long first turn is not unnamed.
    assert title_request == "User: 帮我汇总一下第三季度的销售数据"


def test_deleting_a_conversation_removes_its_environment_lg(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    (tmp_path / ".env").write_text("", encoding="utf-8")
    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/envthread") as ws:
            _receive_until(ws, "history")
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "agent_message")
        client.put(
            "/api/threads/envthread/environment", json={"variables": {"A": "1"}, "secrets": []}
        )
        assert (tmp_path / "state" / "envthread.env.json").exists()
        client.delete("/api/threads/envthread")

    assert not (tmp_path / "state" / "envthread.env.json").exists()


def test_messaging_policy_round_trips_and_refuses_a_bad_one(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        assert client.get("/api/threads/b/messaging").json() == {"mode": "off", "senders": []}

        saved = client.put("/api/threads/b/messaging", json={"mode": "selected", "senders": ["a"]})
        refused = client.put("/api/threads/b/messaging", json={"mode": "everyone", "senders": []})

        assert saved.json() == {"mode": "selected", "senders": ["a"]}
        assert client.get("/api/threads/b/messaging").json() == saved.json()
        assert refused.status_code == 422


def test_deleting_a_conversation_forgets_who_may_message_it(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_model = FakeToolCallingChatModel(responses=[AIMessage(content="hi")])
    with _client_lg(tmp_path, monkeypatch, fake_model) as client:
        with client.websocket_connect("/ws/b") as ws:
            ws.receive_json()
            ws.receive_json()
            ws.send_json({"type": "user_message", "text": "hello"})
            _receive_until(ws, "tasks_changed")
        client.put("/api/threads/b/messaging", json={"mode": "any", "senders": []})

        client.delete("/api/threads/b")

        assert client.get("/api/threads/b/messaging").json()["mode"] == "off"
