import concurrent.futures
import shutil
import tempfile
from pathlib import Path
from typing import Any

import pytest

from coscribe.tools.presentations import (
    PresentationToolkit,
    _check_low_contrast,
    _check_missing_visual_elements,
    _check_text_overlaps,
    build_presentation_tools,
)


def _tools_by_name(root: Path) -> dict[str, object]:
    return {tool.__name__: tool for tool in build_presentation_tools(root)}  # type: ignore[attr-defined]


def _libreoffice_actually_works() -> bool:
    """`shutil.which("soffice")` alone isn't enough to gate the real
    overflow-detection test -- some environments have the soffice binary on
    PATH but it fails to convert anything (broken font config, sandboxing,
    etc.), which would make that test fail rather than skip. Probe with an
    actual trivial write_pptx call and check its own QA result, in a
    throwaway temp dir (skipif conditions run at collection time, before
    the tmp_path fixture exists)."""
    probe_dir = Path(tempfile.mkdtemp(prefix="coscribe_lo_probe_"))
    try:
        result = PresentationToolkit(probe_dir).write_pptx(
            path="probe.pptx", content="# Probe\n- hi"
        )
        return result["qa_skipped_reason"] is None
    finally:
        shutil.rmtree(probe_dir, ignore_errors=True)


TWO_SLIDE_CONTENT = """\
# Intro
- point one
- point two

---

# Data
| name | score |
| --- | --- |
| Alice | 9 |
| Bob | 7 |
"""


def test_write_pptx_then_read_pptx_round_trips_title_and_bullets(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# My Title\n- first\n- second")

    text = tools["read_pptx"](path="deck.pptx")

    assert "## My Title" in text
    assert "- first" in text
    assert "- second" in text


def test_pptx_edits_to_the_same_file_concurrently_do_not_corrupt_it(tmp_path: Path) -> None:
    """Same fix, same reasoning as spreadsheets_tool's identical
    concurrency regression test: any tool here that loads an existing
    .pptx, edits it, and saves it back (add_pptx_chart/set_pptx_notes/
    etc. -- unlike write_pptx itself, which always starts from a blank
    Presentation()) races if two calls target the same file at once,
    since LangGraph runs one AIMessage's tool_calls concurrently on
    separate threads. Fixed via tools/_file_locks.py's per-path lock
    (@locked_by_path); this drives 3 concurrent set_pptx_notes calls
    against 3 different slides of the same deck and asserts all 3
    notes survive with no corruption."""
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](
        path="deck.pptx", content="# One\n- a\n\n---\n\n# Two\n- b\n\n---\n\n# Three\n- c"
    )

    def set_notes(slide: int) -> dict[str, object]:
        return tools["set_pptx_notes"](path="deck.pptx", slide=slide, notes=f"note-{slide}")  # type: ignore[operator]

    with concurrent.futures.ThreadPoolExecutor(max_workers=3) as executor:
        results = list(executor.map(set_notes, [1, 2, 3]))

    assert all("slide" in result for result in results)
    from pptx import Presentation

    prs = Presentation(str(tmp_path / "deck.pptx"))
    assert len(prs.slides) == 3
    notes = {slide.notes_slide.notes_text_frame.text for slide in prs.slides}
    assert notes == {"note-1", "note-2", "note-3"}


def test_write_pptx_default_canvas_is_real_16x9_widescreen(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)


def test_write_pptx_default_canvas_title_placeholder_keeps_its_inherited_height(
    tmp_path: Path,
) -> None:
    # Live-reproduced defect: widening a stock layout's placeholder by
    # setting only .left/.width made python-pptx materialize a brand-new
    # <a:xfrm> with .top/.height defaulting to 0 (silently discarding the
    # values inherited from the slide master), rendering the title clipped
    # to a zero-height box. Both dimensions must survive intact.
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    title = prs.slides[0].shapes.title
    assert title.top > 0
    assert title.height > 0


def test_write_pptx_bold_and_italic_markers_become_real_formatting(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](
        path="deck.pptx",
        content="# Title\nThis is **bold** and *italic* and plain text.",
    )

    prs = Presentation(str(tmp_path / "deck.pptx"))
    body_frame = prs.slides[0].placeholders[1].text_frame
    runs = body_frame.paragraphs[0].runs
    run_texts = [run.text for run in runs]

    assert "".join(run_texts) == "This is bold and italic and plain text."
    assert not any("*" in text for text in run_texts)
    bold_runs = [run for run in runs if run.font.bold]
    italic_runs = [run for run in runs if run.font.italic]
    assert [run.text for run in bold_runs] == ["bold"]
    assert [run.text for run in italic_runs] == ["italic"]


def test_multi_slide_round_trip_and_single_slide_read(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    result = tools["write_pptx"](path="deck.pptx", content=TWO_SLIDE_CONTENT)

    assert result["slide_count"] == 2

    full_text = tools["read_pptx"](path="deck.pptx")
    assert "## Intro" in full_text
    assert "## Data" in full_text
    assert "---" in full_text

    slide_two = tools["read_pptx"](path="deck.pptx", slide=2)
    assert "## Data" not in slide_two  # no heading prefix for a single-slide read
    assert "| Alice | 9 |" in slide_two
    assert "point one" not in slide_two


def test_table_only_slide_round_trips(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](
        path="deck.pptx",
        content="# Data\n| a | b |\n| --- | --- |\n| 1 | 2 |",
    )

    text = tools["read_pptx"](path="deck.pptx")

    assert "| a | b |" in text
    assert "| 1 | 2 |" in text


def test_mixing_table_and_bullets_on_one_slide_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)

    with pytest.raises(ValueError, match="table or bullets"):
        tools["write_pptx"](
            path="deck.pptx",
            content="# Mixed\n- a bullet\n| a | b |\n| --- | --- |\n| 1 | 2 |",
        )


def test_read_pptx_out_of_range_slide_names_real_count(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content=TWO_SLIDE_CONTENT)

    with pytest.raises(ValueError, match="2 slides"):
        tools["read_pptx"](path="deck.pptx", slide=5)


def test_write_pptx_skips_qa_gracefully_when_soffice_missing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("shutil.which", lambda name: None)
    tools = _tools_by_name(tmp_path)

    result = tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    assert result["overflow_warnings"] == []
    assert result["qa_skipped_reason"]
    # The file itself still got written despite QA being unavailable.
    assert (tmp_path / "deck.pptx").is_file()


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed or not functional in this environment",
)
def test_write_pptx_detects_real_overflow_via_libreoffice(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)

    normal = tools["write_pptx"](path="normal.pptx", content="# Short\n- a short bullet")
    assert normal["overflow_warnings"] == []
    assert normal["qa_skipped_reason"] is None

    huge_bullet = " ".join(["overflow"] * 400)
    overflowing = tools["write_pptx"](
        path="overflowing.pptx", content=f"# Too Much\n- {huge_bullet}"
    )
    assert overflowing["overflow_warnings"] != []


def test_write_pptx_skips_preview_without_state_dir(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)  # no state_dir passed

    result = tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    assert result["preview_path"] is None
    assert result["preview_skipped_reason"] == "no state_dir configured"


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed or not functional in this environment",
)
def test_write_pptx_generates_real_preview_via_libreoffice(tmp_path: Path) -> None:
    state_dir = tmp_path / "state"
    tools = {
        tool.__name__: tool
        for tool in build_presentation_tools(tmp_path / "workspace", state_dir=state_dir)  # type: ignore[attr-defined]
    }

    result = tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    assert result["preview_skipped_reason"] is None
    preview_path = state_dir / "previews" / result["preview_path"]
    assert preview_path.is_file()
    assert preview_path.stat().st_size > 0


def test_render_pptx_preview_skips_gracefully_without_state_dir(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)  # no state_dir passed
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    result = tools["render_pptx_preview"](path="deck.pptx")

    assert result["preview_paths"] == []
    assert result["preview_paths_csv"] == ""
    assert result["preview_skipped_reason"] == "no state_dir configured"


def test_render_pptx_preview_skips_gracefully_when_pdftoppm_missing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    state_dir = tmp_path / "state"
    tools = {
        tool.__name__: tool
        for tool in build_presentation_tools(tmp_path / "workspace", state_dir=state_dir)  # type: ignore[attr-defined]
    }
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    # Fakes soffice as *present* (not just falling through to the real
    # shutil.which) -- this test's whole point is isolating "only pdftoppm
    # is missing," and letting soffice's presence depend on whatever
    # happens to be installed on the machine running the test made this a
    # real, deterministic CI failure: GitHub Actions' runner has neither
    # tool installed, so the real shutil.which("soffice") call returned
    # None too, and render_all_page_previews' own soffice check (which
    # runs *before* its pdftoppm check) reported "LibreOffice (soffice)
    # not found" instead of the pdftoppm-specific reason this test
    # actually asserts on.
    real_which = shutil.which

    def fake_which(name: str) -> str | None:
        if name == "pdftoppm":
            return None
        if name == "soffice":
            return "/usr/bin/soffice"
        return real_which(name)

    monkeypatch.setattr("shutil.which", fake_which)
    result = tools["render_pptx_preview"](path="deck.pptx")

    assert result["preview_paths"] == []
    assert result["preview_skipped_reason"] == "poppler-utils (pdftoppm) not found"


def _poppler_and_libreoffice_actually_work() -> bool:
    if shutil.which("pdftoppm") is None:
        return False
    return _libreoffice_actually_works()


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _poppler_and_libreoffice_actually_work(),
    reason="LibreOffice or poppler-utils not installed/functional in this environment",
)
def test_render_pptx_preview_renders_every_slide_via_libreoffice(tmp_path: Path) -> None:
    state_dir = tmp_path / "state"
    tools = {
        tool.__name__: tool
        for tool in build_presentation_tools(tmp_path / "workspace", state_dir=state_dir)  # type: ignore[attr-defined]
    }
    tools["write_pptx"](path="deck.pptx", content=TWO_SLIDE_CONTENT)

    result = tools["render_pptx_preview"](path="deck.pptx")

    assert result["preview_skipped_reason"] is None
    assert len(result["preview_paths"]) == 2
    assert result["preview_paths_csv"] == ",".join(result["preview_paths"])
    for name in result["preview_paths"]:
        preview_path = state_dir / "previews" / name
        assert preview_path.is_file()
        assert preview_path.stat().st_size > 0


def test_render_pptx_preview_flags_missing_visual_elements_without_state_dir_or_libreoffice(
    tmp_path: Path,
) -> None:
    # These two fields are pure python-pptx checks -- no soffice/pdftoppm
    # needed, so they must still come back even when preview rendering is
    # entirely skipped (no state_dir passed here at all).
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- plain bullet, no visuals")

    result = tools["render_pptx_preview"](path="deck.pptx")

    assert result["preview_skipped_reason"] == "no state_dir configured"
    assert result["slides_missing_visual_elements"] == [1]
    assert result["text_overlap_warnings"] == []


def test_check_missing_visual_elements_ignores_slide_background_color(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()

    # A themed slide (dark background) with only unfilled text boxes --
    # background color alone must not count as a visual element.
    plain_slide = prs.slides.add_slide(prs.slide_layouts[6])
    plain_slide.background.fill.solid()
    plain_slide.background.fill.fore_color.rgb = RGBColor(0x1E, 0x27, 0x61)
    text_box = plain_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    text_box.text_frame.text = "Just a title, no shapes"

    # A slide with a real colored (filled) shape -- must count.
    decorated_slide = prs.slides.add_slide(prs.slide_layouts[6])
    icon = decorated_slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(1), Inches(1), Inches(0.5), Inches(0.5)
    )
    icon.fill.solid()
    icon.fill.fore_color.rgb = RGBColor(0xFF, 0x00, 0x00)

    assert _check_missing_visual_elements(prs) == [1]


def test_check_missing_visual_elements_counts_a_decorative_shape_inherited_from_the_layout(
    tmp_path: Path,
) -> None:
    """Real, hand-authored templates commonly draw decorative background
    art (accent bars, curtain motifs) once on the slide layout/master
    rather than copying it onto every slide -- confirmed live against
    coscribe's own bundled "velis" template, whose curtain-motif graphics
    live as GROUP shapes on the slide master and were false-positively
    flagged as "missing visuals" before _slide_has_visual_element learned
    to also check the slide's own layout and that layout's master."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    layout = prs.slide_layouts[6]  # "Blank" -- no shapes of its own by default
    # LayoutShapes has no public add_shape (only SlideShapes does) -- build
    # the shape via a throwaway slide, then move its element onto the
    # layout's own shape tree, matching this module's established
    # "no public API, drop the element directly" idiom (see
    # _remove_placeholder in scripts/build_pptx_templates.py).
    scratch_slide = prs.slides.add_slide(layout)
    decorative = scratch_slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(0.15)
    )
    decorative.fill.solid()
    decorative.fill.fore_color.rgb = RGBColor(0x1E, 0x27, 0x61)
    decorative._element.getparent().remove(decorative._element)  # noqa: SLF001
    layout.shapes._spTree.append(decorative._element)  # noqa: SLF001
    xml_slides = prs.slides._sldIdLst  # noqa: SLF001
    xml_slides.remove(list(xml_slides)[0])

    slide = prs.slides.add_slide(layout)
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    text_box.text_frame.text = "Title only, no shapes of its own"

    assert _check_missing_visual_elements(prs) == []


def test_check_low_contrast_flags_light_text_on_a_matching_light_background(
    tmp_path: Path,
) -> None:
    """The common real-world mistake this exists to catch: light gray body
    text on a white/near-white background, both picked from the same
    palette so neither reads as an obvious mistake in isolation."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "Barely readable"
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    run.font.size = Pt(16)  # ordinary body text -- 4.5:1 threshold applies

    warnings = _check_low_contrast(prs)
    assert len(warnings) == 1
    assert warnings[0]["slide"] == 1
    assert warnings[0]["text"] == "Barely readable"
    assert warnings[0]["contrast_ratio"] < 4.5


def test_check_low_contrast_uses_the_looser_ratio_for_large_bold_text(tmp_path: Path) -> None:
    """WCAG's own large-text carve-out (>=18pt, or >=14pt bold, gets 3.0:1
    instead of 4.5:1) -- a color pair that fails the small-text threshold
    (~3.95:1) but clears the large-text one must NOT be flagged for large
    text, and MUST be flagged for the same pair at ordinary body size."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    large = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    large.text_frame.text = "Large bold headline"
    large_run = large.text_frame.paragraphs[0].runs[0]
    large_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)  # ~3.95:1
    large_run.font.size = Pt(20)
    large_run.font.bold = True

    small = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    small.text_frame.text = "Same color, small body text"
    small_run = small.text_frame.paragraphs[0].runs[0]
    small_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    small_run.font.size = Pt(16)

    warnings = _check_low_contrast(prs)
    assert [w["text"] for w in warnings] == ["Same color, small body text"]


def test_check_low_contrast_skips_what_it_cannot_resolve(tmp_path: Path) -> None:
    """Deliberately conservative: a gradient-filled background, a
    theme/scheme font color, and a run with no explicit color at all must
    all be silently skipped rather than guessed at -- false positives on
    content this check has no real color data for would erode trust in
    every other (real) hit it reports."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.gradient()  # no resolvable solid color anywhere

    gradient_bg_text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    gradient_bg_text.text_frame.text = "On a gradient background"
    gradient_run = gradient_bg_text.text_frame.paragraphs[0].runs[0]
    gradient_run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    gradient_run.font.size = Pt(16)

    # Give this shape its own solid (white) fill so the background side
    # resolves -- only the font side (a theme color) is under test here.
    scheme_box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    scheme_box.fill.solid()
    scheme_box.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    scheme_box.text_frame.text = "Theme-colored text"
    scheme_run = scheme_box.text_frame.paragraphs[0].runs[0]
    scheme_run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
    scheme_run.font.size = Pt(16)

    assert _check_low_contrast(prs) == []


def test_check_low_contrast_resolves_background_from_the_slide_when_shape_has_no_fill(
    tmp_path: Path,
) -> None:
    """A plain textbox (the common case -- write_pptx never fills its own
    text boxes) has to fall back to the slide's own background color, not
    just skip because the shape itself has no fill."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))  # no fill of its own
    box.text_frame.text = "Inherits the slide background"
    run = box.text_frame.paragraphs[0].runs[0]
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    run.font.size = Pt(16)

    warnings = _check_low_contrast(prs)
    assert [w["text"] for w in warnings] == ["Inherits the slide background"]


def test_check_text_overlaps_flags_significantly_overlapping_text_boxes(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    subtitle = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(0.5))
    subtitle.text_frame.text = "Precision Monitoring at Your Fingertips"
    body = slide.shapes.add_textbox(Inches(1), Inches(1.6), Inches(4), Inches(2))
    body.text_frame.text = "Real-Time Tracking: identify power-hungry appliances instantly."

    warnings = _check_text_overlaps(prs)

    assert len(warnings) == 1
    assert warnings[0]["slide"] == 1
    assert "Precision Monitoring" in warnings[0]["text_a"]
    assert "Real-Time Tracking" in warnings[0]["text_b"]


def test_check_text_overlaps_flags_text_overlapping_a_chart(tmp_path: Path) -> None:
    """Real, live-reproduced bug this catches: a chart is a `GraphicFrame`,
    not a text-frame shape, so it was invisible to the text-vs-text-only
    scan even when its own rendered content (axis labels, a chart title)
    visually collided with a nearby placeholder's text -- confirmed via
    add_pptx_chart's own old fixed-position default landing directly on
    top of a slide's body text."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(2))
    body.text_frame.text = "Some real intro text on this slide"

    chart_data = CategoryChartData()  # type: ignore[no-untyped-call]
    chart_data.categories = ["A", "B"]
    chart_data.add_series("S", (1, 2))  # type: ignore[no-untyped-call]
    slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, Inches(1), Inches(1.6), Inches(6), Inches(4), chart_data
    )

    warnings = _check_text_overlaps(prs)

    assert len(warnings) == 1
    assert warnings[0]["slide"] == 1
    assert "Some real intro text" in warnings[0]["text_a"]


def test_check_text_overlaps_does_not_flag_a_small_incidental_overlap(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1))
    title.text_frame.text = "Title"
    # Body starts just barely inside the title's box -- a sliver of
    # incidental overlap, not a real layout bug.
    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(6), Inches(3))
    body.text_frame.text = "Body content well below the title"

    assert _check_text_overlaps(prs) == []


# --- Prefab layout library (layout: icon-list / stat-callout / two-column) ---


def test_layout_icon_list_renders_glyphs_labels_and_stays_clean(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    content = (
        "layout: icon-list FF0000\n"
        "# Feature Highlights\n"
        "- [*] Real-time energy tracking\n"
        "- Mobile app control\n"
    )
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    slide = prs.slides[0]
    icon_texts = [
        shape.text_frame.text for shape in slide.shapes if shape.shape_type == 1  # AUTO_SHAPE
    ]
    label_texts = [
        shape.text_frame.text for shape in slide.shapes if shape.shape_type == 17  # TEXT_BOX
    ]
    assert icon_texts == ["*", "2"]
    assert label_texts == ["Real-time energy tracking", "Mobile app control"]
    assert _check_text_overlaps(prs) == []
    assert _check_missing_visual_elements(prs) == []


def test_layout_omitted_produces_the_same_slide_as_before(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- first\n- second")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    slide = prs.slides[0]
    body_text = slide.placeholders[1].text_frame.text
    assert body_text == "first\nsecond"
    assert len(slide.shapes) == 2  # title + body placeholder only, no new shapes


def test_layout_icon_list_rejects_too_many_items(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    bullets = "\n".join(f"- item {i}" for i in range(7))
    content = f"layout: icon-list\n# Title\n{bullets}\n"

    with pytest.raises(ValueError, match="at most 6 items"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_icon_list_rejects_a_table(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "layout: icon-list\n# Title\n| a | b |\n| --- | --- |\n| 1 | 2 |\n"

    with pytest.raises(ValueError, match="doesn't support tables"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_icon_list_rejects_a_spelled_out_word_glyph(tmp_path: Path) -> None:
    # Live-reproduced defect: a real model wrote "- [dash] Rising costs..."
    # (a plain English word, not a symbol), which rendered as ugly wrapped
    # text inside the small icon circle.
    tools = _tools_by_name(tmp_path)
    content = "layout: icon-list\n# Title\n- [dash] Rising costs\n"

    with pytest.raises(ValueError, match="spelled-out word"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_icon_list_rejects_an_emoji_glyph(tmp_path: Path) -> None:
    # Live-reproduced defect: an emoji glyph (e.g. a "house" or "rocket"
    # character) renders via the font's own built-in color glyph, ignoring
    # the icon circle's accent-color fill -- some icons come out clean,
    # others as a mismatched full-color sticker, inconsistently.
    tools = _tools_by_name(tmp_path)
    content = "layout: icon-list\n# Title\n- [\U0001f680] Launch\n"

    with pytest.raises(ValueError, match="isn't plain ASCII"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_icon_list_allows_short_initials_as_a_glyph(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    content = "layout: icon-list\n# Title\n- [AI] Smart recommendations\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    icon_texts = [
        shape.text_frame.text for shape in prs.slides[0].shapes if shape.shape_type == 1
    ]
    assert icon_texts == ["AI"]


def test_layout_stat_callout_renders_cards_and_stays_clean(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    content = (
        "layout: stat-callout 028090\n"
        "# Q4 Highlights\n"
        "| Stat | Label |\n"
        "| --- | --- |\n"
        "| 3x | Revenue growth |\n"
        "| 12 | New markets |\n"
        "| 98% | Retention |\n"
    )
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    card_texts = sorted(
        shape.text_frame.text
        for shape in prs.slides[0].shapes
        if shape.shape_type == 1  # AUTO_SHAPE
    )
    # label first, then the emphasized stat -- see _add_stat_callout_slide's
    # own docstring (PPTX_DESIGN.md §0) for why this reads label-then-value
    # now instead of value-then-label.
    assert card_texts == ["New markets\n12", "Retention\n98%", "Revenue growth\n3x"]
    assert _check_text_overlaps(prs) == []
    assert _check_missing_visual_elements(prs) == []


def test_layout_stat_callout_rejects_bullets_instead_of_a_table(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "layout: stat-callout\n# Title\n- not a table\n"

    with pytest.raises(ValueError, match="exactly one pipe-table"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_stat_callout_rejects_too_many_rows(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    rows = "\n".join(f"| {i} | label {i} |" for i in range(5))
    content = f"layout: stat-callout\n# Title\n| Stat | Label |\n| --- | --- |\n{rows}\n"

    with pytest.raises(ValueError, match="at most 4 stats"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_stat_callout_rejects_swapped_columns(tmp_path: Path) -> None:
    """Real, live-reproduced mistake: writing the table as
    '| label | value |' (which reads more naturally than '| value |
    label |') silently rendered the long label large/bold/accent-colored
    and the short value small/muted -- the exact opposite of the intended
    hierarchy -- with no error, only a visible defect in the rendered
    deck. Must now raise instead."""
    tools = _tools_by_name(tmp_path)
    content = (
        "layout: stat-callout\n"
        "# Highlights\n"
        "| Label | Value |\n"
        "| --- | --- |\n"
        "| Quarterly revenue growth | +18.4% |\n"
        "| Net revenue retention | 121% |\n"
        "| Customer satisfaction NPS | 62 |\n"
    )

    with pytest.raises(ValueError, match="columns look swapped"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_stat_callout_cards_are_vertically_centered(tmp_path: Path) -> None:
    """Real, live-reported defect caught in a rendered screenshot: cards
    anchored to the top of the content area left a large empty gap for
    the rest of the slide below them, since card_height is deliberately
    much shorter than the full content area (see _add_stat_callout_slide's
    own comment). Cards should split the leftover space above/below
    instead of dumping all of it below."""
    from pptx import Presentation

    from coscribe.tools.presentations import _content_area

    tools = _tools_by_name(tmp_path)
    content = (
        "layout: stat-callout\n"
        "# Highlights\n"
        "| Stat | Label |\n"
        "| --- | --- |\n"
        "| 3x | Revenue growth |\n"
    )
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(tmp_path / "deck.pptx")
    _left, content_top, _width, _height = _content_area(prs)
    card = next(s for s in prs.slides[0].shapes if s.shape_type == 1)
    assert card.top > content_top


def test_layout_two_column_splits_content_at_marker(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    content = (
        "layout: two-column\n"
        "# Before vs After\n"
        "- Manual tracking\n"
        "- No alerts\n"
        ">>>\n"
        "- Real-time dashboards\n"
        "- Proactive alerts\n"
    )
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    slide = prs.slides[0]
    assert slide.shapes.title.text == "Before vs After"
    assert slide.placeholders[1].text_frame.text == "Manual tracking\nNo alerts"
    assert slide.placeholders[2].text_frame.text == "Real-time dashboards\nProactive alerts"


def test_layout_two_column_requires_the_marker(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "layout: two-column\n# Title\n- only one column\n"

    with pytest.raises(ValueError, match="'>>>'"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_two_column_rejects_a_table_on_either_side(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "layout: two-column\n# Title\n| a | b |\n| --- | --- |\n| 1 | 2 |\n>>>\n- right\n"

    with pytest.raises(ValueError, match="doesn't support tables"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_two_column_flags_missing_visual_element_by_design(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    content = "layout: two-column\n# Title\n- left\n>>>\n- right\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    # Documented exception: two-column is plain text by construction (the
    # stock "Two Content" layout has no shape of its own), so this flags
    # unless the model separately adds a chart/image -- intended, not a bug.
    assert _check_missing_visual_elements(prs) == [1]
    assert _check_text_overlaps(prs) == []


def test_layout_unknown_id_raises_naming_valid_ids(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "layout: bogus-id\n# Title\n- body\n"

    with pytest.raises(
        ValueError, match="Unknown layout.*icon-list.*stat-callout.*svg.*two-column"
    ):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_two_column_with_undersized_template_raises(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="template.pptx", content="# Old\n- old bullet")

    # Same technique as test_write_pptx_with_template_missing_required_layouts_raises:
    # strip every layout but the first, leaving fewer than write_pptx needs
    # -- index 3 ("Two Content") is well within what gets stripped here.
    prs = Presentation(tmp_path / "template.pptx")
    layout_master = prs.slide_masters[0]
    sld_layout_id_lst = layout_master.element.find(qn("p:sldLayoutIdLst"))
    for layout_id in list(sld_layout_id_lst)[1:]:
        r_id = layout_id.get(qn("r:id"))
        layout_master.part.drop_rel(r_id)
        sld_layout_id_lst.remove(layout_id)
    prs.save(tmp_path / "template.pptx")

    content = "layout: two-column\n# Title\n- left\n>>>\n- right\n"
    with pytest.raises(ValueError, match="slide layout"):
        tools["write_pptx"](
            path="deck.pptx", content=content, template_path="template.pptx"
        )


def test_layout_svg_renders_background_circle_and_text_and_stays_clean(
    tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="960" height="540" viewBox="0 0 960 540">'
        '<rect x="0" y="0" width="960" height="540" fill="#1E2761"/>'
        '<circle cx="120" cy="270" r="40" fill="FF0000"/>'
        '<text x="220" y="260" font-size="36" fill="#FFFFFF">Real-time tracking</text>'
        '<text x="60" y="80" font-size="48" fill="#FFFFFF">Aurora</text>'
        "</svg>"
    )
    content = f"layout: svg\n{svg_markup}\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    slide = prs.slides[0]
    # The full-canvas rect at index 0 is promoted to the slide background
    # rather than becoming its own shape -- only the circle and two text
    # boxes remain as real shapes.
    assert len(slide.shapes) == 3
    assert slide.background.fill.fore_color.rgb == RGBColor.from_string("1E2761")
    text_shapes = [s for s in slide.shapes if s.has_text_frame and s.shape_type == 17]
    assert sorted(s.text_frame.text for s in text_shapes) == ["Aurora", "Real-time tracking"]
    assert _check_text_overlaps(prs) == []
    assert _check_missing_visual_elements(prs) == []


def test_layout_svg_rx_selects_rounded_rectangle_without_rx_selects_rectangle(
    tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="960" height="540" viewBox="0 0 960 540">'
        '<rect x="10" y="10" width="200" height="80" fill="FF0000"/>'
        '<rect x="10" y="120" width="200" height="80" rx="12" fill="00FF00"/>'
        "</svg>"
    )
    tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    shapes = list(prs.slides[0].shapes)
    assert len(shapes) == 2
    assert shapes[0].auto_shape_type == MSO_SHAPE.RECTANGLE
    assert shapes[1].auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE


def test_layout_svg_accepts_hex_and_named_colors_rejects_unknown_name(
    tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="960" height="540" viewBox="0 0 960 540">'
        '<rect x="10" y="10" width="200" height="80" fill="FF0000"/>'
        '<rect x="10" y="120" width="200" height="80" fill="#00FF00"/>'
        '<rect x="10" y="230" width="200" height="80" fill="navy"/>'
        "</svg>"
    )
    tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    fills = [s.fill.fore_color.rgb for s in prs.slides[0].shapes]
    assert fills == [
        RGBColor.from_string("FF0000"),
        RGBColor.from_string("00FF00"),
        RGBColor.from_string("000080"),
    ]

    bad_svg = (
        '<svg width="100" height="100" viewBox="0 0 100 100">'
        '<rect x="0" y="0" width="50" height="50" fill="notacolor"/>'
        "</svg>"
    )
    with pytest.raises(ValueError, match="not a 6-hex-digit color"):
        tools["write_pptx"](path="deck2.pptx", content=f"layout: svg\n{bad_svg}\n")


def test_layout_svg_rejects_unsupported_element_naming_run_node_script(
    tmp_path: Path,
) -> None:
    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="100" height="100" viewBox="0 0 100 100">'
        '<image href="photo.png" x="0" y="0" width="100" height="100"/>'
        "</svg>"
    )
    with pytest.raises(ValueError, match="run_node_script"):
        tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")


def test_layout_svg_path_straight_lines_renders_a_freeform_shape(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="200" height="200" viewBox="0 0 200 200">'
        '<path d="M10,10 L190,10 L100,190 Z" fill="FF0000"/>'
        "</svg>"
    )
    tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    shapes = list(prs.slides[0].shapes)
    assert len(shapes) == 1
    assert shapes[0].shape_type == MSO_SHAPE_TYPE.FREEFORM
    assert shapes[0].fill.fore_color.rgb == RGBColor.from_string("FF0000")
    assert _check_missing_visual_elements(prs) == []


def test_layout_svg_path_bezier_curves_render_without_error(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="200" height="200" viewBox="0 0 200 200">'
        '<path d="M20,100 C20,20 180,20 180,100 Q180,180 100,180 Z" fill="00FF00"/>'
        "</svg>"
    )
    tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    shape = next(iter(prs.slides[0].shapes))
    assert shape.shape_type == MSO_SHAPE_TYPE.FREEFORM
    # Bounding box spans the full extent of the curve's control points too
    # (a safe overestimate, not a pixel-tight fit -- see module docstring).
    assert shape.width > 0
    assert shape.height > 0


def test_layout_svg_path_relative_commands_match_absolute_equivalent(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    absolute_svg = (
        '<svg width="200" height="200" viewBox="0 0 200 200">'
        '<path d="M10,10 L100,10 L100,100 Z" fill="0000FF"/>'
        "</svg>"
    )
    relative_svg = (
        '<svg width="200" height="200" viewBox="0 0 200 200">'
        '<path d="m10,10 l90,0 l0,90 z" fill="0000FF"/>'
        "</svg>"
    )
    tools["write_pptx"](path="abs.pptx", content=f"layout: svg\n{absolute_svg}\n")
    tools["write_pptx"](path="rel.pptx", content=f"layout: svg\n{relative_svg}\n")

    abs_shape = next(iter(Presentation(str(tmp_path / "abs.pptx")).slides[0].shapes))
    rel_shape = next(iter(Presentation(str(tmp_path / "rel.pptx")).slides[0].shapes))
    assert (abs_shape.left, abs_shape.top, abs_shape.width, abs_shape.height) == (
        rel_shape.left,
        rel_shape.top,
        rel_shape.width,
        rel_shape.height,
    )


def test_layout_svg_path_implicit_repeat_commands_add_extra_segments(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    # "L100,0 100,100" repeats L implicitly for the second pair -- same
    # shape as writing out "L100,0 L100,100" explicitly.
    implicit_svg = (
        '<svg width="200" height="200" viewBox="0 0 200 200">'
        '<path d="M0,0 L100,0 100,100 L0,100 Z" fill="FF0000"/>'
        "</svg>"
    )
    explicit_svg = (
        '<svg width="200" height="200" viewBox="0 0 200 200">'
        '<path d="M0,0 L100,0 L100,100 L0,100 Z" fill="FF0000"/>'
        "</svg>"
    )
    tools["write_pptx"](path="implicit.pptx", content=f"layout: svg\n{implicit_svg}\n")
    tools["write_pptx"](path="explicit.pptx", content=f"layout: svg\n{explicit_svg}\n")

    implicit_shape = next(iter(Presentation(str(tmp_path / "implicit.pptx")).slides[0].shapes))
    explicit_shape = next(iter(Presentation(str(tmp_path / "explicit.pptx")).slides[0].shapes))
    assert (implicit_shape.width, implicit_shape.height) == (
        explicit_shape.width,
        explicit_shape.height,
    )


def test_layout_svg_path_rejects_unsupported_command_naming_run_node_script(
    tmp_path: Path,
) -> None:
    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="100" height="100" viewBox="0 0 100 100">'
        '<path d="M10,10 A50,50 0 0 1 90,90" fill="FF0000"/>'
        "</svg>"
    )
    with pytest.raises(ValueError, match="run_node_script"):
        tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")


def test_layout_svg_path_requires_starting_moveto(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="100" height="100" viewBox="0 0 100 100">'
        '<path d="L90,90" fill="FF0000"/>'
        "</svg>"
    )
    with pytest.raises(ValueError, match="moveto"):
        tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")


def test_layout_svg_path_rejects_bare_moveto_with_no_drawing_command(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="100" height="100" viewBox="0 0 100 100">'
        '<path d="M10,10" fill="FF0000"/>'
        "</svg>"
    )
    with pytest.raises(ValueError, match="at least one drawing command"):
        tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")


def test_layout_svg_path_rejects_missing_d_attribute(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    svg_markup = '<svg width="100" height="100" viewBox="0 0 100 100"><path fill="FF0000"/></svg>'
    with pytest.raises(ValueError, match="needs a d attribute"):
        tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")


def test_layout_svg_linear_gradient_fill_applies_to_a_shape(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.dml import MSO_FILL_TYPE

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="200" height="200" viewBox="0 0 200 200">'
        '<defs><linearGradient id="g1" x1="0%" y1="0%" x2="100%" y2="0%">'
        '<stop offset="0%" stop-color="#FF0000"/>'
        '<stop offset="100%" stop-color="#0000FF"/>'
        "</linearGradient></defs>"
        '<rect x="10" y="10" width="100" height="100" fill="url(#g1)"/>'
        "</svg>"
    )
    tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    shape = next(iter(prs.slides[0].shapes))
    assert shape.fill.type == MSO_FILL_TYPE.GRADIENT
    stops = shape.fill.gradient_stops
    assert len(stops) == 2
    # x1=0,y1=0 -> x2=100%,y2=0 is a pure left-to-right vector, DrawingML's
    # own 0-degree default.
    assert shape.fill.gradient_angle == pytest.approx(0.0)


def test_layout_svg_radial_gradient_fill_applies_to_a_shape(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.dml import MSO_FILL_TYPE

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="200" height="200" viewBox="0 0 200 200">'
        '<radialGradient id="g1">'
        '<stop offset="0%" stop-color="white"/>'
        '<stop offset="100%" stop-color="navy"/>'
        "</radialGradient>"
        '<circle cx="100" cy="100" r="80" fill="url(#g1)"/>'
        "</svg>"
    )
    tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    shape = next(iter(prs.slides[0].shapes))
    assert shape.fill.type == MSO_FILL_TYPE.GRADIENT
    with pytest.raises(ValueError, match="not a linear gradient"):
        _ = shape.fill.gradient_angle


def test_layout_svg_gradient_background_rect_is_promoted_to_slide_background(
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="200" height="200" viewBox="0 0 200 200">'
        '<linearGradient id="bg"><stop offset="0%" stop-color="black"/>'
        '<stop offset="100%" stop-color="white"/></linearGradient>'
        '<rect x="0" y="0" width="200" height="200" fill="url(#bg)"/>'
        '<circle cx="100" cy="100" r="20" fill="red"/>'
        "</svg>"
    )
    tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    slide = prs.slides[0]
    # Full-canvas rect promoted to slide background, not left as a shape.
    assert len(slide.shapes) == 1
    from pptx.enum.dml import MSO_FILL_TYPE

    assert slide.background.fill.type == MSO_FILL_TYPE.GRADIENT


def test_layout_svg_gradient_rejects_unknown_id(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="100" height="100" viewBox="0 0 100 100">'
        '<rect x="0" y="0" width="50" height="50" fill="url(#missing)"/>'
        "</svg>"
    )
    with pytest.raises(ValueError, match="no <linearGradient>/<radialGradient>"):
        tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")


def test_layout_svg_gradient_requires_at_least_two_stops(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="100" height="100" viewBox="0 0 100 100">'
        '<linearGradient id="g1"><stop offset="0%" stop-color="red"/></linearGradient>'
        '<rect x="0" y="0" width="50" height="50" fill="url(#g1)"/>'
        "</svg>"
    )
    with pytest.raises(ValueError, match="at least two"):
        tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")


@pytest.mark.real_libreoffice
@pytest.mark.skipif(not _libreoffice_actually_works(), reason="LibreOffice not usable here")
def test_layout_svg_path_and_gradient_survive_libreoffice_conversion(tmp_path: Path) -> None:
    """Real, non-mocked round trip -- catches OOXML corruption from the
    hand-appended <a:cubicBezTo>/<a:quadBezTo>/<a:gradFill> elements that a
    pure structural check could miss."""
    import subprocess

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="800" height="450" viewBox="0 0 800 450">'
        '<defs><linearGradient id="bg" x1="0%" y1="0%" x2="100%" y2="100%">'
        '<stop offset="0%" stop-color="#1e3a8a"/>'
        '<stop offset="100%" stop-color="#7c3aed"/></linearGradient>'
        '<radialGradient id="blob"><stop offset="0%" stop-color="#ffffff"/>'
        '<stop offset="100%" stop-color="#f472b6"/></radialGradient></defs>'
        '<rect x="0" y="0" width="800" height="450" fill="url(#bg)"/>'
        '<path d="M100,300 C150,150 350,150 400,300 Q450,400 500,300 L650,300 '
        'L650,380 L100,380 Z" fill="url(#blob)"/>'
        "</svg>"
    )
    tools["write_pptx"](path="deck.pptx", content=f"layout: svg\n{svg_markup}\n")

    result = subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(tmp_path / "deck.pptx"),
        ],
        capture_output=True,
        timeout=30,
    )
    assert result.returncode == 0


def test_layout_svg_rejects_malformed_xml(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "layout: svg\n<svg width=\"100\" height=\"100\"><rect></svg>\n"

    with pytest.raises(ValueError, match="not valid XML"):
        tools["write_pptx"](path="deck.pptx", content=content)


def test_layout_mixed_deck_in_one_write_pptx_call(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    svg_markup = (
        '<svg width="960" height="540" viewBox="0 0 960 540">'
        '<rect x="0" y="0" width="960" height="540" fill="1E2761"/>'
        '<circle cx="120" cy="270" r="40" fill="FF0000"/>'
        '<text x="220" y="260" font-size="36" fill="FFFFFF">Tracking</text>'
        "</svg>"
    )
    content = (
        "# Plain Title\n- a bullet\n"
        "---\n"
        "layout: icon-list\n# Highlights\n- one\n- two\n"
        "---\n"
        "layout: stat-callout\n# Stats\n| Stat | Label |\n| --- | --- |\n| 5 | Five |\n"
        "---\n"
        "layout: two-column\n# Compare\n- left\n>>>\n- right\n"
        "---\n"
        f"layout: svg\n{svg_markup}\n"
    )
    result = tools["write_pptx"](path="deck.pptx", content=content)

    assert result["slide_count"] == 5
    prs = Presentation(str(tmp_path / "deck.pptx"))
    assert _check_text_overlaps(prs) == []
    assert _check_missing_visual_elements(prs) == [1, 4]


CHART_TABLE = """\
| Quarter | Revenue | Profit |
| --- | --- | --- |
| Q1 | 100 | 20 |
| Q2 | 120 | 25 |
"""


def test_add_pptx_chart_adds_a_chart_of_the_right_type(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.chart import XL_CHART_TYPE

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["add_pptx_chart"](
        path="deck.pptx", slide=1, chart_type="bar", data=CHART_TABLE, title="Q1 vs Q2"
    )

    prs = Presentation(tmp_path / "deck.pptx")
    charts = [shape for shape in prs.slides[0].shapes if shape.has_chart]
    assert len(charts) == 1
    assert charts[0].chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED


def test_add_pptx_chart_rejects_unknown_chart_type(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="Unknown chart_type"):
        tools["add_pptx_chart"](path="deck.pptx", slide=1, chart_type="scatter", data=CHART_TABLE)


def test_add_pptx_chart_rejects_non_table_data(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="pipe-table"):
        tools["add_pptx_chart"](path="deck.pptx", slide=1, chart_type="bar", data="not a table")


def test_add_pptx_chart_rejects_non_numeric_cell(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    bad_table = "| Quarter | Revenue |\n| --- | --- |\n| Q1 | not-a-number |"
    with pytest.raises(ValueError, match="not numeric"):
        tools["add_pptx_chart"](path="deck.pptx", slide=1, chart_type="bar", data=bad_table)


def test_add_pptx_chart_rejects_multi_series_pie_chart(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="pie"):
        tools["add_pptx_chart"](path="deck.pptx", slide=1, chart_type="pie", data=CHART_TABLE)


def test_add_pptx_chart_out_of_range_slide_names_real_count(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="1 slides"):
        tools["add_pptx_chart"](path="deck.pptx", slide=5, chart_type="bar", data=CHART_TABLE)


def test_add_pptx_chart_on_missing_file_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)

    with pytest.raises(ValueError, match="does not exist"):
        tools["add_pptx_chart"](path="missing.pptx", slide=1, chart_type="bar", data=CHART_TABLE)


def test_add_pptx_chart_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["add_pptx_chart"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


def test_add_pptx_chart_positions_below_existing_body_text_without_overlap(
    tmp_path: Path,
) -> None:
    """Real, live-reproduced bug: the old fixed Inches(1, 1.6, 8, 5) chart
    box landed directly on top of a slide's own short intro sentence.
    A chart added next to real body text must not overlap it, and must
    stay within the slide's own real bounds -- not just some fraction of
    an unrelated content-area constant (the second bug the first fix's
    naive version introduced, see the git history for this test)."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Revenue Trend\n- One short intro line")

    tools["add_pptx_chart"](path="deck.pptx", slide=1, chart_type="line", data=CHART_TABLE)

    prs = Presentation(tmp_path / "deck.pptx")
    slide = prs.slides[0]
    body = next(
        s
        for s in slide.shapes
        if s.has_text_frame and "intro line" in s.text_frame.text and not s.has_chart
    )
    chart_shape = next(s for s in slide.shapes if s.has_chart)
    assert chart_shape.top >= body.top + body.height
    assert chart_shape.top + chart_shape.height <= (prs.slide_height or 0)
    assert _check_text_overlaps(prs) == []
    # Real, live-reproduced defect this last assertion catches: shrinking
    # the body placeholder by setting only `.height` (before this was
    # fixed to set left/top/width too) made python-pptx create a bare new
    # <a:xfrm> with width defaulting to 0 -- the placeholder rendered as
    # a near-zero-width sliver, one character per line, spilling off the
    # slide. A 0-width bbox still "has" a bbox, so no automated overlap/
    # missing-visual check catches this on its own; only a direct width
    # assertion does.
    assert body.width > 0


def test_add_pptx_chart_rejects_when_body_text_leaves_no_room(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    many_bullets = "\n".join(f"- point number {i}" for i in range(20))
    tools["write_pptx"](path="deck.pptx", content=f"# Busy Slide\n{many_bullets}")

    with pytest.raises(ValueError, match="too little room"):
        tools["add_pptx_chart"](path="deck.pptx", slide=1, chart_type="line", data=CHART_TABLE)


def test_add_pptx_chart_font_color_matches_default_black_theme(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["add_pptx_chart"](path="deck.pptx", slide=1, chart_type="bar", data=CHART_TABLE)

    prs = Presentation(tmp_path / "deck.pptx")
    chart = next(s for s in prs.slides[0].shapes if s.has_chart).chart
    assert chart.font.color.rgb == RGBColor(0x00, 0x00, 0x00)  # type: ignore[no-untyped-call]


def test_add_pptx_chart_font_color_matches_custom_dark_theme(tmp_path: Path) -> None:
    """Real, live-reproduced bug: a chart's axis/legend text stayed black
    (python-pptx's own chart default) even on a deck whose `theme=`
    parameter lightened every other piece of text on the slide, making it
    unreadable against a dark background."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](
        path="deck.pptx", content="# Title\n- body", theme="bg=111827,text=F8FAFC"
    )

    tools["add_pptx_chart"](path="deck.pptx", slide=1, chart_type="bar", data=CHART_TABLE)

    prs = Presentation(tmp_path / "deck.pptx")
    chart = next(s for s in prs.slides[0].shapes if s.has_chart).chart
    assert chart.font.color.rgb == RGBColor.from_string("F8FAFC")  # type: ignore[no-untyped-call]


def test_write_pptx_with_template_path_reuses_theme_and_drops_old_slides(
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="template.pptx", content="# Old\n- old bullet\n\n---\n\n# Old 2\n- x")

    tools["write_pptx"](
        path="new.pptx", content="# New Title\n- new bullet", template_path="template.pptx"
    )

    text = tools["read_pptx"](path="new.pptx")
    assert "New Title" in text
    assert "new bullet" in text
    assert "Old" not in text

    prs = Presentation(tmp_path / "new.pptx")
    assert len(prs.slides) == 1

    # Re-adding a slide after the template's own slides were cleared must
    # not collide with an orphaned slide part in the zip package.
    import warnings
    import zipfile

    with warnings.catch_warnings():
        warnings.simplefilter("error")
        with zipfile.ZipFile(tmp_path / "new.pptx") as archive:
            slide_parts = [n for n in archive.namelist() if n.startswith("ppt/slides/slide")]
    assert len(slide_parts) == 1


def test_write_pptx_with_missing_template_path_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)

    with pytest.raises(ValueError, match="does not exist"):
        tools["write_pptx"](path="new.pptx", content="# Hi\n- x", template_path="missing.pptx")


def test_write_pptx_with_template_missing_required_layouts_raises(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="template.pptx", content="# Old\n- old bullet")

    prs = Presentation(tmp_path / "template.pptx")
    layout_master = prs.slide_masters[0]
    # Strip every layout but the first, simulating a minimal custom template.
    sld_layout_id_lst = layout_master.element.find(qn("p:sldLayoutIdLst"))
    for layout_id in list(sld_layout_id_lst)[1:]:
        r_id = layout_id.get(qn("r:id"))
        layout_master.part.drop_rel(r_id)
        sld_layout_id_lst.remove(layout_id)
    prs.save(tmp_path / "template.pptx")

    with pytest.raises(ValueError, match="slide layout"):
        tools["write_pptx"](
            path="new.pptx", content="# Hi\n- x", template_path="template.pptx"
        )


def _make_test_png(path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (10, 10), color="red").save(path)


def test_add_pptx_image_adds_a_picture_shape(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    _make_test_png(tmp_path / "logo.png")

    tools["add_pptx_image"](path="deck.pptx", slide=1, image_path="logo.png")

    prs = Presentation(tmp_path / "deck.pptx")
    pictures = [shape for shape in prs.slides[0].shapes if shape.shape_type == 13]
    assert len(pictures) == 1


def test_add_pptx_image_out_of_range_slide_names_real_count(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    _make_test_png(tmp_path / "logo.png")

    with pytest.raises(ValueError, match="1 slides"):
        tools["add_pptx_image"](path="deck.pptx", slide=5, image_path="logo.png")


def test_add_pptx_image_on_missing_file_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="does not exist"):
        tools["add_pptx_image"](path="deck.pptx", slide=1, image_path="missing.png")


def test_add_pptx_image_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["add_pptx_image"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


def _make_wide_test_png(path: Path) -> None:
    """A non-square image, wider than the default 4:3 slide's own ratio --
    needed to prove set_pptx_background_image's cover-crop math actually
    ran (crop fractions land on a non-default, non-trivial value), unlike
    _make_test_png's square image which a naive stretch-to-fill would also
    "cover" without any crop at all."""
    from PIL import Image

    Image.new("RGB", (400, 100), color="blue").save(path)


def test_set_pptx_background_image_places_picture_behind_existing_shapes(
    tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    _make_wide_test_png(tmp_path / "bg.png")

    tools["set_pptx_background_image"](path="deck.pptx", slide=1, image_path="bg.png")

    prs = Presentation(tmp_path / "deck.pptx")
    sp_tree = prs.slides[0].shapes._spTree
    # First two children are always nvGrpSpPr/grpSpPr; the picture must be
    # the very next one (first *shape* child) so it renders behind the
    # title/body placeholders that were already on the slide.
    children = list(sp_tree)
    assert children[2].tag == qn("p:pic")
    assert all(child.tag != qn("p:pic") for child in children[3:])


def test_set_pptx_background_image_crops_to_cover_without_distortion(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    _make_wide_test_png(tmp_path / "bg.png")  # 400x100, much wider than a 4:3 slide

    tools["set_pptx_background_image"](path="deck.pptx", slide=1, image_path="bg.png")

    prs = Presentation(tmp_path / "deck.pptx")
    picture = next(
        shape for shape in prs.slides[0].shapes if shape.shape_type == 13
    )
    # Wider-than-slide image -> crop left/right, not top/bottom -- and the
    # crop must be a real, non-trivial fraction (not 0.0, which would mean
    # the cover-crop math never ran and the image was just stretched).
    assert picture.crop_left > 0.0
    assert picture.crop_right > 0.0
    assert picture.crop_top == 0.0
    assert picture.crop_bottom == 0.0
    # The shape itself still spans the whole slide -- the crop lives in
    # srcRect (which part of the source image is sampled), not the shape's
    # own on-slide extent.
    assert picture.width == prs.slide_width
    assert picture.height == prs.slide_height


def test_set_pptx_background_image_out_of_range_slide_names_real_count(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    _make_wide_test_png(tmp_path / "bg.png")

    with pytest.raises(ValueError, match="1 slides"):
        tools["set_pptx_background_image"](path="deck.pptx", slide=5, image_path="bg.png")


def test_set_pptx_background_image_on_missing_file_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="does not exist"):
        tools["set_pptx_background_image"](path="deck.pptx", slide=1, image_path="missing.png")


def test_set_pptx_background_image_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["set_pptx_background_image"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


@pytest.mark.real_libreoffice
@pytest.mark.skipif(not _libreoffice_actually_works(), reason="LibreOffice not usable here")
def test_set_pptx_background_image_survives_libreoffice_conversion(tmp_path: Path) -> None:
    """Real, non-mocked round trip -- catches OOXML corruption from the
    spTree element move that a pure structural check could miss."""
    import subprocess

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    _make_wide_test_png(tmp_path / "bg.png")
    tools["set_pptx_background_image"](path="deck.pptx", slide=1, image_path="bg.png")

    result = subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(tmp_path / "deck.pptx"),
        ],
        capture_output=True,
        timeout=30,
    )
    assert result.returncode == 0
    assert (tmp_path / "deck.pdf").is_file()
    assert (tmp_path / "deck.pdf").stat().st_size > 0


def test_add_pptx_scrim_sets_alpha_on_the_fill(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["add_pptx_scrim"](path="deck.pptx", slide=1, opacity=0.4, color="000000")

    prs = Presentation(tmp_path / "deck.pptx")
    rectangles = [shape for shape in prs.slides[0].shapes if shape.shape_type == 1]  # AUTO_SHAPE
    assert len(rectangles) == 1
    srgb_clr = rectangles[0].fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    alpha = srgb_clr.find(qn("a:alpha"))
    assert alpha is not None
    assert alpha.get("val") == "40000"


def test_add_pptx_scrim_defaults_to_opacity_035(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["add_pptx_scrim"](path="deck.pptx", slide=1)

    prs = Presentation(tmp_path / "deck.pptx")
    rectangle = next(shape for shape in prs.slides[0].shapes if shape.shape_type == 1)
    srgb_clr = rectangle.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    assert srgb_clr.find(qn("a:alpha")).get("val") == "35000"
    assert srgb_clr.get("val") == "000000"


def test_add_pptx_scrim_layers_above_background_picture_below_other_shapes(
    tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    _make_wide_test_png(tmp_path / "bg.png")
    tools["set_pptx_background_image"](path="deck.pptx", slide=1, image_path="bg.png")

    tools["add_pptx_scrim"](path="deck.pptx", slide=1)

    prs = Presentation(tmp_path / "deck.pptx")
    sp_tree = prs.slides[0].shapes._spTree
    children = list(sp_tree)
    # Both the scrim and the title/body placeholders are <p:sp> elements --
    # a bare tag match can't tell them apart, so identify the scrim by the
    # one thing unique to it: an <a:alpha> element in its fill.
    alpha_path = ".//" + qn("a:alpha")
    pic_index = [c.tag for c in children].index(qn("p:pic"))
    scrim_index = next(i for i, c in enumerate(children) if c.find(alpha_path) is not None)
    assert scrim_index == pic_index + 1
    # Everything after the scrim is a plain <p:sp> without an alpha element
    # -- the original title/body placeholders, layered above the scrim.
    for element in children[scrim_index + 1 :]:
        assert element.find(alpha_path) is None


def test_add_pptx_scrim_without_background_picture_is_backmost(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.oxml.ns import qn

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["add_pptx_scrim"](path="deck.pptx", slide=1)

    prs = Presentation(tmp_path / "deck.pptx")
    sp_tree = prs.slides[0].shapes._spTree
    children = list(sp_tree)
    group_props_index = [c.tag for c in children].index(qn("p:grpSpPr"))
    assert children[group_props_index + 1].tag == qn("p:sp")


def test_add_pptx_scrim_out_of_range_slide_names_real_count(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="1 slides"):
        tools["add_pptx_scrim"](path="deck.pptx", slide=5)


def test_add_pptx_scrim_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["add_pptx_scrim"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


@pytest.mark.real_libreoffice
@pytest.mark.skipif(not _libreoffice_actually_works(), reason="LibreOffice not usable here")
def test_add_pptx_scrim_survives_libreoffice_conversion(tmp_path: Path) -> None:
    """Real, non-mocked round trip -- catches OOXML corruption from the
    hand-appended <a:alpha> element and the spTree move that a pure
    structural check could miss."""
    import subprocess

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    _make_wide_test_png(tmp_path / "bg.png")
    tools["set_pptx_background_image"](path="deck.pptx", slide=1, image_path="bg.png")
    tools["add_pptx_scrim"](path="deck.pptx", slide=1, opacity=0.4)

    result = subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(tmp_path / "deck.pptx"),
        ],
        capture_output=True,
        timeout=30,
    )
    assert result.returncode == 0
    assert (tmp_path / "deck.pdf").is_file()
    assert (tmp_path / "deck.pdf").stat().st_size > 0


def test_set_pptx_notes_round_trips_via_python_pptx(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["set_pptx_notes"](path="deck.pptx", slide=1, notes="Remember to smile.")

    prs = Presentation(tmp_path / "deck.pptx")
    assert prs.slides[0].notes_slide.notes_text_frame.text == "Remember to smile."


def test_set_pptx_notes_replaces_rather_than_appends(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["set_pptx_notes"](path="deck.pptx", slide=1, notes="First draft")
    tools["set_pptx_notes"](path="deck.pptx", slide=1, notes="Final version")

    prs = Presentation(tmp_path / "deck.pptx")
    assert prs.slides[0].notes_slide.notes_text_frame.text == "Final version"


def test_set_pptx_notes_out_of_range_slide_names_real_count(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="1 slides"):
        tools["set_pptx_notes"](path="deck.pptx", slide=5, notes="x")


def test_set_pptx_notes_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["set_pptx_notes"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


_P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
_MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"
_P15_NS = "http://schemas.microsoft.com/office/powerpoint/2012/main"
_P159_NS = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"


def test_set_pptx_transition_adds_transition_element(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="fade", duration=0.8)

    prs = Presentation(tmp_path / "deck.pptx")
    transition_el = prs.slides[0].element.find(f"{{{_P_NS}}}transition")
    assert transition_el is not None
    assert [c.tag.split("}")[-1] for c in transition_el] == ["fade"]
    # p14:dur is a PowerPoint-2010 extension attribute with no home in the
    # base ECMA-376 schema -- valid OOXML only because mc:Ignorable marks
    # its namespace prefix ignorable (see _mark_mce_ignorable's docstring).
    # A real slide always carries both together, not just the attribute.
    assert (prs.slides[0].element.get(f"{{{_MC_NS}}}Ignorable") or "").split() == ["p14"]


def test_set_pptx_transition_none_removes_any_existing_transition(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="fade")

    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="none")

    prs = Presentation(tmp_path / "deck.pptx")
    assert prs.slides[0].element.find(f"{{{_P_NS}}}transition") is None
    # No p14: content left on the slide -- mc:Ignorable shouldn't still
    # claim there is any.
    assert prs.slides[0].element.get(f"{{{_MC_NS}}}Ignorable") is None


def test_set_pptx_transition_replaces_rather_than_duplicates(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="fade")

    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="wipe")

    prs = Presentation(tmp_path / "deck.pptx")
    transitions = prs.slides[0].element.findall(f"{{{_P_NS}}}transition")
    assert len(transitions) == 1
    assert [c.tag.split("}")[-1] for c in transitions[0]] == ["wipe"]


def test_set_pptx_transition_rejects_unknown_transition(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="Unknown transition"):
        tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="not-a-real-transition")


def test_set_pptx_transition_supports_a_p14_extension_effect(tmp_path: Path) -> None:
    """"vortex" lives in the PowerPoint-2010 extension namespace (p14),
    not the base ECMA-376 "p" namespace fade/push/wipe use -- both the
    effect element itself and the p14:dur attribute need mc:Ignorable."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="vortex", duration=0.8)

    prs = Presentation(tmp_path / "deck.pptx")
    transition_el = prs.slides[0].element.find(f"{{{_P_NS}}}transition")
    assert transition_el is not None
    effect = transition_el[0]
    assert effect.tag == f"{{{_P14_NS}}}vortex"
    assert effect.get("dir") == "r"
    assert (prs.slides[0].element.get(f"{{{_MC_NS}}}Ignorable") or "").split() == ["p14"]


def test_set_pptx_transition_supports_a_p159_extension_effect(tmp_path: Path) -> None:
    """"morph" is the one native transition in the newest (2015/09,
    p159) extension namespace."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="morph")

    prs = Presentation(tmp_path / "deck.pptx")
    transition_el = prs.slides[0].element.find(f"{{{_P_NS}}}transition")
    effect = transition_el[0]
    assert effect.tag == f"{{{_P159_NS}}}morph"
    ignorable = (prs.slides[0].element.get(f"{{{_MC_NS}}}Ignorable") or "").split()
    assert set(ignorable) == {"p14", "p159"}


def test_set_pptx_transition_switching_namespaces_cleans_up_the_old_prefix(
    tmp_path: Path,
) -> None:
    """p14 -> p15 -> base "p": each switch must leave mc:Ignorable naming
    only the prefix(es) the *current* transition actually uses, not a
    stale leftover from whichever effect was set before it."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="vortex")  # p14
    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="fracture")  # p15
    prs = Presentation(tmp_path / "deck.pptx")
    ignorable = (prs.slides[0].element.get(f"{{{_MC_NS}}}Ignorable") or "").split()
    assert set(ignorable) == {"p14", "p15"}

    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="fade")  # base "p"
    prs = Presentation(tmp_path / "deck.pptx")
    ignorable = (prs.slides[0].element.get(f"{{{_MC_NS}}}Ignorable") or "").split()
    assert ignorable == ["p14"]  # only the p14:dur attribute remains foreign


def test_set_pptx_transition_resolves_legacy_aliases(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="wheel")  # alias for "clock"

    prs = Presentation(tmp_path / "deck.pptx")
    transition_el = prs.slides[0].element.find(f"{{{_P_NS}}}transition")
    effect = transition_el[0]
    assert effect.tag == f"{{{_P_NS}}}wheel"
    assert effect.get("spokes") == "1"


def test_set_pptx_transition_is_schema_valid_for_every_registered_transition(
    tmp_path: Path,
) -> None:
    """Every single one of the 48 native transitions + 8 aliases actually
    produces schema-valid OOXML (via the real assert_ooxml_valid call
    set_pptx_transition itself makes) -- not just the few spot-checked
    above. A real, cheap way to catch a typo'd element/attribute name
    across the whole retyped registry at once."""
    from coscribe.tools.presentations import _TRANSITION_ALIASES, _TRANSITION_SPECS

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    for transition in (*_TRANSITION_SPECS, *_TRANSITION_ALIASES):
        tools["set_pptx_transition"](path="deck.pptx", slide=1, transition=transition)


def test_list_pptx_transition_types_matches_the_real_registry(tmp_path: Path) -> None:
    from coscribe.tools.presentations import _TRANSITION_ALIASES, _TRANSITION_SPECS

    tools = _tools_by_name(tmp_path)

    result = tools["list_pptx_transition_types"]()

    assert result == sorted({*_TRANSITION_SPECS, *_TRANSITION_ALIASES})
    assert "none" not in result
    assert len(result) == 48 + 8


def test_set_pptx_transition_out_of_range_slide_names_real_count(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="1 slides"):
        tools["set_pptx_transition"](path="deck.pptx", slide=5, transition="fade")


def test_set_pptx_transition_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["set_pptx_transition"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed or not functional in this environment",
)
def test_set_pptx_transition_survives_libreoffice_conversion(tmp_path: Path) -> None:
    import subprocess

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["set_pptx_transition"](path="deck.pptx", slide=1, transition="fade", duration=0.8)

    result = subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(tmp_path / "deck.pptx"),
        ],
        capture_output=True,
        timeout=30,
        check=True,
    )
    assert result.returncode == 0
    assert (tmp_path / "deck.pdf").is_file()


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed or not functional in this environment",
)
def test_set_pptx_transition_exotic_effects_survive_libreoffice_conversion(
    tmp_path: Path,
) -> None:
    """A mix across all 4 namespace tiers (base "p", p14, p15, p159) plus
    a legacy alias, on one real 5-slide deck -- not just the single base
    "fade" case above. Can't verify the transition *animation* itself in
    a static test (that needs a human watching real PowerPoint/LibreOffice
    play it back), but a real subprocess conversion succeeding, plus a
    warnings-as-errors python-pptx round-trip, is the strongest automated
    check available: it proves the file isn't corrupted the way a
    hand-XML mistake across a 48-entry retyped registry plausibly could
    produce for at least one entry."""
    import subprocess
    import warnings

    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](
        path="deck.pptx",
        content="# S1\n- a\n---\n# S2\n- b\n---\n# S3\n- c\n---\n# S4\n- d\n---\n# S5\n- e",
    )
    for slide, transition in enumerate(["vortex", "fracture", "morph", "fade", "wheel"], start=1):
        tools["set_pptx_transition"](
            path="deck.pptx", slide=slide, transition=transition, duration=0.7
        )

    with warnings.catch_warnings():
        warnings.simplefilter("error")
        prs = Presentation(tmp_path / "deck.pptx")
        assert len(prs.slides) == 5

    result = subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(tmp_path / "deck.pptx"),
        ],
        capture_output=True,
        timeout=30,
        check=True,
    )
    assert result.returncode == 0
    assert (tmp_path / "deck.pdf").is_file()


def test_add_pptx_animation_fade_is_structurally_present(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="fade")

    prs = Presentation(tmp_path / "deck.pptx")
    title_id = prs.slides[0].shapes.title.shape_id
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    assert timing is not None
    matches = [
        cTn
        for cTn in timing.iter(f"{{{_P_NS}}}cTn")
        if cTn.get("presetID") == "10" and cTn.get("presetSubtype") == "0"
    ]
    assert len(matches) == 1
    sp_tgt = matches[0].find(f".//{{{_P_NS}}}spTgt")
    assert sp_tgt.get("spid") == str(title_id)
    assert matches[0].find(f".//{{{_P_NS}}}animEffect") is not None


def test_add_pptx_animation_fly_in_is_structurally_present(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=1, animation="fly-in")

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    matches = [
        cTn
        for cTn in timing.iter(f"{{{_P_NS}}}cTn")
        if cTn.get("presetID") == "2" and cTn.get("presetSubtype") == "4"
    ]
    assert len(matches) == 1
    anim_elements = matches[0].findall(f".//{{{_P_NS}}}anim")
    assert len(anim_elements) == 2


def test_add_pptx_animation_default_trigger_uses_click_effect_node_type(tmp_path: Path) -> None:
    """Regression test for a real, evidence-backed bug found while
    researching trigger-mode support: nodeType on the presetID-bearing
    <p:cTn> encodes which Animation Pane Start mode a row uses
    (on-click -> "clickEffect", with-previous -> "withEffect",
    after-previous -> "afterEffect" -- cross-checked against
    hugohe3/ppt-master's pptx_animations.py, which always overwrites this
    attribute at row-instantiation time based on the caller's actual
    trigger). This file's own previous hardcoded "afterEffect" (itself a
    prior "fix" from "clickEffect", see git history) was wrong for the
    on-click-only trigger this tool exposed at the time -- the default
    `trigger="on-click"` must produce "clickEffect", not "afterEffect"."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="fade")

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    matches = [
        cTn
        for cTn in timing.iter(f"{{{_P_NS}}}cTn")
        if cTn.get("presetID") == "10" and cTn.get("presetSubtype") == "0"
    ]
    assert matches[0].get("nodeType") == "clickEffect"


@pytest.mark.parametrize(
    ("animation", "preset_id", "preset_subtype", "preset_class", "child_tag"),
    [
        ("exit-fade", "10", "0", "exit", "animEffect"),
        ("exit-fly", "2", "4", "exit", "anim"),
        ("emphasis-grow", "6", "0", "emph", "animScale"),
        ("emphasis-spin", "8", "0", "emph", "animRot"),
    ],
)
def test_add_pptx_animation_new_effect_types_are_structurally_present(
    tmp_path: Path,
    animation: str,
    preset_id: str,
    preset_subtype: str,
    preset_class: str,
    child_tag: str,
) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation=animation)

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    matches = [
        cTn
        for cTn in timing.iter(f"{{{_P_NS}}}cTn")
        if cTn.get("presetID") == preset_id and cTn.get("presetSubtype") == preset_subtype
    ]
    assert len(matches) == 1
    assert matches[0].get("presetClass") == preset_class
    assert matches[0].find(f".//{{{_P_NS}}}{child_tag}") is not None


def test_add_pptx_animation_exit_fade_and_entrance_fade_are_distinguishable(
    tmp_path: Path,
) -> None:
    """Regression test for the real disambiguation bug fixed alongside
    the new effect types: exit-fade and fade share the exact same
    presetID/presetSubtype in real PowerPoint XML (both 10/0) -- only
    presetClass ("exit" vs "entr") tells them apart. Adding exit-fade to
    one shape and fade to another on the same slide must produce two
    genuinely distinct rows, not one row two overlapping readback checks
    both accept."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="fade")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=1, animation="exit-fade")

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    same_ids = [
        cTn
        for cTn in timing.iter(f"{{{_P_NS}}}cTn")
        if cTn.get("presetID") == "10" and cTn.get("presetSubtype") == "0"
    ]
    assert len(same_ids) == 2
    assert {cTn.get("presetClass") for cTn in same_ids} == {"entr", "exit"}


def test_add_pptx_animation_two_calls_produce_unique_ids(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="fade")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=1, animation="fly-in")

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    ids = [cTn.get("id") for cTn in timing.iter(f"{{{_P_NS}}}cTn")]
    assert len(ids) == len(set(ids))

    # Both click groups must sit directly under mainSeq's childTnLst.
    main_ctn = timing.find(f".//{{{_P_NS}}}cTn[@nodeType='mainSeq']")
    click_groups = main_ctn.find(f"{{{_P_NS}}}childTnLst").findall(f"{{{_P_NS}}}par")
    assert len(click_groups) == 2


def test_add_pptx_animation_with_previous_chains_into_the_same_group(tmp_path: Path) -> None:
    """with-previous/after-previous append into the LAST existing group
    on the slide as one more step, rather than starting a new top-level
    click group -- the whole point of auto-play is that the presenter
    doesn't click again for these."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="fade")
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=1, animation="fly-in", trigger="with-previous"
    )

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    main_ctn = timing.find(f".//{{{_P_NS}}}cTn[@nodeType='mainSeq']")
    groups = main_ctn.find(f"{{{_P_NS}}}childTnLst").findall(f"{{{_P_NS}}}par")
    assert len(groups) == 1
    group_ctn = groups[0].find(f"{{{_P_NS}}}cTn")
    steps = group_ctn.find(f"{{{_P_NS}}}childTnLst").findall(f"{{{_P_NS}}}par")
    assert len(steps) == 2


def test_add_pptx_animation_with_previous_starts_at_same_time_as_previous_step(
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="fade")
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=1, animation="fly-in", trigger="with-previous"
    )

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    group_path = (
        f".//{{{_P_NS}}}cTn[@nodeType='mainSeq']"
        f"/{{{_P_NS}}}childTnLst/{{{_P_NS}}}par/{{{_P_NS}}}cTn"
    )
    group_ctn = timing.find(group_path)
    steps = group_ctn.find(f"{{{_P_NS}}}childTnLst").findall(f"{{{_P_NS}}}par")
    delays = [
        step.find(f"{{{_P_NS}}}cTn/{{{_P_NS}}}stCondLst/{{{_P_NS}}}cond").get("delay")
        for step in steps
    ]
    assert delays == ["0", "0"]


def test_add_pptx_animation_after_previous_starts_once_previous_duration_elapses(
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=0, animation="fade", duration=0.4
    )
    tools["add_pptx_animation"](
        path="deck.pptx",
        slide=1,
        shape_index=1,
        animation="emphasis-grow",
        duration=0.3,
        trigger="after-previous",
        delay=0.1,
    )

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    group_path = (
        f".//{{{_P_NS}}}cTn[@nodeType='mainSeq']"
        f"/{{{_P_NS}}}childTnLst/{{{_P_NS}}}par/{{{_P_NS}}}cTn"
    )
    group_ctn = timing.find(group_path)
    steps = group_ctn.find(f"{{{_P_NS}}}childTnLst").findall(f"{{{_P_NS}}}par")
    second_cond = steps[1].find(f"{{{_P_NS}}}cTn/{{{_P_NS}}}stCondLst/{{{_P_NS}}}cond")
    # previous (fade) starts at 0, lasts 400ms; +100ms extra delay = 500ms.
    assert second_cond.get("delay") == "500"


def test_add_pptx_animation_after_previous_with_no_existing_animation_auto_anchors(
    tmp_path: Path,
) -> None:
    """The very first animation on a slide still plays even when its own
    Start mode is "After Previous" -- matching real PowerPoint, where
    there's nothing to wait for, so it starts on slide entry instead of
    needing a click."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=0, animation="fade", trigger="after-previous"
    )

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    main_ctn = timing.find(f".//{{{_P_NS}}}cTn[@nodeType='mainSeq']")
    groups = main_ctn.find(f"{{{_P_NS}}}childTnLst").findall(f"{{{_P_NS}}}par")
    assert len(groups) == 1
    conds = groups[0].find(f"{{{_P_NS}}}cTn/{{{_P_NS}}}stCondLst").findall(f"{{{_P_NS}}}cond")
    assert [cond.get("delay") for cond in conds] == ["indefinite", "0"]
    assert conds[1].get("evt") == "onBegin"

    row = groups[0].find(f".//{{{_P_NS}}}cTn[@presetID]")
    assert row.get("nodeType") == "afterEffect"


def test_add_pptx_animation_rejects_unknown_trigger(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="Unknown trigger"):
        tools["add_pptx_animation"](
            path="deck.pptx", slide=1, shape_index=0, animation="fade", trigger="bogus"
        )


def test_add_pptx_animation_by_paragraph_targets_each_non_empty_paragraph(
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- one\n- two\n- three")

    result = tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=1, animation="fade", by_paragraph=True
    )
    assert result["paragraph_count"] == 3

    prs = Presentation(tmp_path / "deck.pptx")
    body_id = list(prs.slides[0].shapes)[1].shape_id
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    p_indexes = set()
    for cTn in timing.iter(f"{{{_P_NS}}}cTn"):
        if cTn.get("presetID") != "10" or cTn.get("presetClass") != "entr":
            continue
        sp_tgt = cTn.find(f".//{{{_P_NS}}}spTgt")
        if sp_tgt is None or sp_tgt.get("spid") != str(body_id):
            continue
        p_rg = sp_tgt.find(f"{{{_P_NS}}}txEl/{{{_P_NS}}}pRg")
        if p_rg is not None:
            p_indexes.add(int(p_rg.get("st")))
    assert p_indexes == {0, 1, 2}


def test_add_pptx_animation_by_paragraph_adds_one_bldp_entry(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- one\n- two")
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=1, animation="fade", by_paragraph=True
    )

    prs = Presentation(tmp_path / "deck.pptx")
    body_id = list(prs.slides[0].shapes)[1].shape_id
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    bld_p = timing.find(f"{{{_P_NS}}}bldLst").findall(f"{{{_P_NS}}}bldP")
    assert len(bld_p) == 1
    assert bld_p[0].get("spid") == str(body_id)
    assert bld_p[0].get("build") == "p"


def test_add_pptx_animation_by_paragraph_rejects_shape_without_text_frame(
    tmp_path: Path,
) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n| a | b |\n| - | - |\n| 1 | 2 |")

    with pytest.raises(ValueError, match="no text frame"):
        tools["add_pptx_animation"](
            path="deck.pptx", slide=1, shape_index=1, animation="fade", by_paragraph=True
        )


def test_add_pptx_animation_rejects_unknown_animation(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="Unknown animation"):
        tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="spin")


def test_add_pptx_animation_rejects_out_of_range_shape_index(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="2 shape"):
        tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=99, animation="fade")


def test_add_pptx_animation_out_of_range_slide_names_real_count(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="1 slides"):
        tools["add_pptx_animation"](path="deck.pptx", slide=5, shape_index=0, animation="fade")


def test_add_pptx_animation_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["add_pptx_animation"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed or not functional in this environment",
)
def test_add_pptx_animation_survives_libreoffice_conversion(tmp_path: Path) -> None:
    import subprocess

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="fade")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=1, animation="fly-in")

    result = subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(tmp_path / "deck.pptx"),
        ],
        capture_output=True,
        timeout=30,
        check=True,
    )
    assert result.returncode == 0
    assert (tmp_path / "deck.pdf").is_file()


def test_add_pptx_animation_is_schema_valid_for_every_registered_animation(
    tmp_path: Path,
) -> None:
    """Every single one of the 203 vendored presets + 6 legacy aliases
    actually produces schema-valid, read-back-verified OOXML (via the
    real assert_ooxml_valid/_verify_animation_readback calls
    add_pptx_animation itself makes) -- not just the few spot-checked
    above. A real, cheap way to catch a bad id-renumbering/spid-retarget/
    duration-scaling edge case across the whole catalog at once, the same
    role test_set_pptx_transition_is_schema_valid_for_every_registered_
    transition plays for the 56-entry transition registry."""
    from coscribe.tools.presentations import _ANIMATIONS

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    for animation in _ANIMATIONS:
        tools["add_pptx_animation"](
            path="deck.pptx", slide=1, shape_index=0, animation=animation, trigger="with-previous"
        )


def test_list_pptx_animation_types_matches_the_real_registry(tmp_path: Path) -> None:
    from coscribe.tools.presentations import _ANIMATION_ALIASES, _load_animation_presets

    tools = _tools_by_name(tmp_path)

    result = tools["list_pptx_animation_types"]()

    assert result == sorted({*_load_animation_presets(), *_ANIMATION_ALIASES})
    assert len(result) == 203 + 6


def test_add_pptx_animation_new_catalog_preset_is_structurally_present(
    tmp_path: Path,
) -> None:
    """A preset from the full 203-entry catalog that was never one of the
    original 6 short-named presets -- proves the general parse/renumber/
    retarget path works for a name that never had hand-written XML, not
    just the 6 that used to."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="path_circle")

    prs = Presentation(tmp_path / "deck.pptx")
    title_id = prs.slides[0].shapes.title.shape_id
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    matches = [
        cTn
        for cTn in timing.iter(f"{{{_P_NS}}}cTn")
        if cTn.get("presetID") == "1" and cTn.get("presetClass") == "path"
    ]
    assert len(matches) == 1
    sp_tgt = matches[0].find(f".//{{{_P_NS}}}spTgt")
    assert sp_tgt.get("spid") == str(title_id)
    assert matches[0].find(f".//{{{_P_NS}}}animMotion") is not None


def test_add_pptx_animation_scales_duration_from_the_presets_own_default(
    tmp_path: Path,
) -> None:
    """entrance_fly's two <p:anim> nodes are each authored at dur=500 (its
    real default_duration_ms) -- requesting duration=2.0 (2000ms) must
    scale both by the same 4x ratio, not just plug 2000 into one of them."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    result = tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=0, animation="entrance_fly", duration=2.0
    )
    assert result["duration_ms"] == 2000

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    anim_durations = {
        cTn.get("dur")
        for cTn in timing.iter(f"{{{_P_NS}}}cTn")
        if cTn.getparent().tag == f"{{{_P_NS}}}cBhvr"
        and cTn.getparent().getparent().tag == f"{{{_P_NS}}}anim"
    }
    assert anim_durations == {"2000"}


def test_add_pptx_animation_non_scalable_preset_ignores_requested_duration(
    tmp_path: Path,
) -> None:
    """entrance_appear isn't duration-adjustable in real PowerPoint either
    (it's an instant toggle) -- a requested duration must be silently
    ignored rather than corrupting its authored 1ms timing, and the
    result must report the real duration actually used, not a value that
    implies the request took effect."""
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    result = tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=0, animation="entrance_appear", duration=5.0
    )

    assert result["duration_ms"] == 1


def test_add_pptx_animation_legacy_alias_and_full_name_produce_identical_xml(
    tmp_path: Path,
) -> None:
    """"fade" (the pre-203-catalog short name) and "entrance_fade" (the
    real preset key it was always secretly implementing, per
    _ANIMATION_ALIASES) must resolve to the exact same preset -- not two
    presets that happen to look similar."""
    import re

    from lxml import etree
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_animation"](path="deck.pptx", slide=1, shape_index=0, animation="fade")
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=1, animation="entrance_fade"
    )

    prs = Presentation(tmp_path / "deck.pptx")
    timing = prs.slides[0].element.find(f"{{{_P_NS}}}timing")
    rows = [
        cTn
        for cTn in timing.iter(f"{{{_P_NS}}}cTn")
        if cTn.get("presetID") == "10" and cTn.get("presetClass") == "entr"
    ]
    assert len(rows) == 2
    # id (per-row running counter) and spid (the two calls deliberately
    # target different shapes) are both expected to differ -- strip both
    # before comparing everything else.
    normalized = [
        re.sub(r' (id|spid)="\d+"', "", etree.tostring(row).decode()) for row in rows
    ]
    assert normalized[0] == normalized[1]


def test_list_pptx_animation_types_is_read_and_no_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["list_pptx_animation_types"])
    assert metadata.risk_category == "READ"
    assert metadata.requires_approval is False


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed or not functional in this environment",
)
def test_add_pptx_animation_all_four_categories_survive_libreoffice_conversion(
    tmp_path: Path,
) -> None:
    """One preset from each of the 4 categories (entrance/emphasis/exit/
    path), mixed triggers, plus by_paragraph, on one real deck -- can't
    verify the animation *plays back* correctly in a static test (that
    needs a human watching real PowerPoint/LibreOffice), but a real
    subprocess conversion succeeding, plus a warnings-as-errors
    python-pptx round-trip, is the strongest automated check available."""
    import subprocess
    import warnings

    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- alpha\n- beta\n- gamma")
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=0, animation="entrance_bounce", trigger="on-click"
    )
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=1, animation="emphasis_teeter",
        trigger="with-previous",
    )
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=1, animation="path_bean",
        trigger="after-previous", delay=0.2,
    )
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=0, animation="exit_bounce", trigger="after-previous"
    )
    tools["add_pptx_animation"](
        path="deck.pptx", slide=1, shape_index=1, animation="entrance_fly",
        by_paragraph=True, trigger="on-click",
    )

    with warnings.catch_warnings():
        warnings.simplefilter("error")
        prs = Presentation(tmp_path / "deck.pptx")
        assert len(prs.slides) == 1

    result = subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(tmp_path / "deck.pptx"),
        ],
        capture_output=True,
        timeout=30,
        check=True,
    )
    assert result.returncode == 0
    assert (tmp_path / "deck.pdf").is_file()


def test_add_pptx_hyperlink_whole_shape(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    result = tools["add_pptx_hyperlink"](
        path="deck.pptx", slide=1, shape_index=0, url="https://example.com"
    )
    assert result == {
        "path": "deck.pptx",
        "slide": 1,
        "shape_index": 0,
        "url": "https://example.com",
        "text": None,
    }

    prs = Presentation(tmp_path / "deck.pptx")
    shape = list(prs.slides[0].shapes)[0]
    assert shape.click_action.hyperlink.address == "https://example.com"


def test_add_pptx_hyperlink_on_one_run(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- one\n- two")

    result = tools["add_pptx_hyperlink"](
        path="deck.pptx", slide=1, shape_index=1, url="https://example.com/one", text="one"
    )
    assert result["text"] == "one"

    prs = Presentation(tmp_path / "deck.pptx")
    body = list(prs.slides[0].shapes)[1]
    runs = [run for p in body.text_frame.paragraphs for run in p.runs]
    linked = {run.text: run.hyperlink.address for run in runs}
    assert linked["one"] == "https://example.com/one"
    assert linked["two"] is None


def test_add_pptx_hyperlink_survives_round_trip(tmp_path: Path) -> None:
    import warnings

    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- one")
    tools["add_pptx_hyperlink"](path="deck.pptx", slide=1, shape_index=0, url="https://example.com")

    with warnings.catch_warnings():
        warnings.simplefilter("error")
        Presentation(str(tmp_path / "deck.pptx"))


def test_list_pptx_shapes_reports_whole_shape_hyperlink(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["add_pptx_hyperlink"](path="deck.pptx", slide=1, shape_index=0, url="https://example.com")

    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    assert shapes[0]["hyperlink"] == "https://example.com"
    assert shapes[1]["hyperlink"] is None


def test_add_pptx_hyperlink_rejects_url_without_scheme(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="must start with one of"):
        tools["add_pptx_hyperlink"](path="deck.pptx", slide=1, shape_index=0, url="example.com")


def test_add_pptx_hyperlink_rejects_empty_url(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="must not be empty"):
        tools["add_pptx_hyperlink"](path="deck.pptx", slide=1, shape_index=0, url="   ")


def test_add_pptx_hyperlink_rejects_missing_run_text(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- one\n- two")

    with pytest.raises(ValueError, match=r"no run with exact text 'nonexistent'"):
        tools["add_pptx_hyperlink"](
            path="deck.pptx",
            slide=1,
            shape_index=1,
            url="https://example.com",
            text="nonexistent",
        )


def test_add_pptx_hyperlink_rejects_text_on_shape_without_text_frame(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# T\n| a | b |\n| - | - |\n| 1 | 2 |")

    with pytest.raises(ValueError, match="no text frame"):
        tools["add_pptx_hyperlink"](
            path="deck.pptx", slide=1, shape_index=1, url="https://example.com", text="a"
        )


def test_add_pptx_hyperlink_rejects_out_of_range_shape_index(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="2 shape"):
        tools["add_pptx_hyperlink"](
            path="deck.pptx", slide=1, shape_index=99, url="https://example.com"
        )


def test_add_pptx_hyperlink_out_of_range_slide_names_real_count(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="1 slides"):
        tools["add_pptx_hyperlink"](
            path="deck.pptx", slide=5, shape_index=0, url="https://example.com"
        )


def test_add_pptx_hyperlink_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["add_pptx_hyperlink"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


def test_read_pptx_theme_colors_returns_all_12_stock_slots(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    result = tools["read_pptx_theme_colors"](path="deck.pptx")
    assert result["path"] == "deck.pptx"
    colors = result["colors"]
    assert set(colors) == {
        "dk1",
        "lt1",
        "dk2",
        "lt2",
        "accent1",
        "accent2",
        "accent3",
        "accent4",
        "accent5",
        "accent6",
        "hlink",
        "folHlink",
    }
    # Stock python-pptx theme's own known default accent1.
    assert colors["accent1"] == "4F81BD"
    assert all(value is not None for value in colors.values())


def test_edit_pptx_theme_colors_changes_only_named_slots(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    before = tools["read_pptx_theme_colors"](path="deck.pptx")["colors"]

    result = tools["edit_pptx_theme_colors"](
        path="deck.pptx", colors="accent1=0F6B5C,accent2=C2410C,hlink=0F6B5C"
    )
    assert result["colors"] == {"accent1": "0F6B5C", "accent2": "C2410C", "hlink": "0F6B5C"}

    after = tools["read_pptx_theme_colors"](path="deck.pptx")["colors"]
    assert after["accent1"] == "0F6B5C"
    assert after["accent2"] == "C2410C"
    assert after["hlink"] == "0F6B5C"
    # Every other slot is untouched.
    untouched = ("dk1", "lt1", "dk2", "lt2", "accent3", "accent4", "accent5", "accent6", "folHlink")
    for slot in untouched:
        assert after[slot] == before[slot]


def test_edit_pptx_theme_colors_applies_to_every_slide_master(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["edit_pptx_theme_colors"](path="deck.pptx", colors="accent1=0F6B5C")

    prs = Presentation(tmp_path / "deck.pptx")
    for master in prs.slide_masters:
        theme_part = master.part.part_related_by(RT.THEME)
        assert b"0F6B5C" in theme_part.blob


def test_edit_pptx_theme_colors_survives_round_trip(tmp_path: Path) -> None:
    import warnings

    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")
    tools["edit_pptx_theme_colors"](path="deck.pptx", colors="accent1=0F6B5C")

    with warnings.catch_warnings():
        warnings.simplefilter("error")
        Presentation(str(tmp_path / "deck.pptx"))


def test_edit_pptx_theme_colors_on_real_bundled_template(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "\n---\n".join(
        ["# Title\nSubtitle", "# First\n- a", "# Second\n- b", "# Closing\nThanks"]
    )
    tools["fill_pptx_template"](
        path="deck.pptx", template_id="bold-statement", content=content
    )

    tools["edit_pptx_theme_colors"](path="deck.pptx", colors="accent1=0F6B5C")
    after = tools["read_pptx_theme_colors"](path="deck.pptx")["colors"]
    assert after["accent1"] == "0F6B5C"


def test_edit_pptx_theme_colors_rejects_unknown_slot(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="Unknown theme color slot"):
        tools["edit_pptx_theme_colors"](path="deck.pptx", colors="notaslot=FFFFFF")


def test_edit_pptx_theme_colors_rejects_bad_hex(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="6-hex-digit"):
        tools["edit_pptx_theme_colors"](path="deck.pptx", colors="accent1=zzzzzz")


def test_edit_pptx_theme_colors_rejects_empty_colors(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- body")

    with pytest.raises(ValueError, match="must not be empty"):
        tools["edit_pptx_theme_colors"](path="deck.pptx", colors="")


def test_edit_pptx_theme_colors_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["edit_pptx_theme_colors"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


def test_read_pptx_theme_colors_is_low_risk(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["read_pptx_theme_colors"])
    assert metadata.requires_approval is False


def test_extra_writable_dir_lets_write_pptx_land_outside_workspace(tmp_path: Path) -> None:
    shared = tmp_path / "shared"
    shared.mkdir()
    workspace = tmp_path / "workspace"
    tools = {
        tool.__name__: tool
        for tool in build_presentation_tools(workspace, extra_writable=[shared])  # type: ignore[attr-defined]
    }

    tools["write_pptx"](path=str(shared / "deck.pptx"), content="# Title\n- body")

    text = tools["read_pptx"](path=str(shared / "deck.pptx"))
    assert "## Title" in text


def test_extra_readable_dir_rejects_write_pptx(tmp_path: Path) -> None:
    downloads = tmp_path / "downloads"
    downloads.mkdir()
    workspace = tmp_path / "workspace"
    tools = {
        tool.__name__: tool
        for tool in build_presentation_tools(workspace, extra_readable=[downloads])  # type: ignore[attr-defined]
    }

    with pytest.raises(PermissionError):
        tools["write_pptx"](path=str(downloads / "deck.pptx"), content="# Title\n- body")


# --- fill_pptx_template ---


def _make_fake_template(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    layout_indexes: list[int],
    *,
    template_id: str = "fake",
    add_decorative_shape: bool = False,
    slide_roles: list[str] | None = None,
) -> dict[str, object]:
    """Build a throwaway .pptx with one slide per entry in layout_indexes
    (a python-pptx stock layout index each), optionally add one decorative
    (non-placeholder) shape to slide 1, monkeypatch
    coscribe.tools.presentations._load_builtin_templates to return exactly
    this one fake TemplateInfo, and return the decorative shape's captured
    (left, top, width, height, fill_rgb) if requested, else {}. `slide_roles`
    defaults to every slide being "content" (the repeatable role) -- tests
    that care about the role-adjustment behavior itself pass an explicit
    list."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    import coscribe.tools.presentations as presentations_module
    from coscribe.tools.pptx_templates import TemplateInfo

    prs = Presentation()
    decorative_shape_info: dict[str, object] = {}
    for index, layout_idx in enumerate(layout_indexes):
        slide = prs.slides.add_slide(prs.slide_layouts[layout_idx])
        if index == 0 and add_decorative_shape:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, 100000, 200000, 300000, 400000
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string("FF00FF")
            decorative_shape_info = {
                "left": shape.left,
                "top": shape.top,
                "width": shape.width,
                "height": shape.height,
                "fill_rgb": str(shape.fill.fore_color.rgb),
            }
    template_path = tmp_path / "fake_template.pptx"
    prs.save(str(template_path))

    info = TemplateInfo(
        id=template_id,
        name="Fake",
        description="fake test template",
        accent="000000",
        slide_count=len(layout_indexes),
        slide_roles=slide_roles if slide_roles is not None else ["content"] * len(layout_indexes),
        dir=tmp_path,
        path=template_path,
        is_widescreen=True,
    )
    monkeypatch.setattr(presentations_module, "_load_builtin_templates", lambda: [info])
    return decorative_shape_info


def test_fill_pptx_template_happy_path_preserves_decorative_shape_and_fills_placeholders(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from pptx import Presentation

    decorative = _make_fake_template(
        tmp_path, monkeypatch, [1], add_decorative_shape=True
    )
    tools = _tools_by_name(tmp_path / "workspace")

    tools["fill_pptx_template"](
        path="out.pptx",
        template_id="fake",
        content="# My Title\n- one\n- two",
    )

    prs = Presentation(str(tmp_path / "workspace" / "out.pptx"))
    slide = prs.slides[0]
    assert slide.shapes.title.text == "My Title"
    body = [
        p
        for p in slide.placeholders
        if p.placeholder_format.idx != slide.shapes.title.placeholder_format.idx
    ][0]
    assert [p.text for p in body.text_frame.paragraphs] == ["one", "two"]

    non_placeholder_shapes = [s for s in slide.shapes if not s.is_placeholder]
    assert len(non_placeholder_shapes) == 1
    shape = non_placeholder_shapes[0]
    assert shape.left == decorative["left"]
    assert shape.top == decorative["top"]
    assert shape.width == decorative["width"]
    assert shape.height == decorative["height"]
    assert str(shape.fill.fore_color.rgb) == decorative["fill_rgb"]


def test_leftover_placeholder_warnings_flags_known_patterns() -> None:
    from pptx import Presentation

    from coscribe.tools.presentations import _leftover_placeholder_warnings

    prs = Presentation()
    cases = [
        "Revenue grew lorem ipsum dolor this quarter",
        "TODO: fill in the real numbers",
        "[insert customer logo here]",
        "xxxxx",
        "This is a placeholder for the chart",
        "Replace this sample text before sending",
    ]
    for text in cases:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = text

    warnings = _leftover_placeholder_warnings(prs)
    assert len(warnings) == len(cases)
    assert [w["slide"] for w in warnings] == list(range(1, len(cases) + 1))


def test_leftover_placeholder_warnings_does_not_flag_real_content() -> None:
    from pptx import Presentation

    from coscribe.tools.presentations import _leftover_placeholder_warnings

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "How to insert a chart: a sample workflow"

    assert _leftover_placeholder_warnings(prs) == []


def test_write_pptx_flags_leftover_todo_in_generated_content(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)

    result = tools["write_pptx"](
        path="deck.pptx", content="# Title\n- TODO: add real figures"
    )

    assert result["placeholder_warnings"] == [{"slide": 1, "text": "TODO"}]


def test_write_pptx_placeholder_warnings_empty_for_clean_content(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)

    result = tools["write_pptx"](path="deck.pptx", content="# Title\n- real content here")

    assert result["placeholder_warnings"] == []


def test_fill_pptx_template_too_few_chunks_for_fixed_slides_raises(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A template with only non-repeatable roles (title + closing, no
    "content") still needs an exact chunk count -- there's nothing to
    duplicate or trim."""
    _make_fake_template(tmp_path, monkeypatch, [1, 1], slide_roles=["title", "closing"])
    tools = _tools_by_name(tmp_path / "workspace")

    with pytest.raises(ValueError, match="needs at least 2 chunk"):
        tools["fill_pptx_template"](
            path="out.pptx", template_id="fake", content="# Only One Chunk\n- x"
        )


def test_fill_pptx_template_repeats_content_slide_when_given_more_chunks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The real point of this feature: a template with one repeatable
    "content" slide accepts more chunks than its own fixed slide count by
    duplicating that slide's design, not by raising."""
    from pptx import Presentation

    _make_fake_template(
        tmp_path, monkeypatch, [1, 1, 1], slide_roles=["title", "content", "closing"]
    )
    tools = _tools_by_name(tmp_path / "workspace")

    content = (
        "# Title\n- t\n---\n# First\n- a\n---\n# Second\n- b\n"
        "---\n# Third\n- c\n---\n# Closing\n- done"
    )
    result = tools["fill_pptx_template"](path="out.pptx", template_id="fake", content=content)
    assert result["slide_count"] == 5

    prs = Presentation(str(tmp_path / "workspace" / "out.pptx"))
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["Title", "First", "Second", "Third", "Closing"]


def test_fill_pptx_template_trims_content_slides_when_given_fewer_chunks(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from pptx import Presentation

    _make_fake_template(
        tmp_path,
        monkeypatch,
        [1, 1, 1, 1],
        slide_roles=["title", "content", "content", "closing"],
    )
    tools = _tools_by_name(tmp_path / "workspace")

    content = "# Title\n- t\n---\n# Only One\n- a\n---\n# Closing\n- done"
    result = tools["fill_pptx_template"](path="out.pptx", template_id="fake", content=content)
    assert result["slide_count"] == 3

    prs = Presentation(str(tmp_path / "workspace" / "out.pptx"))
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["Title", "Only One", "Closing"]


def test_fill_pptx_template_zero_content_chunks_drops_all_content_slides(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from pptx import Presentation

    _make_fake_template(
        tmp_path,
        monkeypatch,
        [1, 1, 1, 1],
        slide_roles=["title", "content", "content", "closing"],
    )
    tools = _tools_by_name(tmp_path / "workspace")

    content = "# Title\n- t\n---\n# Closing\n- done"
    result = tools["fill_pptx_template"](path="out.pptx", template_id="fake", content=content)
    assert result["slide_count"] == 2

    prs = Presentation(str(tmp_path / "workspace" / "out.pptx"))
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["Title", "Closing"]


def test_fill_pptx_template_duplicated_content_slide_keeps_decorative_shape(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A duplicated content slide must carry the same decorative
    (non-placeholder) shape the template's own content slide has --
    otherwise "the design scales with content length" would be false for
    anything but the placeholders."""
    from pptx import Presentation

    decorative = _make_fake_template(
        tmp_path,
        monkeypatch,
        [1, 1],
        add_decorative_shape=True,
        slide_roles=["content", "closing"],
    )
    tools = _tools_by_name(tmp_path / "workspace")

    content = "# First\n- a\n---\n# Second\n- b\n---\n# Closing\n- done"
    tools["fill_pptx_template"](path="out.pptx", template_id="fake", content=content)

    prs = Presentation(str(tmp_path / "workspace" / "out.pptx"))
    for slide in (prs.slides[0], prs.slides[1]):
        non_placeholder = [s for s in slide.shapes if not s.is_placeholder]
        assert len(non_placeholder) == 1
        assert str(non_placeholder[0].fill.fore_color.rgb) == decorative["fill_rgb"]


def test_fill_pptx_template_real_bundled_template_grows_and_shrinks(tmp_path: Path) -> None:
    """Integration-level check against a real bundled template (not the
    throwaway fake harness) -- confirms the manifest's actual
    slide_roles round-trips through the real YAML loader correctly."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)

    grow_content = (
        "# Title\nSubtitle\n---\n# One\n- a\n---\n# Two\n- b\n---\n# Three\n- c\n"
        "---\n# Four\n- d\n---\n# Closing\nThanks\n"
    )
    tools["fill_pptx_template"](
        path="grow.pptx", template_id="bold-statement", content=grow_content
    )
    prs_grow = Presentation(str(tmp_path / "grow.pptx"))
    assert len(prs_grow.slides) == 6

    shrink_content = "# Title\nSubtitle\n---\n# Only\n- a\n---\n# Closing\nThanks\n"
    tools["fill_pptx_template"](
        path="shrink.pptx", template_id="bold-statement", content=shrink_content
    )
    prs_shrink = Presentation(str(tmp_path / "shrink.pptx"))
    assert len(prs_shrink.slides) == 3


def _tools_with_custom_templates_dir(root: Path, templates_dir: Path) -> dict[str, object]:
    return {
        tool.__name__: tool
        for tool in build_presentation_tools(root, custom_templates_dir=templates_dir)  # type: ignore[attr-defined]
    }


def test_extract_pptx_template_classifies_roles_matching_the_real_manifest(
    tmp_path: Path,
) -> None:
    """Extracts from one of coscribe's own bundled templates (used here
    purely as a stand-in "reference deck") -- its own real template.yaml
    is independently-known ground truth to check the role-classification
    heuristic (first=title, last=closing, middle=content) against."""
    import shutil

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    templates_dir = tmp_path / "custom_templates"
    bundled = (
        Path(__file__).parent.parent
        / "src/coscribe/builtin_templates/pptx/bold-statement/template.pptx"
    )
    shutil.copy(bundled, workspace / "reference.pptx")
    tools = _tools_with_custom_templates_dir(workspace, templates_dir)

    result = tools["extract_pptx_template"](
        source_path="reference.pptx",
        template_id="test-extracted",
        name="Test Extracted",
        description="A distilled copy of bold-statement, for testing",
    )

    import re

    assert result["slide_roles"] == ["title", "content", "content", "closing"]
    assert result["slide_count"] == 4
    assert re.match(r"^[0-9A-Fa-f]{6}$", result["accent"])
    assert (templates_dir / "test-extracted" / "template.pptx").is_file()
    assert (templates_dir / "test-extracted" / "template.yaml").is_file()


def test_extract_pptx_template_is_immediately_usable_by_fill_pptx_template(
    tmp_path: Path,
) -> None:
    import shutil

    from pptx import Presentation

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    templates_dir = tmp_path / "custom_templates"
    bundled = (
        Path(__file__).parent.parent
        / "src/coscribe/builtin_templates/pptx/bold-statement/template.pptx"
    )
    shutil.copy(bundled, workspace / "reference.pptx")
    tools = _tools_with_custom_templates_dir(workspace, templates_dir)
    tools["extract_pptx_template"](
        source_path="reference.pptx",
        template_id="my-brand",
        name="My Brand",
        description="test",
    )

    result = tools["fill_pptx_template"](
        path="out.pptx",
        template_id="my-brand",
        content="# Real Title\n- point one\n---\n# Closing\nThanks",
    )

    assert result["template_id"] == "my-brand"
    prs = Presentation(str(workspace / "out.pptx"))
    assert prs.slides[0].shapes.title.text == "Real Title"


def test_extract_pptx_template_rejects_slide_with_no_usable_placeholder(
    tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    templates_dir = tmp_path / "custom_templates"

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "A Title"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])  # blank -- no placeholders at all
    s2.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "plain text"
    prs.save(str(workspace / "messy.pptx"))
    tools = _tools_with_custom_templates_dir(workspace, templates_dir)

    with pytest.raises(ValueError, match="Slide 2 of 'messy.pptx' has no usable"):
        tools["extract_pptx_template"](
            source_path="messy.pptx", template_id="messy", name="Messy", description="test"
        )


def test_extract_pptx_template_rejects_duplicate_id_without_overwrite(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    templates_dir = tmp_path / "custom_templates"
    tools = _tools_with_custom_templates_dir(workspace, templates_dir)
    tools["write_pptx"](path="ref.pptx", content="# Title\n- bullet")
    tools["extract_pptx_template"](
        source_path="ref.pptx", template_id="dup", name="Dup", description="test"
    )

    with pytest.raises(FileExistsError, match="already exists"):
        tools["extract_pptx_template"](
            source_path="ref.pptx", template_id="dup", name="Dup Again", description="test"
        )

    # overwrite=True is the escape hatch
    tools["extract_pptx_template"](
        source_path="ref.pptx",
        template_id="dup",
        name="Dup Replaced",
        description="test",
        overwrite=True,
    )


def test_extract_pptx_template_rejects_bad_template_id(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    templates_dir = tmp_path / "custom_templates"
    tools = _tools_with_custom_templates_dir(workspace, templates_dir)
    tools["write_pptx"](path="ref.pptx", content="# Title\n- bullet")

    with pytest.raises(ValueError, match="lowercase letters/digits/hyphens"):
        tools["extract_pptx_template"](
            source_path="ref.pptx", template_id="Not Valid!", name="X", description="test"
        )


def test_extract_pptx_template_requires_custom_templates_dir_configured(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)  # no custom_templates_dir passed
    tools["write_pptx"](path="ref.pptx", content="# Title\n- bullet")

    with pytest.raises(ValueError, match="No custom templates directory"):
        tools["extract_pptx_template"](
            source_path="ref.pptx", template_id="x", name="X", description="test"
        )


def test_extract_pptx_template_is_write_local_and_gated(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    tools = _tools_with_custom_templates_dir(workspace, tmp_path / "custom_templates")
    metadata = get_tool_metadata(tools["extract_pptx_template"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


def test_fill_pptx_template_chunk_with_heading_but_no_title_placeholder_raises(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _make_fake_template(tmp_path, monkeypatch, [6])  # "Blank" -- no placeholders at all
    tools = _tools_by_name(tmp_path / "workspace")

    with pytest.raises(ValueError, match="no title placeholder"):
        tools["fill_pptx_template"](path="out.pptx", template_id="fake", content="# A Title")


def test_fill_pptx_template_chunk_with_body_but_no_body_placeholder_raises(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _make_fake_template(tmp_path, monkeypatch, [5])  # "Title Only" -- no body placeholder
    tools = _tools_by_name(tmp_path / "workspace")

    with pytest.raises(ValueError, match="no body"):
        tools["fill_pptx_template"](
            path="out.pptx", template_id="fake", content="# Title\n- a bullet"
        )


def test_fill_pptx_template_replaces_existing_placeholder_text_instead_of_appending(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Regression test for a real, live-found bug: coscribe's own bundled
    templates are authored with empty placeholders, so _set_title/
    _fill_content_placeholder appending straight into paragraphs[0] was
    invisible. Any real-world template (confirmed live against an actual
    Microsoft-authored one, "The science of public speaking") ships with
    real sample text already in its placeholders -- appending onto that
    produced garbled runs like "Know your material in advanceNEW BULLET
    A", and any of the template's *own* extra paragraphs past the first
    survived untouched. This pre-fills a fake template's title and body
    placeholder with sample text/multiple paragraphs the way a real
    template would, then asserts fill_pptx_template's output has *only*
    the new content, cleanly."""
    from pptx import Presentation

    _make_fake_template(tmp_path, monkeypatch, [1])
    template_path = tmp_path / "fake_template.pptx"
    prs = Presentation(str(template_path))
    slide = prs.slides[0]
    slide.shapes.title.text_frame.text = "Old Sample Title"
    body = next(
        p
        for p in slide.placeholders
        if p.placeholder_format.idx != slide.shapes.title.placeholder_format.idx
    )
    body.text_frame.text = "Old sample line one"
    body.text_frame.add_paragraph().text = "Old sample line two"
    prs.save(str(template_path))

    tools = _tools_by_name(tmp_path / "workspace")
    tools["fill_pptx_template"](
        path="out.pptx", template_id="fake", content="# New Title\n- new bullet"
    )

    result_prs = Presentation(str(tmp_path / "workspace" / "out.pptx"))
    result_slide = result_prs.slides[0]
    assert result_slide.shapes.title.text == "New Title"
    result_body = next(
        p
        for p in result_slide.placeholders
        if p.placeholder_format.idx != result_slide.shapes.title.placeholder_format.idx
    )
    assert [p.text for p in result_body.text_frame.paragraphs] == ["new bullet"]


def test_find_body_placeholder_skips_a_table_occupying_the_content_placeholder(
    tmp_path: Path,
) -> None:
    """Regression test for a real, live-found crash: a real-world
    template's content placeholder slot can already hold a native table
    (a GraphicFrame, not a text-frame shape) rather than plain text --
    confirmed live against a real Microsoft-authored template's "Agenda"
    slide. Before this fix, _find_body_placeholder still matched it (its
    placeholder_format.type is OBJECT, same as a text one), and
    fill_pptx_template's caller then crashed with an opaque
    AttributeError: 'PlaceholderGraphicFrame' object has no attribute
    'text_frame' trying to write into it -- rather than the same clean,
    actionable "no body/content placeholder" ValueError a slide with no
    eligible placeholder at all already produces. Exercises
    _find_body_placeholder directly against lightweight stand-ins (no real
    OOXML table-in-placeholder XML needed -- python-pptx's own creation
    API in this version can't build one, only read one) rather than
    fill_pptx_template's full public path, which the happy-path/no-body
    tests above already cover end to end."""
    from pptx.enum.shapes import PP_PLACEHOLDER

    from coscribe.tools.presentations import _find_body_placeholder

    class _FakePlaceholderFormat:
        def __init__(self, type_: object, idx: int) -> None:
            self.type = type_
            self.idx = idx

    class _FakeShape:
        def __init__(self, type_: object, idx: int, shape_id: int, has_text_frame: bool) -> None:
            self.placeholder_format = _FakePlaceholderFormat(type_, idx)
            self.shape_id = shape_id
            self.has_text_frame = has_text_frame

    class _FakeShapes:
        def __init__(self, title: object) -> None:
            self.title = title

    class _FakeSlide:
        def __init__(self, title: object, placeholders: list[object]) -> None:
            self.shapes = _FakeShapes(title)
            self.placeholders = placeholders

    title_shape = _FakeShape(PP_PLACEHOLDER.TITLE, idx=0, shape_id=1, has_text_frame=True)
    table_shape = _FakeShape(PP_PLACEHOLDER.OBJECT, idx=1, shape_id=2, has_text_frame=False)
    slide = _FakeSlide(title_shape, [title_shape, table_shape])

    assert _find_body_placeholder(slide) is None


def test_fill_pptx_template_rejects_layout_directive_naming_write_pptx(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _make_fake_template(tmp_path, monkeypatch, [1])
    tools = _tools_by_name(tmp_path / "workspace")

    with pytest.raises(ValueError, match="write_pptx"):
        tools["fill_pptx_template"](
            path="out.pptx",
            template_id="fake",
            content="layout: icon-list\n# Title\n- x",
        )


def test_fill_pptx_template_rejects_table_naming_write_pptx(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _make_fake_template(tmp_path, monkeypatch, [1])
    tools = _tools_by_name(tmp_path / "workspace")

    with pytest.raises(ValueError, match="write_pptx"):
        tools["fill_pptx_template"](
            path="out.pptx",
            template_id="fake",
            content="# Title\n| a | b |\n| --- | --- |\n| 1 | 2 |",
        )


def test_fill_pptx_template_unknown_template_id_raises(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _make_fake_template(tmp_path, monkeypatch, [1], template_id="fake")
    tools = _tools_by_name(tmp_path / "workspace")

    with pytest.raises(ValueError, match="Unknown template"):
        tools["fill_pptx_template"](path="out.pptx", template_id="nope", content="# Title")


def test_fill_pptx_template_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["fill_pptx_template"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


def test_fill_pptx_template_respects_overwrite_false(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _make_fake_template(tmp_path, monkeypatch, [1])
    tools = _tools_by_name(tmp_path / "workspace")
    tools["fill_pptx_template"](path="out.pptx", template_id="fake", content="# Title\n- x")

    with pytest.raises(FileExistsError):
        tools["fill_pptx_template"](
            path="out.pptx", template_id="fake", content="# Title\n- x", overwrite=False
        )


def test_edit_pptx_text_changes_title_and_content_leaving_the_rest_untouched(
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path / "workspace")
    tools["write_pptx"](
        path="deck.pptx",
        content="# Slide One\n- original a\n- original b\n\n---\n\n# Slide Two\n- untouched",
    )

    tools["edit_pptx_text"](
        path="deck.pptx", slide=1, title="New Title", content="- new bullet"
    )

    prs = Presentation(str(tmp_path / "workspace" / "deck.pptx"))
    slide1 = prs.slides[0]
    assert slide1.shapes.title.text == "New Title"
    body = next(
        p
        for p in slide1.placeholders
        if p.placeholder_format.idx != slide1.shapes.title.placeholder_format.idx
    )
    assert [p.text for p in body.text_frame.paragraphs] == ["new bullet"]
    # Slide 2 (not addressed by this call) is exactly as write_pptx left it.
    slide2 = prs.slides[1]
    assert slide2.shapes.title.text == "Slide Two"


def test_edit_pptx_text_title_only_leaves_existing_content_placeholder_alone(
    tmp_path: Path,
) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path / "workspace")
    tools["write_pptx"](path="deck.pptx", content="# Old Title\n- keep me")

    tools["edit_pptx_text"](path="deck.pptx", slide=1, title="Renamed")

    prs = Presentation(str(tmp_path / "workspace" / "deck.pptx"))
    slide = prs.slides[0]
    assert slide.shapes.title.text == "Renamed"
    body = next(
        p
        for p in slide.placeholders
        if p.placeholder_format.idx != slide.shapes.title.placeholder_format.idx
    )
    assert body.text_frame.text == "keep me"


def test_edit_pptx_text_content_only_leaves_existing_title_alone(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path / "workspace")
    tools["write_pptx"](path="deck.pptx", content="# Keep Me\n- old bullet")

    tools["edit_pptx_text"](path="deck.pptx", slide=1, content="- replaced bullet")

    prs = Presentation(str(tmp_path / "workspace" / "deck.pptx"))
    slide = prs.slides[0]
    assert slide.shapes.title.text == "Keep Me"
    body = next(
        p
        for p in slide.placeholders
        if p.placeholder_format.idx != slide.shapes.title.placeholder_format.idx
    )
    assert body.text_frame.text == "replaced bullet"


def test_edit_pptx_text_neither_title_nor_content_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path / "workspace")
    tools["write_pptx"](path="deck.pptx", content="# Title\n- x")

    with pytest.raises(ValueError, match="at least one of title or content"):
        tools["edit_pptx_text"](path="deck.pptx", slide=1)


def test_edit_pptx_text_title_on_slide_with_no_title_placeholder_raises(
    tmp_path: Path,
) -> None:
    tools = _tools_by_name(tmp_path / "workspace")
    # A table-only slide chunk still gets a title placeholder from
    # _TITLE_ONLY_LAYOUT in this codebase's write_pptx, so build a slide
    # with none the same way _make_fake_template's "Blank" layout does.
    from pptx import Presentation

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # "Blank" -- no placeholders at all
    (tmp_path / "workspace").mkdir(parents=True, exist_ok=True)
    prs.save(str(tmp_path / "workspace" / "blank.pptx"))

    with pytest.raises(ValueError, match="no title placeholder"):
        tools["edit_pptx_text"](path="blank.pptx", slide=1, title="New Title")


def test_edit_pptx_text_rejects_a_table_in_content(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path / "workspace")
    tools["write_pptx"](path="deck.pptx", content="# Title\n- x")

    with pytest.raises(ValueError, match="doesn't support tables"):
        tools["edit_pptx_text"](
            path="deck.pptx", slide=1, content="| a | b |\n| --- | --- |\n| 1 | 2 |"
        )


def test_edit_pptx_text_out_of_range_placeholder_index_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path / "workspace")
    tools["write_pptx"](path="deck.pptx", content="# Title\n- x")

    with pytest.raises(ValueError, match="placeholder_index=1"):
        tools["edit_pptx_text"](
            path="deck.pptx", slide=1, content="- y", placeholder_index=1
        )


def test_edit_pptx_text_placeholder_index_reaches_the_second_content_placeholder(
    tmp_path: Path,
) -> None:
    """Regression coverage for the real gap fill_pptx_template still has
    (see its own tests/docstring): a "two column" slide has two
    independent content placeholders, and placeholder_index lets
    edit_pptx_text address the second one specifically -- verified live
    against a real Microsoft-authored template's identical layout shape."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path / "workspace")
    tools["write_pptx"](
        path="deck.pptx",
        content="layout: two-column\n# Two Columns\n- left one\n\n>>>\n\n- right one",
    )

    tools["edit_pptx_text"](
        path="deck.pptx", slide=1, content="- left replaced", placeholder_index=0
    )
    tools["edit_pptx_text"](
        path="deck.pptx", slide=1, content="- right replaced", placeholder_index=1
    )

    prs = Presentation(str(tmp_path / "workspace" / "deck.pptx"))
    slide = prs.slides[0]
    bodies = [
        p
        for p in slide.placeholders
        if p.placeholder_format.idx != slide.shapes.title.placeholder_format.idx
    ]
    assert [p.text_frame.text for p in bodies] == ["left replaced", "right replaced"]


def test_edit_pptx_text_is_write_local_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path / "workspace")
    metadata = get_tool_metadata(tools["edit_pptx_text"])  # type: ignore[arg-type]
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


def test_table_slide_strips_default_style_and_uses_hairline_borders(tmp_path: Path) -> None:
    """Regression test for PPTX_DESIGN.md §0's real, user-reported visual-
    quality defect: a table built with no styling at all renders with
    PowerPoint's default banded-blue table style, which a real user
    directly compared against a competing tool's plainer-but-deliberately-
    designed table and found worse. Asserts the default style is actually
    stripped (first_row/horz_banding off) and every cell's tcPr border
    children are in ECMA-376's required schema order (lnL, lnR, lnT, lnB,
    then a fill element) -- inserting out of order round-trips fine
    through python-pptx's own lenient reader but real PowerPoint enforces
    the schema strictly, so this is the one thing a plain re-read can't
    catch on its own."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    content = (
        "# Regions\n"
        "| Region | Rep | Attainment |\n"
        "| --- | --- | --- |\n"
        "| North | Ada | 112% |\n"
        "| South | Bo | 98% |\n"
    )
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    table = next(s for s in prs.slides[0].shapes if s.has_table).table
    assert table.first_row is False
    assert table.horz_banding is False

    for row_index in range(len(table.rows)):
        for col_index in range(len(table.columns)):
            tc_pr = table.cell(row_index, col_index)._tc.get_or_add_tcPr()
            tags = [child.tag.split("}")[-1] for child in tc_pr]
            border_tags = [t for t in tags if t in ("lnL", "lnR", "lnT", "lnB")]
            assert border_tags == ["lnL", "lnR", "lnT", "lnB"]
            fill_index = next(
                i for i, t in enumerate(tags) if t in ("noFill", "solidFill")
            )
            assert fill_index == len(border_tags)  # every ln* comes before the fill


def test_table_slide_right_aligns_numeric_cells_and_flags_over_100_percent(
    tmp_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.text import PP_ALIGN

    tools = _tools_by_name(tmp_path)
    content = (
        "# Regions\n"
        "| Region | Rep | Attainment |\n"
        "| --- | --- | --- |\n"
        "| North | Ada | 112% |\n"
        "| South | Bo | 98% |\n"
    )
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    table = next(s for s in prs.slides[0].shapes if s.has_table).table

    header_cell = table.cell(0, 0)
    assert header_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT
    assert header_cell.text_frame.paragraphs[0].runs[0].font.bold is True
    # theme_color, not a literal RGB, so a table stays readable if
    # write_pptx's `theme` parameter later sets a dark background --
    # see test_table_text_follows_theme_text_token below.
    assert header_cell.text_frame.paragraphs[0].runs[0].font.color.theme_color == (
        MSO_THEME_COLOR.DARK_1
    )

    region_cell = table.cell(1, 0)  # "North" -- not numeric-looking
    assert region_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT

    over_100_cell = table.cell(1, 2)  # "112%"
    assert over_100_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
    over_100_run = over_100_cell.text_frame.paragraphs[0].runs[0]
    assert over_100_run.font.bold is True
    assert str(over_100_run.font.color.rgb) == "B3261E"

    under_100_cell = table.cell(2, 2)  # "98%"
    under_100_run = under_100_cell.text_frame.paragraphs[0].runs[0]
    assert under_100_run.font.bold is not True
    assert under_100_run.font.color.theme_color == MSO_THEME_COLOR.DARK_1


def test_stat_callout_cards_are_bordered_not_solid_filled_with_hierarchy(
    tmp_path: Path,
) -> None:
    """Regression test for PPTX_DESIGN.md §0's other reported defect: cards
    used to fill solid with the accent color and center two same-weight
    white lines -- a flat block, not a "designed" card. Asserts the new
    shape: white/bordered card, muted label first (smaller), accent-
    colored emphasized stat second (larger), both left-aligned."""
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN

    tools = _tools_by_name(tmp_path)
    content = (
        "layout: stat-callout 1E2761\n"
        "# Highlights\n"
        "| Stat | Label |\n"
        "| --- | --- |\n"
        "| 3x | Revenue growth |\n"
        "| 98% | Retention |\n"
    )
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    cards = [s for s in prs.slides[0].shapes if s.shape_type == 1]
    card = next(c for c in cards if "3x" in c.text_frame.text)

    assert card.fill.type is not None
    assert str(card.fill.fore_color.rgb) == "FFFFFF"
    assert str(card.line.color.rgb) == "E0E0E0"

    label_paragraph, stat_paragraph = card.text_frame.paragraphs
    assert label_paragraph.alignment == PP_ALIGN.LEFT
    assert label_paragraph.text == "Revenue growth"
    label_run = label_paragraph.runs[0]
    assert label_run.font.size.pt == 12
    assert str(label_run.font.color.rgb) == "707070"

    assert stat_paragraph.alignment == PP_ALIGN.LEFT
    assert stat_paragraph.text == "3x"
    stat_run = stat_paragraph.runs[0]
    assert stat_run.font.bold is True
    assert stat_run.font.size.pt == 28
    assert str(stat_run.font.color.rgb) == "1E2761"


def test_table_slide_is_centered_on_the_real_widescreen_slide(tmp_path: Path) -> None:
    """Regression test for a real, user-reported bug caught in a rendered
    screenshot: the table's position used to be hardcoded for python-
    pptx's own default 10in-wide blank Presentation(), but write_pptx
    always resizes the deck to widescreen (13.333in) -- on the real
    width the table actually renders at, that left a large dead gap on
    the right, visibly off-center. Asserts the table's left and right
    margins are now equal."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    content = "# Regions\n| Region | Attainment |\n| --- | --- |\n| North | 112% |\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    table_shape = next(s for s in prs.slides[0].shapes if s.has_table)
    left_margin = table_shape.left
    right_margin = prs.slide_width - (table_shape.left + table_shape.width)
    assert left_margin == right_margin


def test_stat_callout_card_height_is_not_mostly_empty_space(tmp_path: Path) -> None:
    """Regression test for a real, user-reported bug caught in a rendered
    screenshot: a card sized for a near-square aspect ratio left a large
    empty gap below its two short text lines (a label + one emphasized
    number). Asserts the card's height/width ratio is meaningfully
    smaller than the old 0.7 default, without pinning an exact value.
    Uses 3 stats, not 1 -- with only one card spanning the full content
    width, height was already capped by the content area's own height
    regardless of the width-derived multiplier, masking the old 0.7
    ratio entirely (confirmed: this test genuinely fails against the
    pre-fix code only with multiple, narrower cards)."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    content = (
        "layout: stat-callout\n"
        "# Highlights\n"
        "| Stat | Label |\n"
        "| --- | --- |\n"
        "| 3x | Revenue growth |\n"
        "| 12 | New markets |\n"
        "| 98% | Retention |\n"
    )
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    card = next(s for s in prs.slides[0].shapes if s.shape_type == 1)
    assert card.height / card.width < 0.5


def test_plain_paragraph_body_text_has_no_bullet_marker(tmp_path: Path) -> None:
    """Regression test for a real, user-reported bug caught in a rendered
    screenshot: a slide with a single flowing paragraph (no leading -/*,
    so Block(kind="paragraph")) rendered with a meaningless bullet glyph
    in front of it, because the "Title and Content" placeholder this
    lands in auto-bullets every paragraph by inherited layout default
    and nothing ever overrode that for plain-paragraph content. Asserts
    the paragraph's own <a:pPr> now has an explicit <a:buNone/>."""
    from pptx import Presentation
    from pptx.oxml.ns import qn

    tools = _tools_by_name(tmp_path)
    content = "# Context\nA single flowing paragraph, not a list item at all.\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    slide = prs.slides[0]
    body = next(
        p
        for p in slide.placeholders
        if p.placeholder_format.idx != slide.shapes.title.placeholder_format.idx
    )
    p_pr = body.text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    assert p_pr is not None
    assert p_pr.find(qn("a:buNone")) is not None


def test_bullet_body_text_still_has_no_explicit_bullet_override(tmp_path: Path) -> None:
    """A real markdown bullet (leading "-") must NOT get the paragraph-
    kind buNone override -- its existing, not-reported-as-broken
    behavior (inheriting the placeholder's own bullet character) stays
    exactly as it was."""
    from pptx import Presentation
    from pptx.oxml.ns import qn

    tools = _tools_by_name(tmp_path)
    content = "# Context\n- a real bullet point\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    slide = prs.slides[0]
    body = next(
        p
        for p in slide.placeholders
        if p.placeholder_format.idx != slide.shapes.title.placeholder_format.idx
    )
    p_pr = body.text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    assert p_pr is None or p_pr.find(qn("a:buNone")) is None


# --- write_pptx `theme` parameter (deck-wide color/typography identity) ---


def _theme_xml_parts(pptx_path: Path) -> tuple[object, object, object]:
    """(clrScheme element, majorFont element, master) for a saved deck's
    own theme part -- python-pptx has no public API for clrScheme/
    fontScheme (see _apply_theme's docstring), so tests read the raw XML
    the same way the implementation writes it."""
    from lxml import etree
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    prs = Presentation(str(pptx_path))
    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    clr_scheme = root.find(".//a:clrScheme", ns)
    major_font = root.find(".//a:fontScheme/a:majorFont", ns)
    return clr_scheme, major_font, master


def test_theme_token_rejects_malformed_pair(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    with pytest.raises(ValueError, match="Malformed theme token"):
        tools["write_pptx"](path="deck.pptx", content="# Title\n- x", theme="bg")


def test_theme_token_rejects_unknown_key(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    with pytest.raises(ValueError, match="Unknown theme token"):
        tools["write_pptx"](path="deck.pptx", content="# Title\n- x", theme="foo=bar")


def test_theme_token_rejects_non_hex_color(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    with pytest.raises(ValueError, match="6-hex-digit color"):
        tools["write_pptx"](path="deck.pptx", content="# Title\n- x", theme="bg=notacolor")


def test_theme_sets_background_and_clr_scheme_and_font_scheme(tmp_path: Path) -> None:
    """The four color tokens land on the deck's own <a:clrScheme> (not
    just the visible background), and heading_font lands on
    <a:fontScheme>'s majorFont -- so any shape/placeholder that
    references a theme color or inherits the master's font (not just the
    background fill) picks up the new identity too, not only the parts
    of the deck this test can see without reopening the raw XML."""
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](
        path="deck.pptx",
        content="# Title\n- bullet\n",
        theme="bg=0F172A,text=F8FAFC,surface=1E293B,accent=38BDF8,heading_font=Georgia",
    )

    clr_scheme, major_font, master = _theme_xml_parts(tmp_path / "deck.pptx")
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

    def scheme_color(tag: str) -> str:
        el = clr_scheme.find(f"a:{tag}", ns).find("a:srgbClr", ns)
        return el.get("val")

    assert scheme_color("lt1") == "0F172A"
    assert scheme_color("dk1") == "F8FAFC"
    assert scheme_color("lt2") == "1E293B"
    assert scheme_color("accent1") == "38BDF8"
    assert scheme_color("hlink") == "38BDF8"
    assert major_font.find("a:latin", ns).get("typeface") == "Georgia"

    from pptx.dml.color import RGBColor

    assert master.background.fill.type is not None
    assert master.background.fill.fore_color.rgb == RGBColor.from_string("0F172A")


def test_no_theme_parameter_leaves_stock_colors_and_plain_background(tmp_path: Path) -> None:
    """Omitting `theme` entirely (the existing, default call shape) must
    change nothing -- no regression for every deck that doesn't opt in."""
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")

    clr_scheme, _major_font, master = _theme_xml_parts(tmp_path / "deck.pptx")
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    assert clr_scheme.find("a:lt1", ns).find("a:sysClr", ns) is not None
    assert clr_scheme.find("a:accent1", ns).find("a:srgbClr", ns).get("val") == "4F81BD"
    assert str(master.background.fill.type) == "BACKGROUND (5)"


def test_icon_list_and_stat_callout_follow_theme_accent_when_unspecified(
    tmp_path: Path,
) -> None:
    """A `layout: icon-list`/`layout: stat-callout` slide with no
    explicit ACCENTHEX in its own directive line must render its accent
    shapes as a theme_color reference (accent1), not a literal color --
    otherwise write_pptx's `theme` parameter would have no visible effect
    on these two prefab layouts at all, repeating the exact "theme
    editing is inert for decorative shapes" gap independently confirmed
    in a competing tool's own output (see PPTX_DESIGN.md)."""
    from pptx import Presentation
    from pptx.enum.dml import MSO_THEME_COLOR
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    tools = _tools_by_name(tmp_path)
    content = (
        "layout: icon-list\n"
        "# Icons\n"
        "- [A] alpha\n"
        "---\n"
        "layout: stat-callout\n"
        "# Stats\n"
        "| Stat | Label |\n"
        "| --- | --- |\n"
        "| 5 | Five |\n"
    )
    tools["write_pptx"](path="deck.pptx", content=content, theme="accent=38BDF8")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    icon = next(s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
    assert icon.fill.fore_color.theme_color == MSO_THEME_COLOR.ACCENT_1

    stat_run = next(
        run
        for shape in prs.slides[1].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.text == "5"
    )
    assert stat_run.font.color.theme_color == MSO_THEME_COLOR.ACCENT_1


def test_icon_list_explicit_accent_still_overrides_theme(tmp_path: Path) -> None:
    """A slide's own `layout: icon-list ACCENTHEX` must still win over the
    deck's `theme` accent -- an explicit per-slide override, unchanged
    from before this feature existed."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    tools = _tools_by_name(tmp_path)
    content = "layout: icon-list FF0000\n# Icons\n- [A] alpha\n"
    tools["write_pptx"](path="deck.pptx", content=content, theme="accent=38BDF8")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    icon = next(s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)
    assert icon.fill.fore_color.rgb == RGBColor.from_string("FF0000")


def test_write_pptx_default_deck_gets_cjk_typeface(tmp_path: Path) -> None:
    """Regression test for a real, diagnosed-but-previously-unfixed bug
    (PPTX_DESIGN.md §3): every stock python-pptx theme ships an empty
    East-Asian typeface, so Chinese/Japanese/Korean text rendered
    unpredictably depending on the opening machine's own font fallback.
    Applies unconditionally, not just when `theme` is given."""
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# 标题\n- 内容\n")

    _clr_scheme, major_font, _master = _theme_xml_parts(tmp_path / "deck.pptx")
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    assert major_font.find("a:ea", ns).get("typeface") == "Microsoft YaHei"


def test_write_pptx_template_path_deck_keeps_its_own_fonts(tmp_path: Path) -> None:
    """The CJK-typeface fix must NOT touch a user-supplied template_path
    file -- silently rewriting fonts in someone's own uploaded/branded
    template would be an unwanted side effect, unlike write_pptx's own
    freshly-created default deck or coscribe's own bundled templates."""
    from lxml import etree
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    tools = _tools_by_name(tmp_path)
    # Build a template.pptx via write_pptx (gets the CJK fix applied),
    # then blank its ea typeface back out directly via XML -- simulating
    # a real user-authored template that never had one configured, since
    # that's the state this test needs to start from.
    tools["write_pptx"](path="template.pptx", content="# Title\n- bullet\n")
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    prs = Presentation(str(tmp_path / "template.pptx"))
    theme_part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    root.find(".//a:fontScheme/a:majorFont/a:ea", ns).set("typeface", "")
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    prs.save(str(tmp_path / "template.pptx"))

    tools["write_pptx"](
        path="deck.pptx", content="# Title\n- bullet\n", template_path="template.pptx"
    )
    _clr_scheme, major_font, _master = _theme_xml_parts(tmp_path / "deck.pptx")
    assert major_font.find("a:ea", ns).get("typeface") == ""


# --- list_pptx_shapes / edit_pptx_shape (Tier 1: editing an existing deck) ---


def test_list_pptx_shapes_describes_icon_list_shapes(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "layout: icon-list FF0000\n# Icons\n- [A] alpha\n- [B] beta\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    result = tools["list_pptx_shapes"](path="deck.pptx", slide=1)
    assert result["slide"] == 1
    shapes = result["shapes"]
    assert result["shape_count"] == len(shapes) == 5

    title = shapes[0]
    assert title["is_placeholder"] is True
    assert title["text_preview"] == "Icons"

    icon_a = shapes[1]
    assert icon_a["text_preview"] == "A"
    assert icon_a["fill"] == "#FF0000"
    assert icon_a["is_table"] is False
    assert isinstance(icon_a["left_in"], float)
    assert isinstance(icon_a["rotation"], float)


def test_list_pptx_shapes_reports_table_dimensions(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "# Table\n| A | B | C |\n| --- | --- | --- |\n| 1 | 2 | 3 |\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    result = tools["list_pptx_shapes"](path="deck.pptx", slide=1)
    table_shape = next(s for s in result["shapes"] if s["is_table"])
    assert table_shape["table_dimensions"] == {"rows": 2, "cols": 3}
    assert table_shape["fill"] is None  # GraphicFrame has no .fill at all


def test_list_pptx_shapes_theme_color_fill_reports_theme_reference(tmp_path: Path) -> None:
    """An icon-list slide with no explicit accent (follows the deck's
    theme accent1, see write_pptx's `theme` parameter) must report its
    fill as a "theme:" reference, not try (and fail) to resolve it to a
    literal RGB -- a SchemeColor has no .rgb property at all."""
    tools = _tools_by_name(tmp_path)
    content = "layout: icon-list\n# Icons\n- [A] alpha\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    result = tools["list_pptx_shapes"](path="deck.pptx", slide=1)
    icon = next(s for s in result["shapes"] if s["text_preview"] == "A")
    assert icon["fill"] == "theme:ACCENT_1"


def test_list_pptx_shapes_rejects_out_of_range_slide(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    with pytest.raises(ValueError, match="out of range"):
        tools["list_pptx_shapes"](path="deck.pptx", slide=2)


# --- SmartArt detection/extraction + delete_pptx_shape ---
#
# python-pptx has no SmartArt-creation API at all (and neither does any
# other library -- the layout algorithm lives in PowerPoint itself, not
# the file format), so there's no way to build a real SmartArt-bearing
# .pptx through python-pptx's own object model the way every other
# fixture in this file is built. `_write_smartart_fixture` instead
# hand-writes the real OOXML structure directly into a python-pptx-saved
# file's zip package -- the exact <a:graphicData uri=".../diagram">/
# <dgm:relIds r:dm=.../<dgm:pt>/<dgm:t> shape cross-checked against two
# independent, real, working OOXML-consuming implementations (Apache
# POI's XSLFDiagram, LibreOffice's oox diagram filter -- see this
# module's own _SMARTART_NS comment), not guessed.


def _write_smartart_fixture(path: Path) -> None:
    import zipfile

    from pptx import Presentation

    base_path = path.with_suffix(".base.pptx")
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(base_path))

    dgm = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"

    data_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<dgm:dataModel xmlns:dgm="{dgm}" xmlns:a="{a_ns}">
  <dgm:ptLst>
    <dgm:pt modelId="0" type="doc"><dgm:prSet/></dgm:pt>
    <dgm:pt modelId="1"><dgm:t><a:bodyPr/><a:lstStyle/>
      <a:p><a:r><a:t>Plan</a:t></a:r></a:p></dgm:t></dgm:pt>
    <dgm:pt modelId="2"><dgm:t><a:bodyPr/><a:lstStyle/>
      <a:p><a:r><a:t>Build</a:t></a:r><a:r><a:t> phase</a:t></a:r></a:p></dgm:t></dgm:pt>
    <dgm:pt modelId="3"><dgm:t><a:bodyPr/><a:lstStyle/>
      <a:p><a:r><a:t>Ship</a:t></a:r></a:p></dgm:t></dgm:pt>
    <dgm:pt modelId="4" type="parTrans" cxnId="x"><dgm:prSet/></dgm:pt>
  </dgm:ptLst>
  <dgm:cxnLst/>
  <dgm:bg/>
  <dgm:whole/>
</dgm:dataModel>"""
    layout_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<dsp:dataModel xmlns:dsp="http://schemas.microsoft.com/office/drawing/2008/diagram"/>'
    )
    qs_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<dgm:styleDefHdrLst xmlns:dgm="{dgm}"/>'
    )
    colors_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<dgm:colorsDefHdrLst xmlns:dgm="{dgm}"/>'
    )
    graphicframe_xml = f"""<p:graphicFrame
      xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
      xmlns:a="{a_ns}" xmlns:r="{r_ns}">
  <p:nvGraphicFramePr>
    <p:cNvPr id="99" name="SmartArt 1"/>
    <p:cNvGraphicFramePr/>
    <p:nvPr/>
  </p:nvGraphicFramePr>
  <p:xfrm><a:off x="914400" y="914400"/><a:ext cx="4572000" cy="2743200"/></p:xfrm>
  <a:graphic>
    <a:graphicData uri="{dgm}">
      <dgm:relIds xmlns:dgm="{dgm}" r:dm="rIdDM1" r:lo="rIdLO1" r:qs="rIdQS1" r:cs="rIdCS1"/>
    </a:graphicData>
  </a:graphic>
</p:graphicFrame>"""

    with zipfile.ZipFile(base_path, "r") as zin:
        contents = {name: zin.read(name) for name in zin.namelist()}

    slide_xml = contents["ppt/slides/slide1.xml"].decode("utf-8")
    slide_xml = slide_xml.replace("</p:spTree>", graphicframe_xml + "</p:spTree>")

    slide_rels_text = contents["ppt/slides/_rels/slide1.xml.rels"].decode("utf-8")
    rel_type_base = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    new_rels = (
        f'<Relationship Id="rIdDM1" Type="{rel_type_base}/diagramData" '
        'Target="../diagrams/data1.xml"/>'
        f'<Relationship Id="rIdLO1" Type="{rel_type_base}/diagramLayout" '
        'Target="../diagrams/layout1.xml"/>'
        f'<Relationship Id="rIdQS1" Type="{rel_type_base}/diagramQuickStyle" '
        'Target="../diagrams/quickStyle1.xml"/>'
        f'<Relationship Id="rIdCS1" Type="{rel_type_base}/diagramColors" '
        'Target="../diagrams/colors1.xml"/>'
    )
    slide_rels_text = slide_rels_text.replace("</Relationships>", new_rels + "</Relationships>")

    contents["ppt/slides/slide1.xml"] = slide_xml.encode("utf-8")
    contents["ppt/slides/_rels/slide1.xml.rels"] = slide_rels_text.encode("utf-8")
    contents["ppt/diagrams/data1.xml"] = data_xml.encode("utf-8")
    contents["ppt/diagrams/layout1.xml"] = layout_xml.encode("utf-8")
    contents["ppt/diagrams/quickStyle1.xml"] = qs_xml.encode("utf-8")
    contents["ppt/diagrams/colors1.xml"] = colors_xml.encode("utf-8")

    content_types = contents["[Content_Types].xml"].decode("utf-8")
    overrides = (
        '<Override PartName="/ppt/diagrams/data1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml"/>'
        '<Override PartName="/ppt/diagrams/layout1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagramLayout+xml"/>'
        '<Override PartName="/ppt/diagrams/quickStyle1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagramStyle+xml"/>'
        '<Override PartName="/ppt/diagrams/colors1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagramColors+xml"/>'
    )
    content_types = content_types.replace("</Types>", overrides + "</Types>")
    contents["[Content_Types].xml"] = content_types.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in contents.items():
            zout.writestr(name, data)
    base_path.unlink()


def test_list_pptx_shapes_detects_smartart_and_extracts_node_text(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _write_smartart_fixture(tmp_path / "smartart.pptx")

    shapes = tools["list_pptx_shapes"](path="smartart.pptx", slide=1)["shapes"]
    assert len(shapes) == 1
    assert shapes[0]["is_smartart"] is True
    assert shapes[0]["smartart_text"] == ["Plan", "Build phase", "Ship"]


def test_list_pptx_shapes_non_smartart_shape_reports_false_and_none(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")

    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    for shape in shapes:
        assert shape["is_smartart"] is False
        assert shape["smartart_text"] is None


def test_delete_pptx_shape_removes_smartart_and_survives_round_trip(tmp_path: Path) -> None:
    import warnings

    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    _write_smartart_fixture(tmp_path / "smartart.pptx")

    result = tools["delete_pptx_shape"](path="smartart.pptx", slide=1, shape_index=0)
    assert result == {"path": "smartart.pptx", "slide": 1, "shape_index": 0}

    with warnings.catch_warnings():
        warnings.simplefilter("error")
        prs = Presentation(str(tmp_path / "smartart.pptx"))
    assert list(prs.slides[0].shapes) == []


def test_delete_pptx_shape_removes_only_that_shape(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    before = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]

    tools["delete_pptx_shape"](path="deck.pptx", slide=1, shape_index=1)

    after = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    assert len(after) == len(before) - 1
    assert after[0]["index"] == 0
    assert after[0]["text_preview"] == before[0]["text_preview"]


def test_delete_pptx_shape_rejects_out_of_range_shape_index(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")

    with pytest.raises(ValueError, match="2 shape"):
        tools["delete_pptx_shape"](path="deck.pptx", slide=1, shape_index=99)


def test_delete_pptx_shape_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["delete_pptx_shape"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


def test_edit_pptx_shape_changes_geometry_rotation_and_fill(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Emu, Inches

    tools = _tools_by_name(tmp_path)
    content = "layout: icon-list FF0000\n# Icons\n- [A] alpha\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    result = tools["edit_pptx_shape"](
        path="deck.pptx",
        slide=1,
        shape_index=1,
        left_in=2.0,
        top_in=3.0,
        width_in=1.5,
        height_in=1.5,
        rotation=45.0,
        fill_color="00FF00",
    )
    assert result["shape"]["left_in"] == 2.0
    assert result["shape"]["rotation"] == 45.0
    assert result["shape"]["fill"] == "#00FF00"

    prs = Presentation(str(tmp_path / "deck.pptx"))
    icon = list(prs.slides[0].shapes)[1]
    assert icon.left == Inches(2.0)
    assert icon.top == Inches(3.0)
    assert icon.width == Emu(Inches(1.5))
    assert icon.height == Emu(Inches(1.5))
    assert icon.rotation == 45.0
    assert icon.fill.fore_color.rgb == RGBColor.from_string("00FF00")


def test_edit_pptx_shape_partial_update_leaves_other_properties_alone(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    content = "layout: icon-list FF0000\n# Icons\n- [A] alpha\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    before = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"][1]
    tools["edit_pptx_shape"](path="deck.pptx", slide=1, shape_index=1, rotation=90.0)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    icon = list(prs.slides[0].shapes)[1]
    assert icon.rotation == 90.0
    assert round(icon.left / 914400, 2) == before["left_in"]
    assert round(icon.top / 914400, 2) == before["top_in"]
    assert str(icon.fill.fore_color.rgb) == "FF0000"  # untouched


def test_edit_pptx_shape_rejects_no_properties_given(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="layout: icon-list\n# Icons\n- [A] alpha\n")
    with pytest.raises(ValueError, match="at least one property"):
        tools["edit_pptx_shape"](path="deck.pptx", slide=1, shape_index=1)


def test_edit_pptx_shape_rejects_bad_fill_color(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="layout: icon-list\n# Icons\n- [A] alpha\n")
    with pytest.raises(ValueError, match="6-hex-digit color"):
        tools["edit_pptx_shape"](path="deck.pptx", slide=1, shape_index=1, fill_color="#00FF00")


def test_edit_pptx_shape_rejects_out_of_range_shape_index(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    with pytest.raises(ValueError, match="shape_index"):
        tools["edit_pptx_shape"](path="deck.pptx", slide=1, shape_index=99, rotation=10.0)


def test_edit_pptx_shape_rejects_fill_color_on_a_table(tmp_path: Path) -> None:
    """A table (GraphicFrame) has no .fill attribute at all -- must raise
    a clear error, not an unhandled AttributeError."""
    tools = _tools_by_name(tmp_path)
    content = "# Table\n| A | B |\n| --- | --- |\n| 1 | 2 |\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    table_index = next(s["index"] for s in shapes if s["is_table"])
    with pytest.raises(ValueError, match="no fill to set"):
        tools["edit_pptx_shape"](
            path="deck.pptx", slide=1, shape_index=table_index, fill_color="00FF00"
        )


def test_edit_pptx_shape_sets_two_color_gradient_fill(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.dml import MSO_FILL

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="layout: icon-list\n# Icons\n- [A] alpha\n")

    tools["edit_pptx_shape"](
        path="deck.pptx", slide=1, shape_index=1, fill_color="0F172A", fill_color_2="38BDF8"
    )

    prs = Presentation(tmp_path / "deck.pptx")
    shape = list(prs.slides[0].shapes)[1]
    assert shape.fill.type == MSO_FILL.GRADIENT
    stops = shape.fill.gradient_stops
    assert str(stops[0].color.rgb) == "0F172A"
    assert str(stops[1].color.rgb) == "38BDF8"


def test_edit_pptx_shape_gradient_angle_applies(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="layout: icon-list\n# Icons\n- [A] alpha\n")

    tools["edit_pptx_shape"](
        path="deck.pptx",
        slide=1,
        shape_index=1,
        fill_color="0F172A",
        fill_color_2="38BDF8",
        gradient_angle=45.0,
    )

    prs = Presentation(tmp_path / "deck.pptx")
    shape = list(prs.slides[0].shapes)[1]
    assert shape.fill.gradient_angle == 45.0


def test_edit_pptx_shape_fill_color_2_without_fill_color_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="layout: icon-list\n# Icons\n- [A] alpha\n")
    with pytest.raises(ValueError, match="fill_color_2 needs fill_color"):
        tools["edit_pptx_shape"](path="deck.pptx", slide=1, shape_index=1, fill_color_2="38BDF8")


def test_edit_pptx_shape_gradient_angle_without_fill_color_2_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="layout: icon-list\n# Icons\n- [A] alpha\n")
    with pytest.raises(ValueError, match="gradient_angle only applies"):
        tools["edit_pptx_shape"](
            path="deck.pptx", slide=1, shape_index=1, fill_color="0F172A", gradient_angle=45.0
        )


def test_edit_pptx_shape_sets_text_color_on_title(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# **Bold** Title\n- one\n- two")

    tools["edit_pptx_shape"](path="deck.pptx", slide=1, shape_index=0, text_color="2563EB")

    prs = Presentation(tmp_path / "deck.pptx")
    title = prs.slides[0].shapes.title
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            assert str(run.font.color.rgb) == "2563EB"


def test_edit_pptx_shape_text_color_does_not_touch_text_content(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- one\n- two")

    tools["edit_pptx_shape"](path="deck.pptx", slide=1, shape_index=0, text_color="2563EB")

    prs = Presentation(tmp_path / "deck.pptx")
    slide = prs.slides[0]
    assert slide.shapes.title.text == "Title"
    text = tools["read_pptx"](path="deck.pptx", slide=1)
    assert "one" in text
    assert "two" in text


def test_edit_pptx_shape_rejects_text_color_on_a_table(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    content = "# Table\n| A | B |\n| --- | --- |\n| 1 | 2 |\n"
    tools["write_pptx"](path="deck.pptx", content=content)

    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    table_index = next(s["index"] for s in shapes if s["is_table"])
    with pytest.raises(ValueError, match="no text to color"):
        tools["edit_pptx_shape"](
            path="deck.pptx", slide=1, shape_index=table_index, text_color="2563EB"
        )


def test_edit_pptx_shape_rejects_bad_text_color(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    with pytest.raises(ValueError, match="6-hex-digit color"):
        tools["edit_pptx_shape"](path="deck.pptx", slide=1, shape_index=0, text_color="#2563EB")


def test_list_pptx_shapes_describes_gradient_fill(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="layout: icon-list\n# Icons\n- [A] alpha\n")
    tools["edit_pptx_shape"](
        path="deck.pptx", slide=1, shape_index=1, fill_color="0F172A", fill_color_2="38BDF8"
    )

    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    fill = shapes[1]["fill"]
    assert fill == {"type": "gradient", "colors": ["#0F172A", "#38BDF8"], "angle": 90.0}


# --- add_pptx_shape / list_pptx_shape_types (Tier 2: diagram-building shapes) ---


def test_list_pptx_shape_types_returns_sorted_known_names(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    names = tools["list_pptx_shape_types"]()
    assert names == sorted(names)
    assert "RIGHT_ARROW" in names
    assert "FLOWCHART_DECISION" in names
    assert "ROUNDED_RECTANGLE" in names


def test_add_pptx_shape_adds_shape_with_position_size_and_text(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")

    result = tools["add_pptx_shape"](
        path="deck.pptx",
        slide=1,
        shape_type="RIGHT_ARROW",
        left_in=1.0,
        top_in=2.0,
        width_in=3.0,
        height_in=1.5,
        text="Next **step**",
    )

    prs = Presentation(tmp_path / "deck.pptx")
    shapes = list(prs.slides[0].shapes)
    shape = shapes[result["shape_index"]]
    assert shape.auto_shape_type == MSO_SHAPE.RIGHT_ARROW
    assert shape.left == 914400
    assert shape.top == 1828800
    assert shape.width == 2743200
    assert shape.height == 1371600
    assert shape.text_frame.text == "Next step"
    assert shape.text_frame.paragraphs[0].runs[1].font.bold is True


def test_add_pptx_shape_sets_fill_and_line_color(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.dml import MSO_COLOR_TYPE

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")

    result = tools["add_pptx_shape"](
        path="deck.pptx",
        slide=1,
        shape_type="FLOWCHART_DECISION",
        left_in=1.0,
        top_in=1.0,
        width_in=2.0,
        height_in=2.0,
        fill_color="38BDF8",
        line_color="0F172A",
    )

    prs = Presentation(tmp_path / "deck.pptx")
    shape = list(prs.slides[0].shapes)[result["shape_index"]]
    assert shape.fill.fore_color.type == MSO_COLOR_TYPE.RGB
    assert str(shape.fill.fore_color.rgb) == "38BDF8"
    assert str(shape.line.color.rgb) == "0F172A"


def test_add_pptx_shape_returned_index_targets_the_right_shape(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")

    result = tools["add_pptx_shape"](
        path="deck.pptx",
        slide=1,
        shape_type="OVAL",
        left_in=1.0,
        top_in=1.0,
        width_in=1.0,
        height_in=1.0,
    )

    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    assert shapes[result["shape_index"]]["shape_type"] is not None
    assert shapes[result["shape_index"]]["left_in"] == 1.0


def test_add_pptx_shape_rejects_unknown_shape_type(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    with pytest.raises(ValueError, match="Unknown shape_type"):
        tools["add_pptx_shape"](
            path="deck.pptx",
            slide=1,
            shape_type="NOT_A_REAL_SHAPE",
            left_in=1.0,
            top_in=1.0,
            width_in=1.0,
            height_in=1.0,
        )


def test_add_pptx_shape_rejects_bad_fill_color(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    with pytest.raises(ValueError, match="6-hex-digit color"):
        tools["add_pptx_shape"](
            path="deck.pptx",
            slide=1,
            shape_type="OVAL",
            left_in=1.0,
            top_in=1.0,
            width_in=1.0,
            height_in=1.0,
            fill_color="#38BDF8",
        )


# --- add_pptx_formula (native OMML equations, vendored LaTeX->OMML compiler) ---


def test_add_pptx_formula_adds_a_block_equation(tmp_path: Path) -> None:
    from lxml import etree
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")

    result = tools["add_pptx_formula"](
        path="deck.pptx",
        slide=1,
        latex=r"\frac{-b \pm \sqrt{b^2-4ac}}{2a}",
        left_in=1.0,
        top_in=2.0,
        width_in=5.0,
        height_in=1.5,
    )

    prs = Presentation(tmp_path / "deck.pptx")
    shape = list(prs.slides[0].shapes)[result["shape_index"]]
    assert shape.left == 914400
    assert shape.width == 4572000
    p_element = shape.text_frame.paragraphs[0]._p  # noqa: SLF001
    xml = etree.tostring(p_element).decode()
    assert "{http://schemas.microsoft.com/office/drawing/2010/main}m" in [
        child.tag for child in p_element
    ]
    assert "oMathPara" in xml  # display=True -> block equation, centered
    assert "√" not in xml  # real OMML radical markup, not a unicode glyph
    assert result["display"] is True


def test_add_pptx_formula_display_false_is_inline_not_wrapped_in_omathpara(
    tmp_path: Path,
) -> None:
    from lxml import etree
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")

    result = tools["add_pptx_formula"](
        path="deck.pptx",
        slide=1,
        latex=r"x_i^2 + y_i^2 = r^2",
        left_in=1.0,
        top_in=4.0,
        width_in=3.0,
        height_in=0.6,
        display=False,
    )

    prs = Presentation(tmp_path / "deck.pptx")
    shape = list(prs.slides[0].shapes)[result["shape_index"]]
    p_element = shape.text_frame.paragraphs[0]._p  # noqa: SLF001
    xml = etree.tostring(p_element).decode()
    assert "oMathPara" not in xml  # inline -> bare <m:oMath>, no paragraph wrapper
    assert "oMath" in xml
    assert result["display"] is False


def test_add_pptx_formula_marks_mc_ignorable_a14_on_the_slide(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    tools["add_pptx_formula"](
        path="deck.pptx",
        slide=1,
        latex=r"E = mc^2",
        left_in=1.0,
        top_in=1.0,
        width_in=3.0,
        height_in=1.0,
    )

    prs = Presentation(tmp_path / "deck.pptx")
    slide_element = prs.slides[0].element
    mc_ignorable = slide_element.get(
        "{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"
    )
    assert mc_ignorable is not None
    assert "a14" in mc_ignorable.split()


def test_add_pptx_formula_rejects_malformed_latex(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")

    with pytest.raises(ValueError, match="Could not compile LaTeX formula"):
        tools["add_pptx_formula"](
            path="deck.pptx",
            slide=1,
            latex=r"\frac{1}{",  # unbalanced braces
            left_in=1.0,
            top_in=1.0,
            width_in=3.0,
            height_in=1.0,
        )


@pytest.mark.real_libreoffice
@pytest.mark.skipif(
    not _libreoffice_actually_works(),
    reason="LibreOffice not installed or not functional in this environment",
)
def test_add_pptx_formula_renders_through_libreoffice_without_qa_skip(tmp_path: Path) -> None:
    """Not a rendering-correctness assertion (that needs a human eye on a
    PNG, done manually during implementation -- see PPTX_DESIGN.md) --
    just confirms the file this tool produces isn't one LibreOffice itself
    refuses to convert (the class of defect a broken OOXML write causes)."""
    state_dir = tmp_path / "state"
    tools = {
        tool.__name__: tool
        for tool in build_presentation_tools(tmp_path / "workspace", state_dir=state_dir)  # type: ignore[attr-defined]
    }
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    tools["add_pptx_formula"](
        path="deck.pptx",
        slide=1,
        latex=r"\sum_{i=1}^{n} i = \frac{n(n+1)}{2}",
        left_in=1.0,
        top_in=1.0,
        width_in=5.0,
        height_in=1.0,
    )

    result = tools["render_pptx_preview"](path="deck.pptx")
    assert result["preview_skipped_reason"] is None
    assert len(result["preview_paths"]) == 1


def test_add_pptx_formula_is_write_local_and_gated(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["add_pptx_formula"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


# --- crop_pptx_image (Tier 2: general picture cropping) ---


def test_crop_pptx_image_sets_crop_fractions(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    _make_test_png(tmp_path / "red.png")
    tools["add_pptx_image"](
        path="deck.pptx", slide=1, image_path="red.png", left=1.0, top=1.0, width=2.0, height=2.0
    )
    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    pic_index = next(s["index"] for s in shapes if s["is_picture"])

    result = tools["crop_pptx_image"](
        path="deck.pptx",
        slide=1,
        shape_index=pic_index,
        crop_left=0.1,
        crop_right=0.2,
        crop_top=0.05,
        crop_bottom=0.15,
    )
    assert result["crop_left"] == 0.1

    prs = Presentation(tmp_path / "deck.pptx")
    shape = list(prs.slides[0].shapes)[pic_index]
    assert shape.crop_left == pytest.approx(0.1)
    assert shape.crop_right == pytest.approx(0.2)
    assert shape.crop_top == pytest.approx(0.05)
    assert shape.crop_bottom == pytest.approx(0.15)


def test_crop_pptx_image_leaves_position_and_size_untouched(tmp_path: Path) -> None:

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    _make_test_png(tmp_path / "red.png")
    tools["add_pptx_image"](
        path="deck.pptx", slide=1, image_path="red.png", left=1.0, top=1.0, width=2.0, height=2.0
    )
    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    pic_index = next(s["index"] for s in shapes if s["is_picture"])
    before = shapes[pic_index]

    tools["crop_pptx_image"](path="deck.pptx", slide=1, shape_index=pic_index, crop_left=0.3)

    after = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"][pic_index]
    assert after["left_in"] == before["left_in"]
    assert after["top_in"] == before["top_in"]
    assert after["width_in"] == before["width_in"]
    assert after["height_in"] == before["height_in"]


def test_crop_pptx_image_rejects_out_of_range_fraction(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    _make_test_png(tmp_path / "red.png")
    tools["add_pptx_image"](path="deck.pptx", slide=1, image_path="red.png")
    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    pic_index = next(s["index"] for s in shapes if s["is_picture"])

    with pytest.raises(ValueError, match=r"\[0\.0, 1\.0\)"):
        tools["crop_pptx_image"](path="deck.pptx", slide=1, shape_index=pic_index, crop_left=1.0)


def test_crop_pptx_image_rejects_left_plus_right_over_one(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    _make_test_png(tmp_path / "red.png")
    tools["add_pptx_image"](path="deck.pptx", slide=1, image_path="red.png")
    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    pic_index = next(s["index"] for s in shapes if s["is_picture"])

    with pytest.raises(ValueError, match="crop_left"):
        tools["crop_pptx_image"](
            path="deck.pptx", slide=1, shape_index=pic_index, crop_left=0.6, crop_right=0.6
        )


def test_crop_pptx_image_rejects_non_picture_shape(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet")
    with pytest.raises(ValueError, match="isn't a picture"):
        tools["crop_pptx_image"](path="deck.pptx", slide=1, shape_index=0, crop_left=0.1)


# --- edit_pptx_table_cell / merge_pptx_table_cells (Tier 1 #2: table editing) ---


def _table_deck(tmp_path: Path, tools: dict[str, object]) -> int:
    """Writes a 3-row x 3-col table deck (header + 2 data rows), returns
    the table's shape_index."""
    content = "# Table\n| A | B | C |\n| --- | --- | --- |\n| 1 | 2 | 3 |\n| 4 | 5 | 6 |\n"
    tools["write_pptx"](path="deck.pptx", content=content)
    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    return next(s["index"] for s in shapes if s["is_table"])


def test_edit_pptx_table_cell_replaces_one_cell_leaves_others_alone(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    table_index = _table_deck(tmp_path, tools)

    result = tools["edit_pptx_table_cell"](
        path="deck.pptx", slide=1, shape_index=table_index, row=1, col=1, text="CHANGED"
    )
    assert result["row"] == 1
    assert result["col"] == 1

    prs = Presentation(str(tmp_path / "deck.pptx"))
    table = list(prs.slides[0].shapes)[table_index].table
    assert table.cell(1, 1).text == "CHANGED"
    # Neighboring cells untouched.
    assert table.cell(1, 0).text == "1"
    assert table.cell(0, 1).text == "B"


def test_edit_pptx_table_cell_replaces_not_appends(tmp_path: Path) -> None:
    """Same real-world append-vs-replace bug class edit_pptx_text was
    fixed for -- writing to a cell that already has text must replace
    it, not garble it onto the end."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    table_index = _table_deck(tmp_path, tools)
    tools["edit_pptx_table_cell"](
        path="deck.pptx", slide=1, shape_index=table_index, row=0, col=0, text="NEW"
    )

    prs = Presentation(str(tmp_path / "deck.pptx"))
    table = list(prs.slides[0].shapes)[table_index].table
    assert table.cell(0, 0).text == "NEW"


def test_edit_pptx_table_cell_rejects_out_of_range_cell(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    table_index = _table_deck(tmp_path, tools)
    with pytest.raises(ValueError, match="out of range"):
        tools["edit_pptx_table_cell"](
            path="deck.pptx", slide=1, shape_index=table_index, row=9, col=0, text="x"
        )


def test_edit_pptx_table_cell_rejects_non_table_shape(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    table_index = _table_deck(tmp_path, tools)
    title_index = next(
        s["index"]
        for s in tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
        if not s["is_table"]
    )
    assert title_index != table_index
    with pytest.raises(ValueError, match="isn't a table"):
        tools["edit_pptx_table_cell"](
            path="deck.pptx", slide=1, shape_index=title_index, row=0, col=0, text="x"
        )


def test_merge_pptx_table_cells_merges_a_row_span(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    table_index = _table_deck(tmp_path, tools)

    result = tools["merge_pptx_table_cells"](
        path="deck.pptx",
        slide=1,
        shape_index=table_index,
        start_row=0,
        start_col=0,
        end_row=0,
        end_col=2,
    )
    assert result["merged_range"] == {
        "start_row": 0,
        "start_col": 0,
        "end_row": 0,
        "end_col": 2,
    }

    prs = Presentation(str(tmp_path / "deck.pptx"))
    table = list(prs.slides[0].shapes)[table_index].table
    origin = table.cell(0, 0)
    assert origin.is_merge_origin is True
    assert origin.span_width == 3
    assert table.cell(0, 1).is_spanned is True
    assert table.cell(0, 2).is_spanned is True


def test_merge_pptx_table_cells_accepts_either_diagonal_order(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    table_index = _table_deck(tmp_path, tools)

    tools["merge_pptx_table_cells"](
        path="deck.pptx",
        slide=1,
        shape_index=table_index,
        start_row=1,
        start_col=2,
        end_row=0,
        end_col=1,
    )

    prs = Presentation(str(tmp_path / "deck.pptx"))
    table = list(prs.slides[0].shapes)[table_index].table
    assert table.cell(0, 1).span_width == 2
    assert table.cell(0, 1).span_height == 2


def test_merge_pptx_table_cells_rejects_out_of_range_cell(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    table_index = _table_deck(tmp_path, tools)
    with pytest.raises(ValueError, match="out of range"):
        tools["merge_pptx_table_cells"](
            path="deck.pptx",
            slide=1,
            shape_index=table_index,
            start_row=0,
            start_col=0,
            end_row=9,
            end_col=0,
        )


def test_merge_pptx_table_cells_rejects_range_already_containing_a_merge(
    tmp_path: Path,
) -> None:
    tools = _tools_by_name(tmp_path)
    table_index = _table_deck(tmp_path, tools)
    tools["merge_pptx_table_cells"](
        path="deck.pptx",
        slide=1,
        shape_index=table_index,
        start_row=0,
        start_col=0,
        end_row=0,
        end_col=1,
    )
    with pytest.raises(ValueError, match="merged cell"):
        tools["merge_pptx_table_cells"](
            path="deck.pptx",
            slide=1,
            shape_index=table_index,
            start_row=0,
            start_col=0,
            end_row=0,
            end_col=2,
        )


def test_merge_pptx_table_cells_rejects_non_table_shape(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _table_deck(tmp_path, tools)
    title_index = next(
        s["index"]
        for s in tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
        if not s["is_table"]
    )
    with pytest.raises(ValueError, match="isn't a table"):
        tools["merge_pptx_table_cells"](
            path="deck.pptx",
            slide=1,
            shape_index=title_index,
            start_row=0,
            start_col=0,
            end_row=0,
            end_col=1,
        )


# --- replace_pptx_image (Tier 1 #3: image replacement, preserves placement) ---


def test_replace_pptx_image_keeps_position_size_swaps_pixels(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    _make_test_png(tmp_path / "red.png")
    tools["add_pptx_image"](
        path="deck.pptx", slide=1, image_path="red.png", left=1.0, top=1.0, width=2.0, height=2.0
    )

    from PIL import Image

    Image.new("RGB", (50, 50), color="blue").save(tmp_path / "blue.png")

    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    pic_index = next(s["index"] for s in shapes if s["is_picture"])

    result = tools["replace_pptx_image"](
        path="deck.pptx", slide=1, shape_index=pic_index, image_path="blue.png"
    )
    assert result["shape_index"] == pic_index

    prs = Presentation(str(tmp_path / "deck.pptx"))
    shape = list(prs.slides[0].shapes)[pic_index]
    from pptx.util import Inches

    assert shape.left == Inches(1.0)
    assert shape.top == Inches(1.0)
    assert shape.width == Inches(2.0)
    assert shape.height == Inches(2.0)
    assert shape.image.size == (50, 50)


def test_replace_pptx_image_rejects_out_of_range_shape_index(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    _make_test_png(tmp_path / "red.png")
    with pytest.raises(ValueError, match="out of range"):
        tools["replace_pptx_image"](
            path="deck.pptx", slide=1, shape_index=99, image_path="red.png"
        )


def test_replace_pptx_image_rejects_non_picture_shape(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    _make_test_png(tmp_path / "red.png")
    with pytest.raises(ValueError, match="isn't a picture"):
        tools["replace_pptx_image"](
            path="deck.pptx", slide=1, shape_index=0, image_path="red.png"
        )


def test_replace_pptx_image_on_missing_file_raises(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    _make_test_png(tmp_path / "red.png")
    tools["add_pptx_image"](path="deck.pptx", slide=1, image_path="red.png")
    shapes = tools["list_pptx_shapes"](path="deck.pptx", slide=1)["shapes"]
    pic_index = next(s["index"] for s in shapes if s["is_picture"])

    with pytest.raises(ValueError, match="does not exist"):
        tools["replace_pptx_image"](
            path="deck.pptx", slide=1, shape_index=pic_index, image_path="missing.png"
        )


# --- list_pptx_icons / add_pptx_icon (Tier 1 #3: bundled Lucide icons) ---


def test_list_pptx_icons_returns_a_sorted_nonempty_list_of_known_names(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    names = tools["list_pptx_icons"]()
    assert names == sorted(names)
    assert len(names) > 30
    assert "check" in names
    assert "trending-up" in names


def test_add_pptx_icon_inserts_a_square_picture_at_the_given_position(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")

    result = tools["add_pptx_icon"](
        path="deck.pptx",
        slide=1,
        icon_name="check",
        left_in=2.0,
        top_in=3.0,
        size_in=1.25,
    )
    assert result["icon_name"] == "check"

    prs = Presentation(str(tmp_path / "deck.pptx"))
    shapes = list(prs.slides[0].shapes)
    icon = shapes[-1]
    assert icon.shape_type == 13  # PICTURE
    assert icon.left == Inches(2.0)
    assert icon.top == Inches(3.0)
    assert icon.width == Inches(1.25)
    assert icon.height == Inches(1.25)


def test_add_pptx_icon_recolors_to_the_requested_color(tmp_path: Path) -> None:
    from io import BytesIO

    from PIL import Image
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    tools["add_pptx_icon"](
        path="deck.pptx", slide=1, icon_name="check", color="38BDF8"
    )

    prs = Presentation(str(tmp_path / "deck.pptx"))
    icon = list(prs.slides[0].shapes)[-1]
    img = Image.open(BytesIO(icon.image.blob)).convert("RGBA")
    bbox = img.getbbox()
    assert bbox is not None
    # Find one solidly-opaque pixel (avoid anti-aliased edges) and check
    # its RGB matches the requested color exactly.
    found = None
    for y in range(bbox[1], bbox[3]):
        for x in range(bbox[0], bbox[2]):
            pixel = img.getpixel((x, y))
            if pixel[3] == 255:
                found = pixel
                break
        if found:
            break
    assert found is not None
    assert found[:3] == (0x38, 0xBD, 0xF8)


def test_add_pptx_icon_default_color_is_applied_when_omitted(tmp_path: Path) -> None:
    from io import BytesIO

    from PIL import Image
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    tools["add_pptx_icon"](path="deck.pptx", slide=1, icon_name="check")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    icon = list(prs.slides[0].shapes)[-1]
    img = Image.open(BytesIO(icon.image.blob)).convert("RGBA")
    bbox = img.getbbox()
    assert bbox is not None


def test_add_pptx_icon_rejects_unknown_icon_name(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    with pytest.raises(ValueError, match="Unknown icon"):
        tools["add_pptx_icon"](path="deck.pptx", slide=1, icon_name="not-a-real-icon")


def test_add_pptx_icon_rejects_bad_color(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    with pytest.raises(ValueError, match="6-hex-digit color"):
        tools["add_pptx_icon"](path="deck.pptx", slide=1, icon_name="check", color="#38BDF8")


def test_add_pptx_icon_rejects_out_of_range_slide(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    with pytest.raises(ValueError, match="out of range"):
        tools["add_pptx_icon"](path="deck.pptx", slide=2, icon_name="check")


def _solid_opaque_pixel(img: Any) -> tuple[int, int, int, int]:
    """First solidly-opaque pixel (avoid anti-aliased edges), matching
    the same avoid-edge-pixels pattern add_pptx_icon's own recolor test
    already uses."""
    bbox = img.getbbox()
    assert bbox is not None
    for y in range(bbox[1], bbox[3]):
        for x in range(bbox[0], bbox[2]):
            pixel = img.getpixel((x, y))
            if pixel[3] == 255:
                return pixel
    raise AssertionError("no solidly-opaque pixel found")


def test_recolor_pptx_icon_changes_pixel_color_in_place(tmp_path: Path) -> None:
    from io import BytesIO

    from PIL import Image
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    tools["add_pptx_icon"](path="deck.pptx", slide=1, icon_name="check", color="1F2937")

    result = tools["recolor_pptx_icon"](path="deck.pptx", slide=1, shape_index=2, color="c2410c")
    assert result == {
        "path": "deck.pptx",
        "slide": 1,
        "shape_index": 2,
        "color": "C2410C",
        "preview_path": None,
        "preview_skipped_reason": "no state_dir configured",
    }

    prs = Presentation(str(tmp_path / "deck.pptx"))
    icon = list(prs.slides[0].shapes)[2]
    img = Image.open(BytesIO(icon.image.blob)).convert("RGBA")
    assert _solid_opaque_pixel(img)[:3] == (0xC2, 0x41, 0x0C)


def test_recolor_pptx_icon_position_and_size_are_untouched(tmp_path: Path) -> None:
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    tools["add_pptx_icon"](
        path="deck.pptx", slide=1, icon_name="check", left_in=2.0, top_in=3.0, size_in=1.5
    )
    before = Presentation(str(tmp_path / "deck.pptx"))
    icon_before = list(before.slides[0].shapes)[2]
    left, top = icon_before.left, icon_before.top
    width, height = icon_before.width, icon_before.height

    tools["recolor_pptx_icon"](path="deck.pptx", slide=1, shape_index=2, color="C2410C")

    after = Presentation(str(tmp_path / "deck.pptx"))
    icon_after = list(after.slides[0].shapes)[2]
    assert (icon_after.left, icon_after.top, icon_after.width, icon_after.height) == (
        left,
        top,
        width,
        height,
    )


def test_recolor_pptx_icon_works_on_an_already_recolored_icon(tmp_path: Path) -> None:
    from io import BytesIO

    from PIL import Image
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    tools["add_pptx_icon"](path="deck.pptx", slide=1, icon_name="check", color="1F2937")
    tools["recolor_pptx_icon"](path="deck.pptx", slide=1, shape_index=2, color="C2410C")
    tools["recolor_pptx_icon"](path="deck.pptx", slide=1, shape_index=2, color="0F6B5C")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    icon = list(prs.slides[0].shapes)[2]
    img = Image.open(BytesIO(icon.image.blob)).convert("RGBA")
    assert _solid_opaque_pixel(img)[:3] == (0x0F, 0x6B, 0x5C)


def test_recolor_pptx_icon_survives_round_trip(tmp_path: Path) -> None:
    import warnings

    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    tools["add_pptx_icon"](path="deck.pptx", slide=1, icon_name="check")
    tools["recolor_pptx_icon"](path="deck.pptx", slide=1, shape_index=2, color="C2410C")

    with warnings.catch_warnings():
        warnings.simplefilter("error")
        Presentation(str(tmp_path / "deck.pptx"))


def test_recolor_pptx_icon_rejects_non_picture_shape(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")

    with pytest.raises(ValueError, match="isn't a picture"):
        tools["recolor_pptx_icon"](path="deck.pptx", slide=1, shape_index=1, color="C2410C")


def test_recolor_pptx_icon_rejects_bad_color(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    tools["add_pptx_icon"](path="deck.pptx", slide=1, icon_name="check")

    with pytest.raises(ValueError, match="6-hex-digit color"):
        tools["recolor_pptx_icon"](path="deck.pptx", slide=1, shape_index=2, color="#C2410C")


def test_recolor_pptx_icon_rejects_out_of_range_shape_index(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")

    with pytest.raises(ValueError, match="out of range"):
        tools["recolor_pptx_icon"](path="deck.pptx", slide=1, shape_index=99, color="C2410C")


def test_recolor_pptx_icon_is_medium_risk_and_requires_approval(tmp_path: Path) -> None:
    from coscribe.runtime.types import get_tool_metadata

    tools = _tools_by_name(tmp_path)
    metadata = get_tool_metadata(tools["recolor_pptx_icon"])
    assert metadata.risk_category == "WRITE_LOCAL"
    assert metadata.requires_approval is True


# --- delete/duplicate/reorder_pptx_slide (Tier 2 #4: deck structure editing) ---


def _titles(pptx_path: Path) -> list[str | None]:
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    return [s.shapes.title.text if s.shapes.title is not None else None for s in prs.slides]


def _three_slide_deck(tmp_path: Path, tools: dict[str, object]) -> None:
    content = "# Slide A\n- one\n---\n# Slide B\n- two\n---\n# Slide C\n- three\n"
    tools["write_pptx"](path="deck.pptx", content=content)


def test_delete_pptx_slide_removes_only_that_slide(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _three_slide_deck(tmp_path, tools)

    result = tools["delete_pptx_slide"](path="deck.pptx", slide=2)
    assert result["deleted_slide"] == 2
    assert result["slide_count"] == 2
    assert _titles(tmp_path / "deck.pptx") == ["Slide A", "Slide C"]


def test_delete_pptx_slide_rejects_out_of_range(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _three_slide_deck(tmp_path, tools)
    with pytest.raises(ValueError, match="out of range"):
        tools["delete_pptx_slide"](path="deck.pptx", slide=99)


def test_duplicate_pptx_slide_defaults_to_right_after_source(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _three_slide_deck(tmp_path, tools)

    result = tools["duplicate_pptx_slide"](path="deck.pptx", slide=1)
    assert result["new_slide"] == 2
    assert result["slide_count"] == 4
    assert _titles(tmp_path / "deck.pptx") == ["Slide A", "Slide A", "Slide B", "Slide C"]


def test_duplicate_pptx_slide_honors_explicit_insert_at(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _three_slide_deck(tmp_path, tools)

    tools["duplicate_pptx_slide"](path="deck.pptx", slide=1, insert_at=3)
    assert _titles(tmp_path / "deck.pptx") == ["Slide A", "Slide B", "Slide A", "Slide C"]


def test_duplicate_pptx_slide_can_insert_at_the_very_end(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _three_slide_deck(tmp_path, tools)

    tools["duplicate_pptx_slide"](path="deck.pptx", slide=1, insert_at=4)
    assert _titles(tmp_path / "deck.pptx") == ["Slide A", "Slide B", "Slide C", "Slide A"]


def test_duplicate_pptx_slide_rejects_out_of_range_insert_at(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _three_slide_deck(tmp_path, tools)
    with pytest.raises(ValueError, match="out of range"):
        tools["duplicate_pptx_slide"](path="deck.pptx", slide=1, insert_at=99)


def test_duplicate_pptx_slide_copies_shape_content_leaves_original_untouched(
    tmp_path: Path,
) -> None:
    """The copy must have the source's own body text (not just its
    title), and editing the copy afterward must not affect the
    original -- confirms it's a real independent XML part, not a shared
    reference."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- alpha\n- beta\n")

    tools["duplicate_pptx_slide"](path="deck.pptx", slide=1)
    tools["edit_pptx_text"](path="deck.pptx", slide=2, content="- CHANGED")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    original_body = next(
        p for p in prs.slides[0].placeholders if p.placeholder_format.idx != 0
    )
    copy_body = next(p for p in prs.slides[1].placeholders if p.placeholder_format.idx != 0)
    assert "alpha" in original_body.text_frame.text
    assert "CHANGED" in copy_body.text_frame.text
    assert "alpha" not in copy_body.text_frame.text


def test_duplicate_pptx_slide_preserves_picture_and_hyperlink(tmp_path: Path) -> None:
    """Regression test for the real risk this tool's whole design is
    built around: relationship ids on the duplicated slide's own part
    are not guaranteed to match the source's, so every r:id-bearing
    attribute in the copied XML must be correctly remapped -- not just
    the common case where ids happen to line up. A picture (r:embed)
    and a hyperlink (r:id) are two different attribute names in two
    different elements, exercising the blanket remap generically rather
    than one specific tag."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- click me\n")
    _make_test_png(tmp_path / "logo.png")
    tools["add_pptx_image"](path="deck.pptx", slide=1, image_path="logo.png")

    prs = Presentation(str(tmp_path / "deck.pptx"))
    body = next(p for p in prs.slides[0].placeholders if p.placeholder_format.idx != 0)
    body.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.com"
    prs.save(str(tmp_path / "deck.pptx"))

    tools["duplicate_pptx_slide"](path="deck.pptx", slide=1)

    prs2 = Presentation(str(tmp_path / "deck.pptx"))
    dup_slide = prs2.slides[1]
    pictures = [s for s in dup_slide.shapes if s.shape_type == 13]
    assert len(pictures) == 1
    assert pictures[0].image.size == (10, 10)
    dup_body = next(p for p in dup_slide.placeholders if p.placeholder_format.idx != 0)
    assert dup_body.text_frame.paragraphs[0].runs[0].hyperlink.address == "https://example.com"


def test_duplicate_pptx_slide_does_not_carry_over_speaker_notes(tmp_path: Path) -> None:
    """Documented v1 scope cut (see _duplicate_slide_part's docstring),
    not an accident -- pinned down as a test so a future change to this
    behavior is a deliberate decision, not a silent regression either
    way."""
    from pptx import Presentation

    tools = _tools_by_name(tmp_path)
    tools["write_pptx"](path="deck.pptx", content="# Title\n- bullet\n")
    tools["set_pptx_notes"](path="deck.pptx", slide=1, notes="speaker notes here")

    tools["duplicate_pptx_slide"](path="deck.pptx", slide=1)

    prs = Presentation(str(tmp_path / "deck.pptx"))
    dup_slide = prs.slides[1]
    assert dup_slide.has_notes_slide is False


def test_reorder_pptx_slide_moves_to_new_position(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _three_slide_deck(tmp_path, tools)

    result = tools["reorder_pptx_slide"](path="deck.pptx", slide=3, new_position=1)
    assert result == {
        "path": "deck.pptx",
        "old_position": 3,
        "new_position": 1,
    }
    assert _titles(tmp_path / "deck.pptx") == ["Slide C", "Slide A", "Slide B"]


def test_reorder_pptx_slide_to_its_own_position_is_a_no_op(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _three_slide_deck(tmp_path, tools)

    tools["reorder_pptx_slide"](path="deck.pptx", slide=2, new_position=2)
    assert _titles(tmp_path / "deck.pptx") == ["Slide A", "Slide B", "Slide C"]


def test_reorder_pptx_slide_rejects_out_of_range_new_position(tmp_path: Path) -> None:
    tools = _tools_by_name(tmp_path)
    _three_slide_deck(tmp_path, tools)
    with pytest.raises(ValueError, match="out of range"):
        tools["reorder_pptx_slide"](path="deck.pptx", slide=1, new_position=99)
