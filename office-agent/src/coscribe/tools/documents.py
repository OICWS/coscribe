"""Word (.docx) and PDF read/write tools, scoped to a single workspace root.

Both writers accept the same lightweight markdown subset (headings via
``#``/``##``/``###``, ``-``/``*`` bullets, ``1.`` numbered items, and
``| a | b |`` pipe tables) so the model can produce either format from one
mental model. Editing an existing document is read -> modify the text ->
write(overwrite=True), the same whole-file-overwrite contract
``write_file`` already uses.

``read_docx`` converts via ``mammoth`` (docx -> html) + ``markdownify``
(html -> markdown) rather than walking the document tree by hand -- the
same pair of well-established libraries microsoft/markitdown itself uses
for docx, without markitdown's own mandatory ``magika`` dependency (which
drags in onnxruntime + numpy for file-type sniffing this codebase doesn't
need, since each tool already knows its own format). Its output is real
markdown -- a superset of the writers' subset (arbitrary heading depth,
links, bold/italic) -- rather than an exact mirror of what the writers
accept, so round-tripping a generated file through read_docx won't be
byte-for-byte identical, just semantically equivalent. ``read_pdf``/
``search_pdf`` extract text with ``pypdfium2`` (Chrome's PDFium) -- see
``_pdf_page_texts`` for why not ``pdfplumber``.

``mammoth``/``pypdfium2``/``docx``/``markdownify``/``reportlab`` are all
imported lazily, inside each method that actually needs them, rather than
at module level -- a real, measured startup-time cost (see
runtime_lg/README.md's startup-time section): every one of these tools is
registered into the Coordinator's tool list at agent-build time
regardless of whether it's ever called this session (the model needs
every tool's name/docstring/schema up front), but none of their
*implementations* are needed until a specific tool is actually invoked.
None of the tool-facing function signatures below reference any of these
libraries' own types (only `str`/`int`/`dict`/`list`), so LangChain's
tool-schema introspection never forces the lazy import open -- confirmed
by this file's own tests still passing with these deferred.

``write_docx``'s optional ``template_path`` loads an existing .docx as the
starting point instead of a blank ``Document()``, then clears its body --
keeping only ``sectPr`` (page size/margins), removing every other child --
before writing the new content. Verified empirically (round-tripped
through python-docx) that removing anything else, or removing ``sectPr``
too, breaks the document; this is the one thing in this file that isn't a
pure library call.

``write_docx`` validates its finished in-memory document tree against the
real ECMA-376 ``wml.xsd`` schema (``_ooxml_validate.assert_wml_valid``)
immediately before ``document.save()`` -- the same "never serialize an
unvalidated write" discipline ``presentations.py`` already established
for its own hand-built OOXML (tracked-changes markup, TOC field,
comment-range anchors, and template body-clearing are all hand-mutated
XML python-docx has no schema-checked API for). A failure raises before
the file is touched, rather than producing a `.docx` that happens to open
in Word/LibreOffice's own forgiving parsers today but is not actually a
conformant document.

``insert_docx_text``/``delete_docx_text``/
``replace_docx_text``/``accept_docx_tracked_changes``/
``reject_docx_tracked_changes`` edit an *existing* docx in place with real
``<w:ins>``/``<w:del>`` tracked-changes markup, using
``_native_docx_tracks.tracks.TracksMixin`` (vendored from
`SecurityRonin/docx-mcp`, MIT -- see that package's NOTICE.md) for the
hard part: cascading context-anchored fuzzy text matching, multi-run-
spanning delete/insert with ``rPr`` inheritance, and word-level diff-
minimised replace. Unlike that project's own tool surface (which
addresses a paragraph by its ``w14:paraId``, discovered via a separate
call), these tools locate the target paragraph themselves from the text
being edited -- ``_locate_paragraph`` tries every paragraph and requires
exactly one unambiguous match -- so the caller never needs to know or
pass a paragraph id. ``_ensure_para_ids`` assigns a ``w14:paraId`` (a
Word 2010+ extension the vendored mixin's own paragraph lookup is keyed
on) to every paragraph that doesn't already have one -- real Word always
sets this, ``python-docx`` never does -- and marks the namespace
``mc:Ignorable`` on the document root the same way real Word-authored
files do, so the new attribute validates against ``wml.xsd`` the same
way any other Word 2010+ extension already does (see
``_ooxml_validate.py``'s ``_strip_mce_ignorable``).
"""

from __future__ import annotations

import random
import re
import threading
import zipfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape as _xml_escape

from ..runtime.types import tool_metadata
from ._native_docx_tracks import _constants as _tracks_constants
from ._native_docx_tracks.tracks import TracksMixin, _resolve
from ._ooxml_validate import assert_wml_valid
from ._thumbnail import render_thumbnail
from ._workspace import WorkspaceScope
from .files import DEFAULT_IGNORES

_MC_IGNORABLE = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"

_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_NUMBER_RE = re.compile(r"^\d+\.\s+(.*)$")
_SEPARATOR_CELL_RE = re.compile(r"^:?-{2,}:?$")
_TOC_LINE_RE = re.compile(r"^\[TOC\]$", re.IGNORECASE)
_COMMENT_MARKER_RE = re.compile(r"\s*\{\{comment:\s*(.+?)\}\}\s*$", re.IGNORECASE)

_PAGE_SIZES_EMU: dict[str, tuple[int, int]] = {
    # (width, height) in EMU, portrait orientation, at python-docx's own
    # 914400-EMU-per-inch convention -- Inches()/Mm() both just multiply
    # into this unit, using raw EMU here avoids importing docx.shared at
    # module level (kept a lazy, function-local import like every other
    # `docx` usage in this file, see the module docstring).
    "letter": (7772400, 10058400),  # 8.5in x 11in
    "legal": (7772400, 12801600),  # 8.5in x 14in
    "a4": (7560310, 10692130),  # 210mm x 297mm
    "a3": (10692130, 15120620),  # 297mm x 420mm
}


def _extract_comment(text: str) -> tuple[str, str | None]:
    """Split a trailing ``{{comment: ...}}`` marker off a block's text.

    Returns ``(visible_text, comment_text_or_None)``. Only recognized on
    headings/paragraphs/bullets/numbered items -- table cells are taken
    literally, so a marker inside one is not a comment (documented v1
    scope limit, see ``write_docx``'s docstring)."""
    match = _COMMENT_MARKER_RE.search(text)
    if not match:
        return text, None
    return text[: match.start()].rstrip(), match.group(1).strip()


@dataclass
class Block:
    kind: str  # "heading" | "paragraph" | "bullet" | "number" | "table" | "toc"
    text: str = ""
    level: int = 1
    ordinal: int = 1
    rows: list[list[str]] = field(default_factory=list)
    comment: str | None = None


_PDFIUM_LOCK = threading.Lock()


def _pdf_page_texts(
    file_path: Path, start_page: int = 1, end_page: int = 0
) -> list[tuple[int, str]]:
    """(1-based page number, text) pairs, via pdfium rather than
    pdfplumber: ~100x faster (650 pages in ~1.6s vs ~130s measured), and
    pdfplumber's pure-Python parsing holds the GIL for that whole time,
    stalling the web server's event loop along with it. PDFium itself is
    not thread-safe, hence the lock -- tool calls can run concurrently."""
    import pypdfium2 as pdfium

    pages: list[tuple[int, str]] = []
    with _PDFIUM_LOCK:
        document = pdfium.PdfDocument(str(file_path))
        try:
            last = len(document) if end_page <= 0 else min(end_page, len(document))
            for index in range(max(start_page, 1), last + 1):
                page = document[index - 1]
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_range()
                finally:
                    textpage.close()
                    page.close()
                pages.append((index, text.replace("\r\n", "\n")))
        finally:
            document.close()
    return pages


def _split_table_row(line: str) -> list[str] | None:
    if not (line.startswith("|") and line.endswith("|") and len(line) >= 2):
        return None
    return [cell.strip() for cell in line[1:-1].split("|")]


def _is_separator_row(cells: list[str]) -> bool:
    return all(_SEPARATOR_CELL_RE.match(cell) for cell in cells)


def parse_blocks(content: str) -> list[Block]:
    """Parse the shared markdown subset into structured blocks."""
    blocks: list[Block] = []
    paragraph_lines: list[str] = []
    table_rows: list[list[str]] = []
    number_counter = 0

    def flush_paragraph() -> None:
        text = " ".join(paragraph_lines).strip()
        if text:
            text, comment = _extract_comment(text)
            blocks.append(Block(kind="paragraph", text=text, comment=comment))
        paragraph_lines.clear()

    def flush_table() -> None:
        if table_rows:
            blocks.append(Block(kind="table", rows=[list(row) for row in table_rows]))
            table_rows.clear()

    for raw_line in content.splitlines():
        line = raw_line.strip()

        cells = _split_table_row(line)
        if cells is not None:
            flush_paragraph()
            if not _is_separator_row(cells):
                table_rows.append(cells)
            continue
        flush_table()

        if not line:
            flush_paragraph()
            number_counter = 0
            continue

        if _TOC_LINE_RE.match(line):
            flush_paragraph()
            number_counter = 0
            blocks.append(Block(kind="toc"))
            continue

        heading_match = _HEADING_RE.match(line)
        if heading_match:
            flush_paragraph()
            number_counter = 0
            level = min(len(heading_match.group(1)), 3)
            text, comment = _extract_comment(heading_match.group(2).strip())
            blocks.append(Block(kind="heading", level=level, text=text, comment=comment))
            continue

        bullet_match = _BULLET_RE.match(line)
        if bullet_match:
            flush_paragraph()
            number_counter = 0
            text, comment = _extract_comment(bullet_match.group(1).strip())
            blocks.append(Block(kind="bullet", text=text, comment=comment))
            continue

        number_match = _NUMBER_RE.match(line)
        if number_match:
            flush_paragraph()
            number_counter += 1
            text, comment = _extract_comment(number_match.group(1).strip())
            blocks.append(
                Block(kind="number", text=text, ordinal=number_counter, comment=comment)
            )
            continue

        number_counter = 0
        paragraph_lines.append(line)

    flush_paragraph()
    flush_table()
    return blocks


_INLINE_MARKUP_RE = re.compile(r"\*\*(?P<bold>.+?)\*\*|\*(?P<italic1>.+?)\*|_(?P<italic2>.+?)_")


def parse_inline_runs(text: str) -> list[tuple[str, bool, bool]]:
    """Split ``**bold**``/``*italic*``/``_italic_`` markers out of a block's
    text into ``(text, bold, italic)`` runs, so a writer can style each
    fragment with its own run instead of emitting the literal asterisks/
    underscores as visible text -- the model reaches for this markdown
    convention constantly even though ``parse_blocks`` above never taught it
    to the writers, so it used to come through as literal ``**bold**``
    characters in generated docx/pptx. A lone, unterminated marker (e.g. a
    bullet reading "the file *.txt") is left as plain text rather than
    swallowed as unterminated markup, since the regex only matches when a
    closing marker is also present."""
    runs: list[tuple[str, bool, bool]] = []
    pos = 0
    for match in _INLINE_MARKUP_RE.finditer(text):
        if match.start() > pos:
            runs.append((text[pos : match.start()], False, False))
        if match.group("bold") is not None:
            runs.append((match.group("bold"), True, False))
        else:
            italic_text = match.group("italic1")
            if italic_text is None:
                italic_text = match.group("italic2")
            runs.append((italic_text, False, True))
        pos = match.end()
    if pos < len(text):
        runs.append((text[pos:], False, False))
    return runs or [(text, False, False)]


def _add_inline_runs(paragraph: Any, text: str) -> None:
    """Add ``text`` to a python-docx paragraph as one run per
    ``parse_inline_runs`` fragment, so ``**bold**``/``*italic*`` render as
    real formatting instead of literal markup characters."""
    for run_text, bold, italic in parse_inline_runs(text):
        run = paragraph.add_run(run_text)
        if bold:
            run.bold = True
        if italic:
            run.italic = True


def _clear_body(document: Any) -> None:
    """Remove every existing paragraph/table/etc. from a loaded template's
    body, keeping only ``sectPr`` (section properties: page size, margins)
    -- verified empirically that removing it too breaks the document, and
    that removing everything else and re-adding new content afterward
    round-trips cleanly through python-docx."""
    for child in list(document.element.body):
        if not child.tag.endswith("}sectPr"):
            document.element.body.remove(child)


def _mark_body_deleted(document: Any, author: str, change_ids: _ChangeIdCounter) -> None:
    """Template-editing + ``track_changes=True`` variant of ``_clear_body``:
    instead of silently discarding the template's existing paragraphs,
    wrap each one's runs in ``<w:del>``/``<w:delText>`` and mark its
    paragraph mark deleted too, so Word shows the old content as a
    reviewable tracked deletion rather than making it vanish outright.
    Keeps ``sectPr`` like ``_clear_body``. Tables in the template are still
    removed outright (documented v1 scope limit -- table-level tracked
    deletion needs cell/row markup this tool doesn't generate)."""
    for child in list(document.element.body):
        if child.tag.endswith("}sectPr"):
            continue
        if child.tag.endswith("}p"):
            _mark_paragraph_deleted(child, author, change_ids)
        else:
            document.element.body.remove(child)


def _add_toc(document: Any) -> Any:
    """Insert a Word TOC field (``TOC \\o "1-3" \\h \\z \\u``) at the current
    end of the document -- python-docx has no native TOC API, this is the
    standard raw-OXML field-code recipe (fldChar begin/separate/end +
    instrText), matching the docx skill's own guidance that headings must
    use built-in ``HeadingLevel``-equivalent styles (``add_heading`` already
    does that). The placeholder text only shows until Word recalculates the
    field -- ``_enable_update_fields`` makes that happen automatically on
    open instead of requiring the user to press F9."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    r_element = run._r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose Update Field to generate the table of contents."

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")

    r_element.append(begin)
    r_element.append(instr)
    r_element.append(separate)
    r_element.append(placeholder)
    r_element.append(end)
    return paragraph


def _enable_update_fields(document: Any) -> None:
    """Set ``<w:updateFields w:val="true"/>`` in settings.xml so Word
    recalculates every field (our TOC included) automatically on open,
    instead of showing stale placeholder text until the user manually
    updates it."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    settings = document.settings.element
    if settings.find(qn("w:updateFields")) is not None:
        return
    element = OxmlElement("w:updateFields")
    element.set(qn("w:val"), "true")
    settings.append(element)


def _apply_page_setup(document: Any, page_size: str, orientation: str) -> None:
    """Set page width/height/orientation on every section -- python-docx
    exposes this natively (``section.page_width``/``page_height``/
    ``orientation``), it's just never been wired into ``write_docx`` before
    now. Empty ``page_size``/``orientation`` leaves that axis at whatever
    the blank document or template already has."""
    from docx.enum.section import WD_ORIENT
    from docx.shared import Emu

    if page_size:
        key = page_size.strip().lower()
        if key not in _PAGE_SIZES_EMU:
            raise ValueError(
                f"Unknown page_size {page_size!r}. Supported: "
                f"{', '.join(sorted(_PAGE_SIZES_EMU))}"
            )
    if orientation:
        orientation_key = orientation.strip().lower()
        if orientation_key not in ("portrait", "landscape"):
            raise ValueError(
                f"Unknown orientation {orientation!r}. Use 'portrait' or 'landscape'."
            )
    else:
        orientation_key = ""

    for section in document.sections:
        width, height = (
            _PAGE_SIZES_EMU[page_size.strip().lower()]
            if page_size
            else (section.page_width, section.page_height)
        )
        if orientation_key == "landscape":
            width, height = max(width, height), min(width, height)
            section.orientation = WD_ORIENT.LANDSCAPE
        elif orientation_key == "portrait":
            width, height = min(width, height), max(width, height)
            section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Emu(int(width))
        section.page_height = Emu(int(height))


class _ChangeIdCounter:
    """Shared, monotonically increasing ``w:id`` source for every
    ``<w:ins>``/``<w:del>`` element a single ``write_docx`` call produces --
    Word requires these unique across the whole document."""

    def __init__(self) -> None:
        self._next = 0

    def next(self) -> int:
        value = self._next
        self._next += 1
        return value


def _mark_paragraph_deleted(
    paragraph_element: Any, author: str, change_ids: _ChangeIdCounter
) -> None:
    """In place: convert every run's ``<w:t>`` to ``<w:delText>``, wrap each
    run in ``<w:del>``, and mark the paragraph mark itself deleted (so
    accepting the change merges this paragraph into the next one, matching
    Word's own behavior for a fully-deleted paragraph and the docx skill's
    documented XML shape for paragraph deletions)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    date = datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")

    p_pr = paragraph_element.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        paragraph_element.insert(0, p_pr)
    r_pr = p_pr.find(qn("w:rPr"))
    if r_pr is None:
        r_pr = OxmlElement("w:rPr")
        p_pr.append(r_pr)
    mark_del = OxmlElement("w:del")
    mark_del.set(qn("w:id"), str(change_ids.next()))
    mark_del.set(qn("w:author"), author)
    mark_del.set(qn("w:date"), date)
    r_pr.insert(0, mark_del)

    for run in list(paragraph_element.findall(qn("w:r"))):
        texts = run.findall(qn("w:t"))
        if not texts:
            continue
        for text_element in texts:
            text_element.tag = qn("w:delText")
            text_element.set(qn("xml:space"), "preserve")
        run_index = list(paragraph_element).index(run)
        wrapper = OxmlElement("w:del")
        wrapper.set(qn("w:id"), str(change_ids.next()))
        wrapper.set(qn("w:author"), author)
        wrapper.set(qn("w:date"), date)
        paragraph_element.remove(run)
        wrapper.append(run)
        paragraph_element.insert(run_index, wrapper)


def _mark_paragraph_inserted(
    paragraph_element: Any, author: str, change_ids: _ChangeIdCounter
) -> None:
    """In place: wrap every direct ``<w:r>`` run child in ``<w:ins>``, so
    Word shows this (newly written) paragraph's text as a tracked insertion
    a reviewer can accept/reject, instead of silently-already-final
    content. Tables are not wrapped this way (documented v1 scope limit --
    see ``write_docx``'s docstring)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    date = datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")
    for run in list(paragraph_element.findall(qn("w:r"))):
        run_index = list(paragraph_element).index(run)
        wrapper = OxmlElement("w:ins")
        wrapper.set(qn("w:id"), str(change_ids.next()))
        wrapper.set(qn("w:author"), author)
        wrapper.set(qn("w:date"), date)
        paragraph_element.remove(run)
        wrapper.append(run)
        paragraph_element.insert(run_index, wrapper)


def _new_comment_anchor(paragraph_element: Any, comment_id: str) -> None:
    """Wrap the whole paragraph in ``<w:commentRangeStart>``/
    ``<w:commentRangeEnd>`` + append the ``<w:commentReference>`` run, so
    the comment (written separately to ``comments.xml``) anchors to this
    entire paragraph/bullet/heading. No sub-paragraph anchor in this tool's
    v1 -- a comment always spans the whole block it's attached to."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), comment_id)
    p_pr = paragraph_element.find(qn("w:pPr"))
    paragraph_element.insert(1 if p_pr is not None else 0, start)

    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), comment_id)
    paragraph_element.append(end)

    ref_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "CommentReference")
    r_pr.append(r_style)
    ref_run.append(r_pr)
    comment_ref = OxmlElement("w:commentReference")
    comment_ref.set(qn("w:id"), comment_id)
    ref_run.append(comment_ref)
    paragraph_element.append(ref_run)


def _build_comments_xml(comments: list[tuple[str, str, str]]) -> bytes:
    """``comments``: list of ``(comment_id, author, text)``. Builds the
    standalone ``word/comments.xml`` part -- small and flat enough that
    plain string templating (with proper XML escaping) is simpler and just
    as safe here as a full XML-tree build."""
    now = datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")
    parts = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">',
    ]
    for comment_id, author, text in comments:
        initials = "".join(part[0] for part in author.split()[:2]).upper() or "?"
        parts.append(
            f'<w:comment w:id="{comment_id}" w:author="{_xml_escape(author)}" '
            f'w:date="{now}" w:initials="{_xml_escape(initials)}">'
            f"<w:p><w:r><w:t xml:space=\"preserve\">{_xml_escape(text)}</w:t></w:r></w:p>"
            f"</w:comment>"
        )
    parts.append("</w:comments>")
    return "".join(parts).encode("utf-8")


def _write_comments_part(file_path: Path, comments_xml: bytes) -> None:
    """Append ``word/comments.xml`` + its content-type override + a
    relationship to an already-saved .docx. python-docx has no public API
    for adding a part type it doesn't itself know about, so this is the
    same "unzip, patch, rezip" approach the docx skill documents for adding
    comments to an existing file -- applied here right after this tool's
    own from-scratch save. ``word/document.xml`` itself is never re-parsed
    here: python-docx already wrote the comment-range anchors correctly as
    part of the normal save, since ``_new_comment_anchor`` mutated the
    in-memory tree before ``document.save()`` was called."""
    from lxml import etree

    with zipfile.ZipFile(file_path, "r") as zin:
        entries = {name: zin.read(name) for name in zin.namelist()}

    entries["word/comments.xml"] = comments_xml

    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    content_types = etree.fromstring(entries["[Content_Types].xml"])
    override = etree.SubElement(content_types, f"{{{ct_ns}}}Override")
    override.set("PartName", "/word/comments.xml")
    override.set(
        "ContentType",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
    )
    entries["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    pr_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    rels_path = "word/_rels/document.xml.rels"
    rels_root = etree.fromstring(entries[rels_path])
    max_n = 0
    for rel in rels_root.findall(f"{{{pr_ns}}}Relationship"):
        rid = rel.get("Id") or ""
        if rid.startswith("rId") and rid[3:].isdigit():
            max_n = max(max_n, int(rid[3:]))
    relationship = etree.SubElement(rels_root, f"{{{pr_ns}}}Relationship")
    relationship.set("Id", f"rId{max_n + 1}")
    relationship.set(
        "Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
    )
    relationship.set("Target", "comments.xml")
    entries[rels_path] = etree.tostring(
        rels_root, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    with zipfile.ZipFile(file_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, payload in entries.items():
            zout.writestr(name, payload)


def _ensure_para_ids(document_element: Any) -> None:
    """Assign a ``w14:paraId`` to every ``<w:p>`` that doesn't already have
    one. ``_native_docx_tracks.tracks.TracksMixin``'s whole paragraph-lookup
    contract is keyed on this attribute -- real Word always sets it, but
    ``python-docx`` never does, so a freshly-written or freshly-edited
    document needs it backfilled before any tracked-change tool can address
    a paragraph. ``w14`` is a Word 2010+ extension namespace, not part of
    the base ``wml.xsd``, so this also adds it to the document root's
    ``mc:Ignorable`` list when missing -- the same declaration real
    Word-authored files carry, and the one ``_strip_mce_ignorable`` already
    knows how to treat as extension content rather than a schema defect."""
    existing: set[str] = set()
    missing: list[Any] = []
    for paragraph in document_element.iter(f"{_tracks_constants.W}p"):
        para_id = paragraph.get(f"{_tracks_constants.W14}paraId")
        if para_id:
            existing.add(para_id.upper())
        else:
            missing.append(paragraph)
    if not missing:
        return

    ignorable_tokens = (document_element.get(_MC_IGNORABLE) or "").split()
    if "w14" not in ignorable_tokens:
        ignorable_tokens.append("w14")
        document_element.set(_MC_IGNORABLE, " ".join(ignorable_tokens))

    for paragraph in missing:
        while True:
            candidate = f"{random.randint(1, 0x7FFFFFFF):08X}"
            if candidate not in existing:
                existing.add(candidate)
                break
        paragraph.set(f"{_tracks_constants.W14}paraId", candidate)


def _locate_paragraph(
    document_element: Any,
    find: str,
    context_before: str,
    context_after: str,
    *,
    ignore_case: bool,
) -> str:
    """Return the ``w14:paraId`` of the single paragraph where
    ``tracks.py``'s own per-paragraph ``_resolve`` would unambiguously
    locate ``find`` (respecting ``context_before``/``context_after``) --
    tried against every paragraph in the document, since coscribe's own
    tools address text by its content, not by a paragraph id the caller has
    no way to already know. Raises ``ValueError`` if no paragraph matches,
    or if more than one does."""
    if not find:
        raise ValueError("find text is empty")
    matches: list[str] = []
    ambiguous_within: list[str] = []
    for paragraph in document_element.iter(f"{_tracks_constants.W}p"):
        para_id = paragraph.get(f"{_tracks_constants.W14}paraId", "")
        try:
            _resolve(
                document_element,
                paragraph,
                find,
                context_before,
                context_after,
                ignore_case=ignore_case,
            )
        except ValueError as exc:
            if "Ambiguous" in str(exc):
                ambiguous_within.append(f"paragraph {para_id}: {exc}")
            continue
        matches.append(para_id)

    if not matches:
        if ambiguous_within:
            raise ValueError(
                f"Text {find!r} is ambiguous within a paragraph; provide "
                f"context_before/context_after. " + "; ".join(ambiguous_within)
            )
        raise ValueError(f"Text {find!r} not found in the document")
    if len(matches) > 1:
        raise ValueError(
            f"Text {find!r} found in {len(matches)} different paragraphs; "
            "provide context_before or context_after to disambiguate"
        )
    return matches[0]


class _TrackedDocxHost(TracksMixin):
    """Adapts ``TracksMixin``'s small host contract to an already-open
    python-docx ``Document`` -- coscribe already has the tree in memory via
    ``Document(path)``/``document.save()``, so it doesn't need docx-mcp's
    own unzip/parse/repack lifecycle (see ``_native_docx_tracks/NOTICE.md``).
    Only ``_require`` and ``_mark`` are adapter-specific; ``_find_para``/
    ``_next_markup_id``/``_make_run`` are the same plain functions
    ``tracks.py`` itself was extracted alongside (see ``_constants.py``),
    bound here via ``staticmethod`` since ``tracks.py`` calls them as
    ``self.<name>(...)``."""

    _find_para = staticmethod(_tracks_constants._find_para)
    _next_markup_id = staticmethod(_tracks_constants._next_markup_id)
    _make_run = staticmethod(_tracks_constants._make_run)

    def __init__(self, document: Any) -> None:
        self._document_element = document.element
        _ensure_para_ids(self._document_element)

    def _require(self, rel_path: str) -> Any:
        if rel_path != "word/document.xml":
            raise RuntimeError(f"{rel_path} not supported by tracked-change edits")
        return self._document_element

    def _mark(self, rel_path: str) -> None:
        pass  # python-docx re-serializes the whole tree on save() regardless


class DocumentToolkit:
    def __init__(
        self,
        root: str | Path,
        *,
        state_dir: str | Path | None = None,
        extra_readable: Sequence[str | Path] = (),
        extra_writable: Sequence[str | Path] = (),
    ) -> None:
        self._scope = WorkspaceScope(
            root, extra_readable=extra_readable, extra_writable=extra_writable
        )
        self._state_dir = Path(state_dir) if state_dir is not None else None

    def _check_readable(self, path: str) -> Path:
        file_path = self._scope.resolve(path)
        if not file_path.exists():
            raise ValueError(f"File does not exist: {path}")
        if not file_path.is_file():
            raise ValueError(f"Path is not a file: {path}")
        return file_path

    def _check_writable(self, path: str, overwrite: bool) -> Path:
        file_path = self._scope.resolve(path, write=True)
        if file_path.exists() and file_path.is_dir():
            raise ValueError(f"Path is a directory: {path}")
        if file_path.exists() and not overwrite:
            raise FileExistsError(f"File already exists: {path}")
        file_path.parent.mkdir(parents=True, exist_ok=True)
        return file_path

    def _check_editable(self, path: str) -> Path:
        """Like `_check_readable`, but resolved through the write-permitted
        scope -- the tracked-change tools edit an existing file in place, so
        they need both: the file must already exist (unlike `write_docx`,
        these tools have nothing to write if it doesn't), and the location
        must be one the caller is allowed to write back to."""
        file_path = self._scope.resolve(path, write=True)
        if not file_path.exists():
            raise ValueError(f"File does not exist: {path}")
        if not file_path.is_file():
            raise ValueError(f"Path is not a file: {path}")
        return file_path

    def read_docx(self, path: str) -> str:
        import mammoth
        from markdownify import markdownify

        file_path = self._check_readable(path)
        with file_path.open("rb") as docx_file:
            html = mammoth.convert_to_html(docx_file).value
        return markdownify(
            html, heading_style="ATX", bullets="-", table_infer_header=True
        ).strip()

    def write_docx(
        self,
        path: str,
        content: str,
        overwrite: bool = True,
        template_path: str = "",
        page_size: str = "",
        orientation: str = "",
        track_changes: bool = False,
        change_author: str = "Coscribe",
        comment_author: str = "Coscribe",
    ) -> dict[str, object]:
        from docx import Document

        file_path = self._check_writable(path, overwrite)
        change_ids = _ChangeIdCounter()
        if template_path:
            template_file = self._check_readable(template_path)
            document = Document(str(template_file))
            if track_changes:
                _mark_body_deleted(document, change_author, change_ids)
            else:
                _clear_body(document)
        else:
            document = Document()

        has_toc = False
        pending_comments: list[tuple[Any, str]] = []
        for block in parse_blocks(content):
            paragraph = None
            if block.kind == "heading":
                paragraph = document.add_heading("", level=block.level)
                _add_inline_runs(paragraph, block.text)
            elif block.kind == "bullet":
                paragraph = document.add_paragraph(style="List Bullet")
                _add_inline_runs(paragraph, block.text)
            elif block.kind == "number":
                paragraph = document.add_paragraph(style="List Number")
                _add_inline_runs(paragraph, block.text)
            elif block.kind == "toc":
                paragraph = _add_toc(document)
                has_toc = True
            elif block.kind == "table" and block.rows:
                n_cols = max(len(row) for row in block.rows)
                table = document.add_table(rows=len(block.rows), cols=n_cols)
                table.style = "Table Grid"
                for row_index, row in enumerate(block.rows):
                    for col_index in range(n_cols):
                        cell_text = row[col_index] if col_index < len(row) else ""
                        _add_inline_runs(table.cell(row_index, col_index).paragraphs[0], cell_text)
            else:
                paragraph = document.add_paragraph()
                _add_inline_runs(paragraph, block.text)

            if paragraph is None:
                continue
            if track_changes:
                _mark_paragraph_inserted(paragraph._p, change_author, change_ids)
            if block.comment:
                pending_comments.append((paragraph._p, block.comment))

        if has_toc:
            _enable_update_fields(document)
        if page_size or orientation:
            _apply_page_setup(document, page_size, orientation)
        for index, (paragraph_element, _) in enumerate(pending_comments):
            _new_comment_anchor(paragraph_element, str(index))

        assert_wml_valid(document.element, "write_docx")
        document.save(str(file_path))

        if pending_comments:
            comments_xml = _build_comments_xml(
                [
                    (str(index), comment_author, text)
                    for index, (_, text) in enumerate(pending_comments)
                ]
            )
            _write_comments_part(file_path, comments_xml)

        preview_path, preview_skipped_reason = render_thumbnail(file_path, self._state_dir)
        return {
            "path": self._scope.relative(file_path),
            "bytes_written": file_path.stat().st_size,
            "preview_path": preview_path,
            "preview_skipped_reason": preview_skipped_reason,
            "has_toc": has_toc,
            "comments_added": len(pending_comments),
            "tracked_changes": track_changes,
        }

    def _open_for_tracked_edit(self, path: str) -> tuple[Path, Any, _TrackedDocxHost]:
        from docx import Document

        file_path = self._check_editable(path)
        document = Document(str(file_path))
        return file_path, document, _TrackedDocxHost(document)

    def _save_tracked_edit(self, file_path: Path, document: Any) -> dict[str, object]:
        assert_wml_valid(document.element, "tracked_docx_edit")
        document.save(str(file_path))
        preview_path, preview_skipped_reason = render_thumbnail(file_path, self._state_dir)
        return {
            "path": self._scope.relative(file_path),
            "bytes_written": file_path.stat().st_size,
            "preview_path": preview_path,
            "preview_skipped_reason": preview_skipped_reason,
        }

    def insert_docx_text(
        self,
        path: str,
        text: str,
        context_before: str,
        context_after: str = "",
        track_changes: bool = True,
        author: str = "Coscribe",
        ignore_case: bool = False,
    ) -> dict[str, object]:
        if not context_before:
            raise ValueError(
                "context_before is required -- it names the existing text to "
                "insert after; context_after alone cannot locate an insertion point"
            )
        file_path, document, host = self._open_for_tracked_edit(path)
        para_id = _locate_paragraph(
            host._document_element, context_before, "", context_after, ignore_case=ignore_case
        )
        result = host.insert_text(
            para_id,
            text,
            author=author,
            context_before=context_before,
            context_after=context_after,
            ignore_case=ignore_case,
            tracked=track_changes,
        )
        result.update(self._save_tracked_edit(file_path, document))
        return result

    def delete_docx_text(
        self,
        path: str,
        text: str,
        context_before: str = "",
        context_after: str = "",
        track_changes: bool = True,
        author: str = "Coscribe",
        ignore_case: bool = False,
    ) -> dict[str, object]:
        file_path, document, host = self._open_for_tracked_edit(path)
        para_id = _locate_paragraph(
            host._document_element, text, context_before, context_after, ignore_case=ignore_case
        )
        result = host.delete_text(
            para_id,
            text,
            author=author,
            context_before=context_before,
            context_after=context_after,
            ignore_case=ignore_case,
            tracked=track_changes,
        )
        result.update(self._save_tracked_edit(file_path, document))
        return result

    def replace_docx_text(
        self,
        path: str,
        find: str,
        replace: str,
        context_before: str = "",
        context_after: str = "",
        track_changes: bool = True,
        author: str = "Coscribe",
        ignore_case: bool = False,
    ) -> dict[str, object]:
        file_path, document, host = self._open_for_tracked_edit(path)
        para_id = _locate_paragraph(
            host._document_element, find, context_before, context_after, ignore_case=ignore_case
        )
        result = host.replace_text(
            para_id,
            find=find,
            replace=replace,
            author=author,
            context_before=context_before,
            context_after=context_after,
            ignore_case=ignore_case,
            tracked=track_changes,
        )
        result.update(self._save_tracked_edit(file_path, document))
        return result

    def accept_docx_tracked_changes(self, path: str, author: str = "") -> dict[str, object]:
        file_path, document, host = self._open_for_tracked_edit(path)
        result = host.accept_changes(author=author or None)
        result.update(self._save_tracked_edit(file_path, document))
        return result

    def reject_docx_tracked_changes(self, path: str, author: str = "") -> dict[str, object]:
        file_path, document, host = self._open_for_tracked_edit(path)
        result = host.reject_changes(author=author or None)
        result.update(self._save_tracked_edit(file_path, document))
        return result

    def read_pdf(self, path: str, start_page: int = 1, end_page: int = 0) -> str:
        file_path = self._check_readable(path)
        pages = []
        for index, text in _pdf_page_texts(file_path, start_page, end_page):
            pages.append(f"--- Page {index} ---\n{text.strip()}")
        return "\n\n".join(pages)

    def search_pdf(
        self, query: str, path: str = ".", pattern: str = "*.pdf", regex: bool = False
    ) -> list[dict[str, object]]:
        base = self._scope.resolve(path)
        if not base.exists():
            raise ValueError(f"Path does not exist: {path}")
        if regex:
            try:
                compiled = re.compile(query)
            except re.error as exc:
                raise ValueError(f"Invalid regular expression {query!r}: {exc}") from exc

            def is_match(line: str) -> bool:
                return compiled.search(line) is not None
        else:

            def is_match(line: str) -> bool:
                return query in line

        candidates = [base] if base.is_file() else sorted(base.rglob(pattern))
        matches: list[dict[str, object]] = []
        for item in candidates:
            if not item.is_file():
                continue
            relative_parts = item.relative_to(self._scope.root).parts
            if any(part in DEFAULT_IGNORES for part in relative_parts):
                continue
            for page_number, text in _pdf_page_texts(item):
                for line in text.splitlines():
                    if is_match(line):
                        matches.append(
                            {
                                "path": self._scope.relative(item),
                                "page": page_number,
                                "text": line,
                            }
                        )
                    if len(matches) >= 100:
                        return matches
        return matches

    def write_pdf(self, path: str, content: str, overwrite: bool = True) -> dict[str, object]:
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        file_path = self._check_writable(path, overwrite)
        styles = getSampleStyleSheet()
        bullet_style = ParagraphStyle("CoscribeBullet", parent=styles["Normal"], leftIndent=18)

        flowables: list[Any] = []
        for block in parse_blocks(content):
            if block.kind == "heading":
                heading_style = styles[f"Heading{block.level}"]
                flowables.append(Paragraph(_xml_escape(block.text), heading_style))
            elif block.kind == "bullet":
                # Plain ASCII, not "•" -- reportlab's default font maps
                # the unicode bullet glyph in a way pdfplumber/pypdf both
                # mis-render on read-back (a stray "(cid:127)" or a dropped
                # character), confirmed while testing the write/read round trip.
                flowables.append(Paragraph(f"- {_xml_escape(block.text)}", bullet_style))
            elif block.kind == "number":
                numbered_text = f"{block.ordinal}. {_xml_escape(block.text)}"
                flowables.append(Paragraph(numbered_text, bullet_style))
            elif block.kind == "table" and block.rows:
                table = Table(block.rows)
                table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                            ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ]
                    )
                )
                flowables.append(table)
            else:
                flowables.append(Paragraph(_xml_escape(block.text), styles["Normal"]))
            flowables.append(Spacer(1, 6))

        SimpleDocTemplate(str(file_path)).build(flowables)
        preview_path, preview_skipped_reason = render_thumbnail(file_path, self._state_dir)
        return {
            "path": self._scope.relative(file_path),
            "bytes_written": file_path.stat().st_size,
            "preview_path": preview_path,
            "preview_skipped_reason": preview_skipped_reason,
        }


def build_document_tools(
    root: str | Path,
    *,
    state_dir: str | Path | None = None,
    extra_readable: Sequence[str | Path] = (),
    extra_writable: Sequence[str | Path] = (),
) -> list[Callable[..., Any]]:
    """Return the tool callables the Coordinator agent can call, bound to `root`
    (plus any user-configured extra_readable/extra_writable directories)."""
    toolkit = DocumentToolkit(
        root, state_dir=state_dir, extra_readable=extra_readable, extra_writable=extra_writable
    )

    def read_docx(path: str) -> str:
        """Read a Word (.docx) file under the workspace as markdown.

        Headings, lists, tables, links, and bold/italic all come through as
        their markdown equivalents.

        Args:
            path: file to read, relative to the workspace root
        """
        return toolkit.read_docx(path=path)

    def write_docx(
        path: str,
        content: str,
        overwrite: bool = True,
        template_path: str = "",
        page_size: str = "",
        orientation: str = "",
        track_changes: bool = False,
        change_author: str = "Coscribe",
        comment_author: str = "Coscribe",
    ) -> dict[str, object]:
        """Create a Word (.docx) file under the workspace from structured text.

        `content` uses a light markdown subset: `#`/`##`/`###` for headings,
        `-`/`*` for bullet items, `1.` for numbered items, and
        `| cell | cell |` rows for tables. Plain lines become paragraphs.
        Two extra markers on top of that subset:

        - A line that is exactly `[TOC]` inserts a real Word table-of-
          contents field at that position, built from the document's own
          heading levels. It shows placeholder text until Word recalculates
          it, which happens automatically the first time the file is opened
          -- no manual "Update Field" needed.
        - A trailing `{{comment: your comment text}}` on a heading,
          paragraph, bullet, or numbered-item line attaches a real Word
          comment (visible in Word's Comments pane) anchored to that whole
          line, authored as `comment_author`. Not recognized inside table
          cells.

        The response's `preview_path` (when LibreOffice is installed) names
        a rendered thumbnail of the first page the user can see in the
        chat UI -- not something to fetch or parse yourself.

        `template_path`, if given, is an existing .docx file (e.g. one with
        the user's own letterhead/styles/page setup) to build the new
        document on: its styles and page setup (size, margins) carry over.
        Without `track_changes`, its own body content is removed before the
        new content from `content` is added. With `track_changes=True`, the
        template's existing paragraphs are kept but marked as a tracked
        deletion instead (visible to the user as a reviewable redline, not
        silently discarded) -- tables in the template are still removed
        outright either way.

        `page_size` ("letter", "legal", "a4", or "a3") and `orientation`
        ("portrait" or "landscape") override the blank document's or
        template's page setup; leave either empty to keep the default.

        `track_changes=True` wraps every paragraph/bullet/heading this call
        writes in a tracked insertion (`w:ins`, authored as `change_author`)
        instead of writing it as already-final content -- use this when the
        document should go to a human as a reviewable redline they can
        accept/reject in Word, rather than a silent rewrite. Tables are
        always written as final content, not tracked, even with
        `track_changes=True` (documented scope limit -- table-level tracked
        changes need row/cell markup this tool doesn't generate).

        Args:
            path: file to write, relative to the workspace root
            content: document content in the markdown subset described above
            overwrite: whether to replace the file if it already exists
            template_path: optional .docx file to use as the document's
                starting styles/page-setup template, relative to the
                workspace root
            page_size: "letter", "legal", "a4", "a3", or empty to leave
                the current page size unchanged
            orientation: "portrait", "landscape", or empty to leave the
                current orientation unchanged
            track_changes: write this call's content as tracked
                insertions (and, with template_path, mark the template's
                old content as a tracked deletion instead of removing it)
            change_author: author name recorded on tracked-change markup
                when track_changes is True
            comment_author: author name recorded on `{{comment: ...}}`
                markers
        """
        return toolkit.write_docx(
            path=path,
            content=content,
            overwrite=overwrite,
            template_path=template_path,
            page_size=page_size,
            orientation=orientation,
            track_changes=track_changes,
            change_author=change_author,
            comment_author=comment_author,
        )

    def insert_docx_text(
        path: str,
        text: str,
        context_before: str,
        context_after: str = "",
        track_changes: bool = True,
        author: str = "Coscribe",
        ignore_case: bool = False,
    ) -> dict[str, object]:
        """Insert text into an existing Word (.docx) file, in place.

        Edits the real document tree directly -- headers/footers, other
        paragraphs, tables, and images are left completely untouched,
        unlike `write_docx` (which regenerates the whole document from a
        markdown subset and would silently drop anything that subset can't
        represent). The target location is found from `context_before` --
        the existing text to insert immediately after -- not a paragraph
        number or id; this must be unique in the document, or use
        `context_after` (text immediately following the insertion point)
        to disambiguate.

        With `track_changes=True` (the default), the inserted text is
        wrapped in `<w:ins>` tracked-changes markup Word shows in its
        Review pane, rather than becoming final text immediately -- use
        `track_changes=False` for a plain, silent edit.

        Args:
            path: existing .docx file to edit, relative to the workspace root
            text: the text to insert
            context_before: existing document text to insert immediately
                after -- required, and must be unique (with context_after
                to disambiguate if it appears more than once)
            context_after: existing text immediately following the
                insertion point, to disambiguate when context_before alone
                is not unique
            track_changes: insert as a reviewable tracked change (default)
                rather than a plain, already-final edit
            author: author name recorded on the tracked-change markup, when
                track_changes is True
            ignore_case: match context_before/context_after case-insensitively
        """
        return toolkit.insert_docx_text(
            path=path,
            text=text,
            context_before=context_before,
            context_after=context_after,
            track_changes=track_changes,
            author=author,
            ignore_case=ignore_case,
        )

    def delete_docx_text(
        path: str,
        text: str,
        context_before: str = "",
        context_after: str = "",
        track_changes: bool = True,
        author: str = "Coscribe",
        ignore_case: bool = False,
    ) -> dict[str, object]:
        """Delete text from an existing Word (.docx) file, in place.

        Edits the real document tree directly -- headers/footers, other
        paragraphs, tables, and images are left completely untouched,
        unlike `write_docx` (which regenerates the whole document from a
        markdown subset and would silently drop anything that subset can't
        represent). `text` must be unique in the document; use
        `context_before`/`context_after` (surrounding text) to disambiguate
        if it appears more than once.

        With `track_changes=True` (the default), the text is wrapped in
        `<w:del>` tracked-changes markup Word shows in its Review pane
        (strikethrough) rather than being removed immediately -- use
        `track_changes=False` to remove it outright.

        Args:
            path: existing .docx file to edit, relative to the workspace root
            text: the exact text to delete
            context_before: text immediately before `text`, to disambiguate
                if it appears more than once
            context_after: text immediately after `text`, to disambiguate
                if it appears more than once
            track_changes: mark as a reviewable tracked deletion (default)
                rather than removing the text outright
            author: author name recorded on the tracked-change markup, when
                track_changes is True
            ignore_case: match text/context case-insensitively
        """
        return toolkit.delete_docx_text(
            path=path,
            text=text,
            context_before=context_before,
            context_after=context_after,
            track_changes=track_changes,
            author=author,
            ignore_case=ignore_case,
        )

    def replace_docx_text(
        path: str,
        find: str,
        replace: str,
        context_before: str = "",
        context_after: str = "",
        track_changes: bool = True,
        author: str = "Coscribe",
        ignore_case: bool = False,
    ) -> dict[str, object]:
        """Replace text in an existing Word (.docx) file, in place.

        Edits the real document tree directly -- headers/footers, other
        paragraphs, tables, and images are left completely untouched,
        unlike `write_docx` (which regenerates the whole document from a
        markdown subset and would silently drop anything that subset can't
        represent). `find` must be unique in the document; use
        `context_before`/`context_after` to disambiguate if it appears
        more than once.

        With `track_changes=True` (the default), marks only the actually-
        changed portion of `find` as a tracked deletion+insertion (e.g.
        replacing "red" with "blue" in "the red car" tracks just "red" ->
        "blue", not the whole sentence) -- use `track_changes=False` for a
        plain, silent edit.

        Args:
            path: existing .docx file to edit, relative to the workspace root
            find: the exact existing text to replace
            replace: the new text
            context_before: text immediately before `find`, to disambiguate
                if it appears more than once
            context_after: text immediately after `find`, to disambiguate
                if it appears more than once
            track_changes: mark as a reviewable tracked change (default)
                rather than a plain, already-final edit
            author: author name recorded on the tracked-change markup, when
                track_changes is True
            ignore_case: match find/context case-insensitively
        """
        return toolkit.replace_docx_text(
            path=path,
            find=find,
            replace=replace,
            context_before=context_before,
            context_after=context_after,
            track_changes=track_changes,
            author=author,
            ignore_case=ignore_case,
        )

    def accept_docx_tracked_changes(path: str, author: str = "") -> dict[str, object]:
        """Accept tracked changes in an existing Word (.docx) file.

        Insertions become permanent text; deletions are permanently
        removed. Leave `author` empty to accept every tracked change in the
        document, or name one author to accept only that author's changes
        and leave everyone else's still pending review.

        Args:
            path: existing .docx file to edit, relative to the workspace root
            author: accept only this author's changes, or empty for all
        """
        return toolkit.accept_docx_tracked_changes(path=path, author=author)

    def reject_docx_tracked_changes(path: str, author: str = "") -> dict[str, object]:
        """Reject tracked changes in an existing Word (.docx) file.

        Insertions are discarded; deletions are restored. Leave `author`
        empty to reject every tracked change in the document, or name one
        author to reject only that author's changes and leave everyone
        else's still pending review.

        Args:
            path: existing .docx file to edit, relative to the workspace root
            author: reject only this author's changes, or empty for all
        """
        return toolkit.reject_docx_tracked_changes(path=path, author=author)

    def read_pdf(path: str, start_page: int = 1, end_page: int = 0) -> str:
        """Extract the text contents of a PDF file under the workspace, page by page.

        For a long PDF, find the relevant pages with `search_pdf` first and
        read only that range, rather than the whole document.

        Args:
            path: file to read, relative to the workspace root
            start_page: first page to read (1-based)
            end_page: last page to read, inclusive; 0 means through the last page
        """
        return toolkit.read_pdf(path=path, start_page=start_page, end_page=end_page)

    def write_pdf(path: str, content: str, overwrite: bool = True) -> dict[str, object]:
        """Create a PDF file under the workspace from structured text.

        Uses the same markdown subset as `write_docx` (headings, bullets,
        numbered items, pipe tables, and plain paragraphs).

        The response's `preview_path` (when LibreOffice is installed) names
        a rendered thumbnail of the first page the user can see in the
        chat UI -- not something to fetch or parse yourself.

        Args:
            path: file to write, relative to the workspace root
            content: document content in the markdown subset described above
            overwrite: whether to replace the file if it already exists
        """
        return toolkit.write_pdf(path=path, content=content, overwrite=overwrite)

    def search_pdf(
        query: str, path: str = ".", pattern: str = "*.pdf", regex: bool = False
    ) -> list[dict[str, object]]:
        """Search PDF files under the workspace line by line, returning each
        matching line with its page number (at most 100 matches).

        Finds matches without needing to read_pdf an entire (possibly large)
        document into context first.

        Args:
            query: substring to search for, or a Python regular expression when regex is true
            path: a PDF file, or a directory to search under, relative to the workspace root
            pattern: glob pattern to filter PDF files by when path is a directory
            regex: treat query as a regular expression instead of a literal substring
        """
        return toolkit.search_pdf(query=query, path=path, pattern=pattern, regex=regex)

    return [
        tool_metadata(read_docx, risk_category="READ", category="documents"),
        tool_metadata(write_docx, risk_category="WRITE_LOCAL", category="documents"),
        tool_metadata(insert_docx_text, risk_category="WRITE_LOCAL", category="documents"),
        tool_metadata(delete_docx_text, risk_category="WRITE_LOCAL", category="documents"),
        tool_metadata(replace_docx_text, risk_category="WRITE_LOCAL", category="documents"),
        tool_metadata(
            accept_docx_tracked_changes, risk_category="WRITE_LOCAL", category="documents"
        ),
        tool_metadata(
            reject_docx_tracked_changes, risk_category="WRITE_LOCAL", category="documents"
        ),
        tool_metadata(read_pdf, risk_category="READ", category="documents"),
        tool_metadata(write_pdf, risk_category="WRITE_LOCAL", category="documents"),
        tool_metadata(search_pdf, risk_category="READ", category="documents"),
    ]
