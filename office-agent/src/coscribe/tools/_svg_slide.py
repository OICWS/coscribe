"""A small, deliberately limited SVG-subset -> native PowerPoint shapes
converter, powering ``write_pptx``'s ``layout: svg`` directive
(``presentations.py``).

This is **not** a general SVG renderer. Live research this session (see
``builtin_skills/pptx/SKILL.md``'s "Prefab layouts" section for the full
story) confirmed SVG is a genuinely better intermediate format than raw
pptxgenjs/DrawingML for a model to author a custom slide in -- it shares
DrawingML's own absolute-coordinate 2D vector model, and is far better
represented in LLM training data -- but building (or vendoring) a *complete*
SVG-to-OOXML converter was deliberately rejected as out of scope: a real
open-source implementation (``hugohe3/ppt-master``) was evaluated and its
core technique (SVG as intermediate format, converted deterministically to
native DrawingML shapes) was verified to work end-to-end via a live spike,
but its own code carries a deliberate integrity gate requiring its entire
118MB distribution to ship intact, which is disproportionate to what this
project needs and entangles licensing/attribution concerns beyond a plain
MIT credit. This module applies the same general *technique* -- confirmed
to work by that spike -- as a small, fully independent implementation.
Nothing here is copied from any other project.

Extended (after the decision to build decorative graphics via code instead
of a diffusion image-generation API -- deterministic, zero licensing risk,
zero per-call cost, and stays on-theme) to also cover ``<path>`` (a fixed
subset of path commands) and gradient fills, the two primitives real flat-
illustration/abstract-decoration SVGs actually rely on. Still not a general
renderer: anything else raises ``ValueError`` naming ``run_node_script`` as
the real fallback, the same "fail loud, name the alternative" pattern
already used by ``add_pptx_chart`` (bar/line/pie only) and
``add_pptx_animation`` (fade/fly-in only) elsewhere in this package.

Supported elements:

- ``<svg width height viewBox?>`` -- the root; establishes the coordinate
  scale. A ``viewBox`` takes priority over ``width``/``height`` when both
  are given (matches real SVG semantics).
- ``<rect x y width height rx? fill?>`` -- a rounded rectangle if ``rx`` is
  present, else a plain rectangle. A rect at (0, 0) covering the whole
  canvas, first in document order, is promoted to the slide's own
  background fill instead of becoming a shape (mirrors what the spike
  showed looks clean, and keeps ``_check_missing_visual_elements`` honest --
  a background-only rect shouldn't itself count as "a visual element").
- ``<circle cx cy r fill?>`` / ``<ellipse cx cy rx ry fill?>``.
- ``<text x y font-family? font-size? font-weight? fill?>text</text>`` --
  one run per element (no ``<tspan>`` support).
- ``<path d fill?>`` -- ``d`` supports ``M/L/H/V/C/Q/Z`` (absolute
  uppercase or relative lowercase), including SVG's own implicit-repeat
  grammar (extra coordinate pairs after ``M``/``L`` without a new command
  letter). No arcs (``A``) and no shorthand curves (``S``/``T``) -- a
  model wanting those should fall back to ``run_node_script``. Converted
  to a native DrawingML freeform shape (``<a:custGeom>``); python-pptx's
  own ``FreeformBuilder`` only supports straight line segments, so
  ``<a:cubicBezTo>``/``<a:quadBezTo>`` are hand-appended, the same
  "reach into python-pptx's oxml layer directly" pattern already used
  elsewhere in this package (e.g. ``add_pptx_scrim``'s hand-appended
  ``<a:alpha>``). The shape's bounding box is computed from *all* path
  points including bezier control points (a safe overestimate per the
  convex-hull property, not a pixel-tight fit -- acceptable for a
  decorative shape).
- ``fill="url(#id)"`` on any of the above (plus the background-promoted
  ``<rect>``), referencing a ``<linearGradient id>``/``<radialGradient
  id>`` -- found anywhere in the document, ``<defs>`` wrapper or not, via
  a full-tree scan, since ``<defs>`` itself is never rendered. Linear
  gradients honor ``x1/y1/x2/y2`` (fractions or percentages, SVG's own
  ``objectBoundingBox`` default) as a direction vector; radial gradients
  are always centered (no off-center focus point support -- v1 scope).
  ``<stop offset stop-color>`` only, no ``stop-opacity``.

The whole SVG canvas is scaled *uniformly* (never stretched non-uniformly)
to fit the deck's real ``slide_width``/``slide_height``, centered with
letterboxing if the aspect ratios differ -- deliberately not a hardcoded
16:9/4:3 assumption, since the model doesn't know which template it's
writing into.
"""

from __future__ import annotations

import math
import re
from collections.abc import Callable
from dataclasses import dataclass
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from pptx.presentation import Presentation as PresentationType

# A modest, fixed set of CSS color keywords likely to show up in
# model-authored SVG -- not the full CSS named-color list (147 entries),
# just enough common ground that a model reaching for "white"/"navy"/"gold"
# instead of a hex code doesn't hit a wall. Anything else must be a 6-hex-
# digit color, same convention write_pptx's other layouts already use.
_NAMED_COLORS = {
    "white": "FFFFFF", "black": "000000", "red": "FF0000", "green": "008000",
    "blue": "0000FF", "yellow": "FFFF00", "orange": "FFA500", "purple": "800080",
    "gray": "808080", "grey": "808080", "silver": "C0C0C0", "navy": "000080",
    "teal": "008080", "olive": "808000", "maroon": "800000", "lime": "00FF00",
    "aqua": "00FFFF", "cyan": "00FFFF", "magenta": "FF00FF", "fuchsia": "FF00FF",
    "pink": "FFC0CB", "brown": "A52A2A", "gold": "FFD700", "indigo": "4B0082",
    "violet": "EE82EE", "coral": "FF7F50", "crimson": "DC143C",
    "darkblue": "00008B", "darkgreen": "006400", "darkred": "8B0000",
    "lightgray": "D3D3D3", "lightgrey": "D3D3D3", "beige": "F5F5DC",
}

_HEX_COLOR_RE = re.compile(r"^#?([0-9A-Fa-f]{6})$")
_URL_REF_RE = re.compile(r"^url\(#([^)]+)\)$")
# write_pptx's own default canvas (see presentations.py's
# _WIDESCREEN_WIDTH_EMU) -- 16:9 widescreen, 13.333in x 7.5in. This is a
# defensive fallback only: add_svg_slide always operates on a prs already
# created by write_pptx or loaded by fill_pptx_template, so
# prs.slide_width/height are never actually None in practice.
_DEFAULT_SLIDE_WIDTH_EMU = 12192000
_DEFAULT_SLIDE_HEIGHT_EMU = 6858000
_EMU_PER_POINT = 12700
_BLANK_LAYOUT = 6

# <path> "d" parsing -- a fixed, deliberately small command subset (see
# module docstring). Local path-space coordinates are scaled up by this
# factor before being stored as DrawingML custGeom integers, so
# fractional/sub-unit curve control points (common in hand-tuned bezier
# paths) don't lose precision to integer truncation.
_PATH_UNIT_SCALE = 1000
# Matches *any* letter (not just the supported ones) so an unsupported
# command (e.g. "A" for arcs) is tokenized as a command token and rejected
# with a clear error in _parse_svg_path -- rather than being silently
# dropped by the regex, which would misparse its numeric arguments as
# implicit-repeat coordinates of whatever command came before it.
_PATH_TOKEN_RE = re.compile(r"[A-Za-z]|-?\d+\.?\d*(?:[eE][-+]?\d+)?|-?\.\d+(?:[eE][-+]?\d+)?")
_PATH_ARITY = {"M": 2, "L": 2, "H": 1, "V": 1, "C": 6, "Q": 4, "Z": 0}


@dataclass
class _GradientDef:
    """A parsed ``<linearGradient>``/``<radialGradient>``, keyed by its
    ``id`` and looked up whenever an element's ``fill`` is ``url(#id)``."""

    stops: list[tuple[float, str]]  # (position 0.0-1.0, 6-hex-digit color)
    radial: bool
    angle_deg: float  # clockwise from horizontal-right; unused when radial


def _parse_number(value: str, attr_name: str) -> float:
    text = value.strip()
    if text.endswith("px"):
        text = text[:-2]
    try:
        return float(text)
    except ValueError:
        raise ValueError(
            f"layout: svg attribute {attr_name}={value!r} is not a plain number."
        ) from None


def _parse_color(value: str) -> str:
    text = value.strip()
    match = _HEX_COLOR_RE.match(text)
    if match:
        return match.group(1).upper()
    named = _NAMED_COLORS.get(text.lower())
    if named:
        return named
    raise ValueError(
        f"layout: svg color {value!r} is not a 6-hex-digit color or one of "
        f"this project's recognized CSS color names."
    )


def _svg_canvas_size(root: Any) -> tuple[float, float]:
    view_box = root.get("viewBox")
    if view_box:
        parts = view_box.replace(",", " ").split()
        if len(parts) == 4:
            return _parse_number(parts[2], "viewBox width"), _parse_number(
                parts[3], "viewBox height"
            )
    width = root.get("width")
    height = root.get("height")
    if width and height:
        return _parse_number(width, "width"), _parse_number(height, "height")
    raise ValueError(
        "layout: svg's root <svg> element needs a viewBox or width/height "
        "attribute to establish the coordinate scale."
    )


def _localname(element: Any) -> str:
    from lxml import etree

    return str(etree.QName(element).localname)


def _is_full_canvas_rect(rect: Any, svg_w: float, svg_h: float) -> bool:
    x = _parse_number(rect.get("x", "0"), "x")
    y = _parse_number(rect.get("y", "0"), "y")
    w = _parse_number(rect.get("width", "0"), "width")
    h = _parse_number(rect.get("height", "0"), "height")
    return x <= 0 and y <= 0 and w >= svg_w and h >= svg_h


def _fraction(element: Any, name: str, default: float) -> float:
    raw = element.get(name)
    if raw is None:
        return default
    raw = raw.strip()
    if raw.endswith("%"):
        return _parse_number(raw[:-1], name) / 100.0
    return _parse_number(raw, name)


def _linear_gradient_angle(element: Any) -> float:
    """DrawingML's own ``<a:lin ang>`` is a clockwise angle from
    horizontal-right in a y-down coordinate system -- exactly SVG's own
    convention too, so the gradient vector's angle converts directly with
    no axis flip needed."""
    x1, y1 = _fraction(element, "x1", 0.0), _fraction(element, "y1", 0.0)
    x2, y2 = _fraction(element, "x2", 1.0), _fraction(element, "y2", 0.0)
    dx, dy = x2 - x1, y2 - y1
    if dx == 0 and dy == 0:
        return 0.0
    return math.degrees(math.atan2(dy, dx)) % 360.0


def _parse_gradient_stops(gradient_element: Any) -> list[tuple[float, str]]:
    gradient_id = gradient_element.get("id")
    stops: list[tuple[float, str]] = []
    for stop in gradient_element:
        if _localname(stop) != "stop":
            continue
        offset_raw = (stop.get("offset") or "0").strip()
        position = (
            _parse_number(offset_raw[:-1], "offset") / 100.0
            if offset_raw.endswith("%")
            else _parse_number(offset_raw, "offset")
        )
        color_value = stop.get("stop-color")
        if not color_value:
            raise ValueError(
                f"layout: svg gradient id={gradient_id!r} has a <stop> with "
                "no stop-color attribute."
            )
        stops.append((position, _parse_color(color_value)))
    if len(stops) < 2:
        raise ValueError(
            f"layout: svg gradient id={gradient_id!r} needs at least two "
            "<stop> elements."
        )
    return stops


def _collect_gradients(root: Any) -> dict[str, _GradientDef]:
    """Full-tree scan (not just root's direct children) so gradients work
    whether or not the model wraps them in a <defs> element -- <defs> and
    its contents are never slide-visible on their own."""
    gradients: dict[str, _GradientDef] = {}
    for element in root.iter():
        tag = _localname(element)
        if tag not in ("linearGradient", "radialGradient"):
            continue
        gradient_id = element.get("id")
        if not gradient_id:
            continue
        stops = _parse_gradient_stops(element)
        gradients[gradient_id] = _GradientDef(
            stops=stops,
            radial=(tag == "radialGradient"),
            angle_deg=0.0 if tag == "radialGradient" else _linear_gradient_angle(element),
        )
    return gradients


def _apply_gradient_fill(fill_format: Any, gradient: _GradientDef) -> None:
    """Hand-builds the whole `<a:gradFill>` rather than using python-pptx's
    high-level gradient API: that API only supports exactly 2 stops
    (matching its own default-gradient template) and raises on radial
    gradients entirely (`gradient_angle` explicitly documents this). Calls
    `.solid()` first only to get a real fill child inserted at the correct
    schema position via python-pptx's own ordering logic, then swaps it
    for the hand-built `<a:gradFill>` in place."""
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    fill_format.solid()
    xPr = fill_format._xPr
    old_fill = xPr.find(qn("a:solidFill"))

    grad_fill = OxmlElement("a:gradFill")
    gs_lst = OxmlElement("a:gsLst")
    for position, hex_color in gradient.stops:
        gs = OxmlElement("a:gs")
        gs.set("pos", str(int(round(position * 100000))))
        srgb_clr = OxmlElement("a:srgbClr")
        srgb_clr.set("val", hex_color)
        gs.append(srgb_clr)
        gs_lst.append(gs)
    grad_fill.append(gs_lst)

    if gradient.radial:
        path = OxmlElement("a:path")
        path.set("path", "circle")
        fill_to_rect = OxmlElement("a:fillToRect")
        for attr in ("l", "t", "r", "b"):
            fill_to_rect.set(attr, "50000")
        path.append(fill_to_rect)
        grad_fill.append(path)
    else:
        lin = OxmlElement("a:lin")
        lin.set("ang", str(int(round(gradient.angle_deg * 60000)) % 21600000))
        lin.set("scaled", "1")
        grad_fill.append(lin)

    xPr.replace(old_fill, grad_fill)


def _apply_background(slide: Any, rect: Any, gradients: dict[str, _GradientDef]) -> None:
    from pptx.dml.color import RGBColor

    fill_value = rect.get("fill")
    if not fill_value:
        return
    fill_value = fill_value.strip()
    url_match = _URL_REF_RE.match(fill_value)
    if url_match:
        gradient = _lookup_gradient(url_match.group(1), gradients, "background <rect>")
        _apply_gradient_fill(slide.background.fill, gradient)
        return
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(  # type: ignore[no-untyped-call]
        _parse_color(fill_value)
    )


def _lookup_gradient(
    gradient_id: str, gradients: dict[str, _GradientDef], where: str
) -> _GradientDef:
    gradient = gradients.get(gradient_id)
    if gradient is None:
        raise ValueError(
            f"layout: svg {where} references fill=\"url(#{gradient_id})\", but no "
            f"<linearGradient>/<radialGradient> with id={gradient_id!r} was found "
            "in this SVG."
        )
    return gradient


def _apply_fill(shape: Any, element: Any, gradients: dict[str, _GradientDef]) -> None:
    from pptx.dml.color import RGBColor

    fill_value = element.get("fill")
    if not fill_value:
        raise ValueError(
            f"layout: svg <{_localname(element)}> needs a fill attribute -- "
            "an unfilled shape is invisible and would silently disappear."
        )
    fill_value = fill_value.strip()
    url_match = _URL_REF_RE.match(fill_value)
    if url_match:
        gradient = _lookup_gradient(
            url_match.group(1), gradients, f"<{_localname(element)}>"
        )
        _apply_gradient_fill(shape.fill, gradient)
        shape.line.fill.background()
        return
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(_parse_color(fill_value))  # type: ignore[no-untyped-call]
    shape.line.fill.background()


def _add_rect(
    slide: Any,
    element: Any,
    to_x: Callable[[float], int],
    to_y: Callable[[float], int],
    to_len: Callable[[float], int],
    gradients: dict[str, _GradientDef],
) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    x = _parse_number(element.get("x", "0"), "x")
    y = _parse_number(element.get("y", "0"), "y")
    w = _parse_number(element.get("width", "0"), "width")
    h = _parse_number(element.get("height", "0"), "height")
    if w <= 0 or h <= 0:
        raise ValueError("layout: svg <rect> needs positive width/height.")
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if element.get("rx") else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, to_x(x), to_y(y), to_len(w), to_len(h))
    _apply_fill(shape, element, gradients)


def _add_ellipse(
    slide: Any,
    element: Any,
    to_x: Callable[[float], int],
    to_y: Callable[[float], int],
    to_len: Callable[[float], int],
    gradients: dict[str, _GradientDef],
) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    if _localname(element) == "circle":
        cx = _parse_number(element.get("cx", "0"), "cx")
        cy = _parse_number(element.get("cy", "0"), "cy")
        r = _parse_number(element.get("r", "0"), "r")
        rx = ry = r
    else:
        cx = _parse_number(element.get("cx", "0"), "cx")
        cy = _parse_number(element.get("cy", "0"), "cy")
        rx = _parse_number(element.get("rx", "0"), "rx")
        ry = _parse_number(element.get("ry", "0"), "ry")
    if rx <= 0 or ry <= 0:
        raise ValueError(f"layout: svg <{_localname(element)}> needs a positive radius.")
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, to_x(cx - rx), to_y(cy - ry), to_len(2 * rx), to_len(2 * ry)
    )
    _apply_fill(shape, element, gradients)


def _add_text(
    slide: Any,
    element: Any,
    to_x: Callable[[float], int],
    to_y: Callable[[float], int],
    to_len: Callable[[float], int],
    to_pt: Callable[[float], float],
    slide_width: int,
) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    text = (element.text or "").strip()
    if not text:
        raise ValueError("layout: svg <text> needs non-empty text content.")
    x = _parse_number(element.get("x", "0"), "x")
    y = _parse_number(element.get("y", "0"), "y")
    font_size = _parse_number(element.get("font-size", "24"), "font-size")
    # SVG's (x, y) is the text baseline, not a box's top-left corner --
    # approximate the top of the box by backing off by roughly one
    # ascent's worth of the font size. Not pixel-exact SVG baseline
    # semantics, an acceptable v1 approximation, verified by live rendering
    # rather than assumed correct.
    box_top = y - font_size * 0.8
    box_left = to_x(x)
    # SVG <text> has no box width of its own, only a start point -- stretch
    # the box out to the slide's right edge.
    box_width = max(1, slide_width - box_left)

    textbox = slide.shapes.add_textbox(box_left, to_y(box_top), box_width, to_len(font_size * 1.4))
    text_frame = textbox.text_frame
    text_frame.word_wrap = True
    text_frame.margin_left = text_frame.margin_right = 0
    text_frame.margin_top = text_frame.margin_bottom = 0
    run = text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(to_pt(font_size))
    run.font.bold = element.get("font-weight", "").strip().lower() in ("bold", "700", "800", "900")
    fill_value = element.get("fill")
    if fill_value:
        run.font.color.rgb = RGBColor.from_string(_parse_color(fill_value))  # type: ignore[no-untyped-call]


def _tokenize_path(d: str) -> list[str]:
    return _PATH_TOKEN_RE.findall(d.strip())


def _parse_svg_path(d: str) -> list[tuple[str, bool, list[float]]]:
    """Returns a list of ``(command_letter_uppercase, is_relative,
    params)``. Implements SVG's implicit-repeat grammar: extra coordinate
    groups after a command letter (most commonly after ``M``/``L``) are
    treated as repeats of that command (an ``M``'s repeats are implicit
    ``L``s, per spec) without a new letter."""
    tokens = _tokenize_path(d)
    if not tokens:
        raise ValueError("layout: svg <path> needs a non-empty d attribute.")
    if not (tokens[0].isalpha() and tokens[0].upper() == "M"):
        raise ValueError(
            "layout: svg <path> data must start with a moveto (M or m) command."
        )

    commands: list[tuple[str, bool, list[float]]] = []
    i = 0
    current: str | None = None
    is_relative = False
    while i < len(tokens):
        tok = tokens[i]
        if tok.isalpha():
            if tok.upper() not in _PATH_ARITY:
                raise ValueError(
                    f"layout: svg <path> command {tok!r} is not supported -- only "
                    "M/L/H/V/C/Q/Z (upper- or lowercase) are. Use run_node_script "
                    "instead for arcs or shorthand curves."
                )
            current = tok.upper()
            is_relative = tok.islower()
            i += 1
        elif current is None:
            raise ValueError(
                "layout: svg <path> data must start with a moveto (M or m) command."
            )
        elif current == "Z":
            raise ValueError("layout: svg <path> has extra data after a Z (close) command.")
        else:
            current = "L" if current == "M" else current
        arity = _PATH_ARITY[current]
        if current == "Z":
            commands.append(("Z", is_relative, []))
            continue
        params: list[float] = []
        for _ in range(arity):
            if i >= len(tokens) or tokens[i].isalpha():
                raise ValueError(
                    f"layout: svg <path> command {current!r} needs {arity} number(s)."
                )
            params.append(_parse_number(tokens[i], f"path {current} parameter"))
            i += 1
        commands.append((current, is_relative, params))
    return commands


def _resolve_path_segments(
    commands: list[tuple[str, bool, list[float]]],
) -> tuple[list[tuple[str, list[tuple[float, float]]]], float, float, float, float]:
    """Resolves relative coordinates to absolute and returns a flat list of
    (kind, points) drawing segments plus the bounding box over *every*
    point involved, including bezier control points -- a safe overestimate
    of the visible shape's extent (a curve never leaves the convex hull of
    its control points), not a pixel-tight fit."""
    segments: list[tuple[str, list[tuple[float, float]]]] = []
    cx = cy = 0.0
    sx = sy = 0.0
    min_x = min_y = float("inf")
    max_x = max_y = float("-inf")

    def track(x: float, y: float) -> None:
        nonlocal min_x, min_y, max_x, max_y
        min_x, max_x = min(min_x, x), max(max_x, x)
        min_y, max_y = min(min_y, y), max(max_y, y)

    for letter, relative, params in commands:
        if letter == "M":
            x, y = params
            if relative:
                x, y = x + cx, y + cy
            cx, cy, sx, sy = x, y, x, y
            track(x, y)
            segments.append(("move", [(x, y)]))
        elif letter == "L":
            x, y = params
            if relative:
                x, y = x + cx, y + cy
            cx, cy = x, y
            track(x, y)
            segments.append(("line", [(x, y)]))
        elif letter == "H":
            (x,) = params
            x = x + cx if relative else x
            cx = x
            track(cx, cy)
            segments.append(("line", [(cx, cy)]))
        elif letter == "V":
            (y,) = params
            y = y + cy if relative else y
            cy = y
            track(cx, cy)
            segments.append(("line", [(cx, cy)]))
        elif letter == "C":
            x1, y1, x2, y2, x, y = params
            if relative:
                x1, y1, x2, y2, x, y = x1 + cx, y1 + cy, x2 + cx, y2 + cy, x + cx, y + cy
            track(x1, y1)
            track(x2, y2)
            track(x, y)
            segments.append(("cubic", [(x1, y1), (x2, y2), (x, y)]))
            cx, cy = x, y
        elif letter == "Q":
            x1, y1, x, y = params
            if relative:
                x1, y1, x, y = x1 + cx, y1 + cy, x + cx, y + cy
            track(x1, y1)
            track(x, y)
            segments.append(("quad", [(x1, y1), (x, y)]))
            cx, cy = x, y
        else:  # "Z"
            track(sx, sy)
            segments.append(("close", []))
            cx, cy = sx, sy

    if min_x == float("inf"):
        raise ValueError("layout: svg <path> produced no points.")
    return segments, min_x, min_y, max_x, max_y


def _add_path(
    slide: Any,
    element: Any,
    to_x: Callable[[float], int],
    to_y: Callable[[float], int],
    to_len: Callable[[float], int],
    gradients: dict[str, _GradientDef],
) -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    d_attr = element.get("d")
    if not d_attr:
        raise ValueError("layout: svg <path> needs a d attribute.")
    segments, min_x, min_y, max_x, max_y = _resolve_path_segments(_parse_svg_path(d_attr))
    if not any(kind in ("line", "cubic", "quad") for kind, _ in segments):
        raise ValueError(
            "layout: svg <path> needs at least one drawing command (L/H/V/C/Q) "
            "after its initial moveto -- a bare moveto has no visible shape."
        )

    bbox_w = max(max_x - min_x, 0.0)
    bbox_h = max(max_y - min_y, 0.0)
    emu_left, emu_top = to_x(min_x), to_y(min_y)
    emu_width = max(to_len(bbox_w), 1)
    emu_height = max(to_len(bbox_h), 1)
    path_w = max(round(bbox_w * _PATH_UNIT_SCALE), 1)
    path_h = max(round(bbox_h * _PATH_UNIT_SCALE), 1)

    def local(x: float, y: float) -> tuple[int, int]:
        return round((x - min_x) * _PATH_UNIT_SCALE), round((y - min_y) * _PATH_UNIT_SCALE)

    sp = slide.shapes._spTree.add_freeform_sp(emu_left, emu_top, emu_width, emu_height)
    path_el = sp.add_path(w=path_w, h=path_h)
    for kind, points in segments:
        if kind == "move":
            path_el.add_moveTo(*local(*points[0]))
        elif kind == "line":
            path_el.add_lnTo(*local(*points[0]))
        elif kind == "close":
            path_el.add_close()
        elif kind in ("cubic", "quad"):
            # python-pptx's FreeformBuilder only supports straight line
            # segments -- <a:cubicBezTo>/<a:quadBezTo> have no public API,
            # hand-appended here the same way add_pptx_scrim hand-appends
            # <a:alpha> elsewhere in this package. Both element types have
            # an empty `successors` list in python-pptx's own oxml schema
            # (verified by reading CT_Path2D's ZeroOrMore declarations), so
            # a plain append -- matching what add_moveTo/add_lnTo/add_close
            # already do -- lands correctly in document (drawing) order.
            tag = "a:cubicBezTo" if kind == "cubic" else "a:quadBezTo"
            bez_el = OxmlElement(tag)
            for x, y in points:
                pt = OxmlElement("a:pt")
                lx, ly = local(x, y)
                pt.set("x", str(lx))
                pt.set("y", str(ly))
                bez_el.append(pt)
            path_el.append(bez_el)

    shape = slide.shapes._shape_factory(sp)
    _apply_fill(shape, element, gradients)


def add_svg_slide(prs: PresentationType, svg_markup: str) -> None:
    """Parse ``svg_markup`` (one complete ``<svg>...</svg>`` document) and
    append one new slide built from real, native shapes -- the entry point
    for ``write_pptx``'s ``layout: svg`` directive."""
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
    try:
        root = etree.fromstring(svg_markup.encode("utf-8"), parser=parser)
    except etree.XMLSyntaxError as exc:
        raise ValueError(f"layout: svg content is not valid XML: {exc}") from None

    if _localname(root) != "svg":
        raise ValueError(
            f"layout: svg content must be a single <svg> root element, got "
            f"<{_localname(root)}>."
        )

    svg_w, svg_h = _svg_canvas_size(root)
    if svg_w <= 0 or svg_h <= 0:
        raise ValueError("layout: svg's canvas width/height must be positive.")

    gradients = _collect_gradients(root)

    slide_width = prs.slide_width or _DEFAULT_SLIDE_WIDTH_EMU
    slide_height = prs.slide_height or _DEFAULT_SLIDE_HEIGHT_EMU
    scale = min(slide_width / svg_w, slide_height / svg_h)
    offset_x = (slide_width - svg_w * scale) / 2
    offset_y = (slide_height - svg_h * scale) / 2

    def to_emu_x(x: float) -> int:
        return round(offset_x + x * scale)

    def to_emu_y(y: float) -> int:
        return round(offset_y + y * scale)

    def to_emu_len(length: float) -> int:
        return round(length * scale)

    def to_pt(length: float) -> float:
        return length * scale / _EMU_PER_POINT

    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT])

    children = list(root)
    background_index = None
    # A gradient definition (bare or <defs>-wrapped) never renders on its
    # own, so it doesn't count as "the first child" when deciding whether
    # that first real, visible child is a full-canvas rect eligible for
    # background promotion.
    _definitional_tags = ("defs", "linearGradient", "radialGradient")
    first_visible_index = next(
        (i for i, c in enumerate(children) if _localname(c) not in _definitional_tags), None
    )
    if (
        first_visible_index is not None
        and _localname(children[first_visible_index]) == "rect"
        and _is_full_canvas_rect(children[first_visible_index], svg_w, svg_h)
    ):
        background_index = first_visible_index

    for index, child in enumerate(children):
        tag = _localname(child)
        if index == background_index:
            _apply_background(slide, child, gradients)
            continue
        if tag in ("defs", "linearGradient", "radialGradient"):
            # Never slide-visible on their own -- already collected above.
            continue
        if tag == "rect":
            _add_rect(slide, child, to_emu_x, to_emu_y, to_emu_len, gradients)
        elif tag in ("circle", "ellipse"):
            _add_ellipse(slide, child, to_emu_x, to_emu_y, to_emu_len, gradients)
        elif tag == "text":
            _add_text(slide, child, to_emu_x, to_emu_y, to_emu_len, to_pt, slide_width)
        elif tag == "path":
            _add_path(slide, child, to_emu_x, to_emu_y, to_emu_len, gradients)
        else:
            raise ValueError(
                f"layout: svg does not support <{tag}> -- only <rect>, "
                "<circle>, <ellipse>, <text>, and <path> are supported "
                "(plus <linearGradient>/<radialGradient> for fill="
                "\"url(#id)\"). Use run_node_script instead for anything "
                "needing images, groups, arcs, or shorthand curves."
            )
