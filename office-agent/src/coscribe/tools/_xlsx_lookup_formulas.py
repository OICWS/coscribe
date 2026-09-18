"""Native, coscribe-owned evaluation of `XLOOKUP`/`XMATCH` -- the two of
`spreadsheets.py`'s six blocked "LibreOffice can't evaluate this" functions
that are commonly used in their plain *scalar* form (one lookup value, one
result), as opposed to `FILTER`/`UNIQUE`/`SORT`/`SEQUENCE`, which are
"dynamic array" functions that spill a result across multiple cells --
real Excel-365-only OOXML machinery (a `t="array"` `<f>`, cell metadata
referencing a new `xl/metadata.xml` part) that stays out of scope here
deliberately: confirmed via research that a safe, real reference
implementation exists to build that from later (`jmcnamara/XlsxWriter`,
BSD-2-Clause), but it's a separate, larger piece of work than this module.

**Why this exists at all, instead of just unblocking these two functions
and trusting `_recalc_xlsx`'s LibreOffice pass**: confirmed empirically
(not assumed) that LibreOffice, given `=_xlfn.XLOOKUP(...)`, doesn't
crash or corrupt anything -- it lowercases the function name in the
stored formula text and writes the cell as a real OOXML error
(`t="e"`, `<v>#NAME?</v>`), while correctly recalculating every *other*
formula in the same file. So the gap is real and precisely scoped: these
two functions specifically, nothing else.

**Why direct XML patching, not something through openpyxl's own API**:
confirmed by reading `openpyxl.cell._writer.etree_write_cell`'s actual
source -- for *any* formula cell (not just these two functions), the
writer emits `<f>` with the formula text and then unconditionally
discards the value before ever reaching the `<v>`-writing branch. There
is no openpyxl API, public or private, that produces a formula cell with
a real cached value. This is exactly why `_recalc_xlsx` shells out to
LibreOffice for every other formula instead of trying to compute and
write cached values through openpyxl directly -- this module does the
same "compute it ourselves, then patch the saved XML" thing, just for
the two functions LibreOffice itself can't compute.

**Deliberately scoped, not a formula-language parser**: only recognizes
a formula whose *entire* body is one top-level call to XLOOKUP or XMATCH
(optionally `_xlfn.`/`_xlfn._xlws.`-prefixed, case-insensitive -- matches
LibreOffice's own lowercased output too). `=1+XLOOKUP(...)` or
`=XLOOKUP(...)+1` are left alone. Each argument must be a literal
(a quoted string, number, or TRUE/FALSE) or a plain, optionally
sheet-qualified cell/range reference (`A2`, `A2:A10`, `A:A`,
`'My Sheet'!B:B`) -- an argument that's itself an expression (another
function call, an operator) makes the whole formula skipped rather than
guessed at. None of this is a regression versus today: a formula this
module declines to evaluate is left exactly as LibreOffice's own pass
already left it (an OOXML `#NAME?` error cell), not silently wrong.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, cast

_XLFN_PREFIXES = ("_XLFN._XLWS.", "_XLFN.")
_SUPPORTED_FUNCTIONS = frozenset({"XLOOKUP", "XMATCH"})

_CALL_RE = re.compile(r"^([A-Za-z_][A-Za-z0-9_.]*)\((.*)\)$", re.DOTALL)

_REF_RE = re.compile(
    r"^(?:'(?P<qsheet>[^']+)'!|(?P<sheet>[A-Za-z_][A-Za-z0-9_]*)!)?"
    r"(?P<ref>"
    r"\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?"  # A2 or A2:B10
    r"|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3}"  # A:A whole column(s)
    r"|\$?\d+:\$?\d+"  # 2:2 whole row(s)
    r")$"
)

_SML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


class FormulaError:
    """A real Excel error result (`#N/A`, `#VALUE!`, ...) -- distinct from a
    plain string so `_format_cell_value` writes `t="e"`, not `t="str"`."""

    __slots__ = ("code",)

    def __init__(self, code: str) -> None:
        self.code = code

    def __repr__(self) -> str:
        return f"FormulaError({self.code!r})"


class _Sentinel:
    __slots__ = ("label",)

    def __init__(self, label: str) -> None:
        self.label = label

    def __repr__(self) -> str:
        return f"<{self.label}>"


_NO_MATCH = _Sentinel("no-match")
_NOT_GIVEN = _Sentinel("not-given")
_UNRESOLVABLE = _Sentinel("unresolvable")


# ── Formula-text parsing ───────────────────────────────────────────────────


def _split_top_level_args(text: str) -> list[str]:
    """Split a comma-separated Excel argument list at top-level commas only
    -- respecting nested parens and double-quoted strings (Excel escapes an
    embedded quote inside a string as `""`)."""
    if not text.strip():
        return []
    args: list[str] = []
    current: list[str] = []
    depth = 0
    in_quotes = False
    i = 0
    n = len(text)
    while i < n:
        ch = text[i]
        if in_quotes:
            current.append(ch)
            if ch == '"':
                if i + 1 < n and text[i + 1] == '"':
                    current.append('"')
                    i += 2
                    continue
                in_quotes = False
            i += 1
            continue
        if ch == '"':
            in_quotes = True
            current.append(ch)
        elif ch == "(":
            depth += 1
            current.append(ch)
        elif ch == ")":
            depth -= 1
            current.append(ch)
        elif ch == "," and depth == 0:
            args.append("".join(current))
            current = []
        else:
            current.append(ch)
        i += 1
    args.append("".join(current))
    return args


def parse_single_call(formula_text: str) -> tuple[str, list[str]] | None:
    """`formula_text` includes the leading `=`. Returns `(function_name,
    [raw_arg_strings])` if the whole formula is one top-level call to a
    supported function, else `None`."""
    if not formula_text.startswith("="):
        return None
    body = formula_text[1:].strip()
    upper = body.upper()
    for prefix in _XLFN_PREFIXES:
        if upper.startswith(prefix):
            body = body[len(prefix) :]
            break
    match = _CALL_RE.match(body)
    if not match:
        return None
    name = match.group(1).upper()
    if name not in _SUPPORTED_FUNCTIONS:
        return None
    args = [a.strip() for a in _split_top_level_args(match.group(2))]
    return name, args


# ── Argument resolution ─────────────────────────────────────────────────────


def _parse_literal(text: str) -> Any:
    """Returns the literal's Python value, or `_UNRESOLVABLE` if `text`
    isn't a plain literal (quoted string / number / TRUE / FALSE)."""
    if len(text) >= 2 and text[0] == '"' and text[-1] == '"':
        return text[1:-1].replace('""', '"')
    if text.upper() == "TRUE":
        return True
    if text.upper() == "FALSE":
        return False
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        pass
    return _UNRESOLVABLE


def _resolve_reference(
    ref_text: str, current_sheet: str, values_workbook: Any
) -> list[Any] | None:
    """Resolve a (possibly sheet-qualified) cell/range reference against
    `values_workbook` (an already-`data_only=True`-loaded workbook, so a
    formula cell in range reads back its LibreOffice-recalculated cached
    value). Returns a flat list of cell values in reading order, or `None`
    if the reference doesn't parse or names an unknown sheet."""
    from openpyxl.utils import column_index_from_string

    match = _REF_RE.match(ref_text)
    if not match:
        return None
    sheet_name = match.group("qsheet") or match.group("sheet") or current_sheet
    if sheet_name not in values_workbook.sheetnames:
        return None
    sheet = values_workbook[sheet_name]
    ref = (match.group("ref") or "").replace("$", "")

    if ":" not in ref:
        column_letters = "".join(c for c in ref if c.isalpha())
        row_digits = "".join(c for c in ref if c.isdigit())
        col = column_index_from_string(column_letters)
        row = int(row_digits)
        return [sheet.cell(row=row, column=col).value]

    left, right = ref.split(":")
    min_col: int
    min_row: int
    max_col: int
    max_row: int
    if left.isalpha() and right.isalpha():
        min_col, max_col = column_index_from_string(left), column_index_from_string(right)
        min_row, max_row = 1, sheet.max_row
    elif left.isdigit() and right.isdigit():
        min_row, max_row = int(left), int(right)
        min_col, max_col = 1, sheet.max_column
    else:
        from openpyxl.utils.cell import range_boundaries

        raw_min_col, raw_min_row, raw_max_col, raw_max_row = range_boundaries(ref)
        if raw_min_col is None or raw_min_row is None or raw_max_col is None or raw_max_row is None:
            return None
        min_col, min_row, max_col, max_row = raw_min_col, raw_min_row, raw_max_col, raw_max_row

    values: list[Any] = []
    for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        values.extend(cell.value for cell in row)
    return values


def _resolve_arg(arg_text: str, current_sheet: str, values_workbook: Any) -> Any:
    """Returns a plain scalar, a `list` (for a multi-cell range), or
    `_UNRESOLVABLE` if `arg_text` is neither a plain literal nor a plain
    reference this module understands (e.g. a nested function call or an
    operator expression)."""
    if _REF_RE.match(arg_text):
        resolved = _resolve_reference(arg_text, current_sheet, values_workbook)
        if resolved is None:
            return _UNRESOLVABLE
        return resolved[0] if len(resolved) == 1 else resolved
    return _parse_literal(arg_text)


# ── Excel comparison semantics ──────────────────────────────────────────────


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _excel_eq(a: Any, b: Any) -> bool:
    """Excel's default equality: text compared case-insensitively; a
    number and its string form do NOT match (`42 <> "42"` in a lookup,
    real Excel behavior); booleans only match booleans."""
    if isinstance(a, bool) or isinstance(b, bool):
        return isinstance(a, bool) and isinstance(b, bool) and a == b
    if _is_number(a) and _is_number(b):
        return bool(a == b)
    if isinstance(a, str) and isinstance(b, str):
        return a.casefold() == b.casefold()
    return False


def _wildcard_match(value: Any, pattern: str) -> bool:
    """`*` = any run of characters, `?` = exactly one character -- Excel's
    `~`-escape for a literal `*`/`?` is not supported (a documented, rare
    edge case, not silently wrong: an unescaped wildcard in the pattern is
    simply treated as a wildcard, matching Excel's own behavior for that
    much at least)."""
    if not isinstance(value, str):
        return False
    escaped = re.escape(pattern).replace(r"\*", ".*").replace(r"\?", ".")
    return re.fullmatch(escaped, value, re.IGNORECASE) is not None


def xmatch_index(
    lookup_value: Any, array: list[Any], match_mode: int, search_mode: int
) -> int | _Sentinel:
    """Returns the 0-based index of `lookup_value` in `array`, or
    `_NO_MATCH`. `search_mode` 2/-2 (Excel's binary-search modes) are
    executed as a linear scan in that same direction -- identical result
    to a real binary search for an exact match_mode; Excel's own docs
    already describe the binary modes as needing pre-sorted data and
    otherwise being unreliable, so this gives up only unneeded speed on
    very large ranges, not correctness."""
    order = range(len(array)) if search_mode in (1, 2) else range(len(array) - 1, -1, -1)

    if match_mode == 2:
        if not isinstance(lookup_value, str):
            return _NO_MATCH
        for i in order:
            if _wildcard_match(array[i], lookup_value):
                return i
        return _NO_MATCH

    if match_mode == 0 or not _is_number(lookup_value):
        for i in order:
            if _excel_eq(array[i], lookup_value):
                return i
        return _NO_MATCH

    # match_mode -1 (exact or next smaller) / 1 (exact or next larger).
    best_index: int | None = None
    best_value: float | None = None
    for i in order:
        value = array[i]
        if not _is_number(value):
            continue
        if value == lookup_value:
            return i
        if match_mode == -1 and value < lookup_value:
            if best_value is None or value > best_value:
                best_index, best_value = i, value
        elif match_mode == 1 and value > lookup_value:
            if best_value is None or value < best_value:
                best_index, best_value = i, value
    return best_index if best_index is not None else _NO_MATCH


def _as_int_arg(value: Any, default: int) -> int | None:
    if value is _NOT_GIVEN:
        return default
    if isinstance(value, bool) or not _is_number(value):
        return None
    return int(value)


def evaluate_xmatch(args: list[Any]) -> Any:
    """`args`: already-resolved positional values for
    `XMATCH(lookup_value, lookup_array, [match_mode], [search_mode])`.
    Returns a 1-based match position (Excel convention), or a
    `FormulaError`."""
    lookup_value, lookup_array = args[0], args[1]
    match_mode = _as_int_arg(args[2] if len(args) > 2 else _NOT_GIVEN, 0)
    search_mode = _as_int_arg(args[3] if len(args) > 3 else _NOT_GIVEN, 1)
    if not isinstance(lookup_array, list):
        return FormulaError("#VALUE!")
    if match_mode not in (-1, 0, 1, 2) or search_mode not in (-2, -1, 1, 2):
        return FormulaError("#VALUE!")
    index = xmatch_index(lookup_value, lookup_array, match_mode, search_mode)
    if index is _NO_MATCH:
        return FormulaError("#N/A")
    return cast(int, index) + 1


def evaluate_xlookup(args: list[Any]) -> Any:
    """`args`: already-resolved positional values for `XLOOKUP(lookup_value,
    lookup_array, return_array, [if_not_found], [match_mode],
    [search_mode])`. Returns the matched entry from `return_array`, or
    `if_not_found` / a `FormulaError`."""
    lookup_value, lookup_array, return_array = args[0], args[1], args[2]
    if_not_found = args[3] if len(args) > 3 else _NOT_GIVEN
    match_mode = _as_int_arg(args[4] if len(args) > 4 else _NOT_GIVEN, 0)
    search_mode = _as_int_arg(args[5] if len(args) > 5 else _NOT_GIVEN, 1)
    if not isinstance(lookup_array, list) or not isinstance(return_array, list):
        return FormulaError("#VALUE!")
    if len(lookup_array) != len(return_array):
        return FormulaError("#VALUE!")
    if match_mode not in (-1, 0, 1, 2) or search_mode not in (-2, -1, 1, 2):
        return FormulaError("#VALUE!")
    index = xmatch_index(lookup_value, lookup_array, match_mode, search_mode)
    if index is _NO_MATCH:
        if if_not_found is not _NOT_GIVEN:
            return if_not_found
        return FormulaError("#N/A")
    return return_array[cast(int, index)]


_EVALUATORS = {"XMATCH": evaluate_xmatch, "XLOOKUP": evaluate_xlookup}
_MIN_ARGS = {"XMATCH": 2, "XLOOKUP": 3}
_MAX_ARGS = {"XMATCH": 4, "XLOOKUP": 6}


def evaluate_formula(formula_text: str, current_sheet: str, values_workbook: Any) -> Any:
    """Top-level entry point for one cell's formula text. Returns the
    computed value/`FormulaError`, or `_UNRESOLVABLE` if this formula isn't
    a supported single-call XLOOKUP/XMATCH (the caller leaves such cells
    untouched -- same as today)."""
    parsed = parse_single_call(formula_text)
    if parsed is None:
        return _UNRESOLVABLE
    name, raw_args = parsed
    if not (_MIN_ARGS[name] <= len(raw_args) <= _MAX_ARGS[name]):
        return _UNRESOLVABLE
    resolved_args = [_resolve_arg(a, current_sheet, values_workbook) for a in raw_args]
    if any(a is _UNRESOLVABLE for a in resolved_args):
        return _UNRESOLVABLE
    return _EVALUATORS[name](resolved_args)


# ── XML patching ─────────────────────────────────────────────────────────


def _format_cell_value(result: Any) -> tuple[str, str | None]:
    """Returns `(text, t_attribute_or_None)` matching the exact convention
    a real LibreOffice-recalculated cell already uses (confirmed by reading
    a recalculated file's raw XML directly): `t="e"` for an error, `t="b"`
    for a boolean (`"1"`/`"0"`), `t="n"` for a number, `t="str"` for an
    inline-formula-result string (not `t="s"`, which indexes the shared-
    string table -- a formula's string result is never shared-string-typed
    in real Excel/LibreOffice output)."""
    if isinstance(result, FormulaError):
        return result.code, "e"
    if isinstance(result, bool):
        return ("1" if result else "0"), "b"
    if _is_number(result):
        if isinstance(result, float) and result.is_integer():
            return str(int(result)), "n"
        return str(result), "n"
    return str(result), "str"


def _sheet_archive_paths(entries: dict[str, bytes]) -> dict[str, str]:
    """`{sheet_name: "xl/worksheets/sheetN.xml"}`, resolved the correct way
    (via `xl/workbook.xml`'s `<sheet r:id=...>` + `xl/_rels/workbook.xml.rels`'s
    `<Relationship Id=... Target=...>`) rather than assuming sheet order
    matches file-naming order, which the OOXML spec never guarantees and a
    file that's been through other edits genuinely might not preserve."""
    from lxml import etree

    workbook_root = etree.fromstring(entries["xl/workbook.xml"])
    rels_root = etree.fromstring(entries["xl/_rels/workbook.xml.rels"])
    rel_targets = {
        rel.get("Id"): rel.get("Target")
        for rel in rels_root.iter(f"{{{_PKG_RELS_NS}}}Relationship")
    }
    paths: dict[str, str] = {}
    for sheet_el in workbook_root.iter(f"{{{_SML_NS}}}sheet"):
        name = sheet_el.get("name")
        rel_id = sheet_el.get(f"{{{_R_NS}}}id")
        target = rel_targets.get(rel_id) if rel_id else None
        if not (name and target):
            continue
        # A relationship Target is either absolute-from-package-root
        # (leading "/", e.g. openpyxl's own "/xl/worksheets/sheet1.xml" --
        # confirmed by reading a real openpyxl-saved file's rels, not
        # assumed) or relative to xl/ (the directory containing
        # xl/workbook.xml, the part that owns this .rels file) -- both are
        # legal OPC forms, real files use either.
        paths[name] = target[1:] if target.startswith("/") else f"xl/{target}"
    return paths


def _patch_cell(sheet_root: Any, address: str, result: Any) -> bool:
    """Sets/replaces the cached `<v>` (and `t=` type) on `<c r="address">`
    inside `sheet_root`, right after its existing `<f>` child (ECMA-376's
    required child order). Returns whether the cell/formula was found."""
    from lxml import etree

    for cell_el in sheet_root.iter(f"{{{_SML_NS}}}c"):
        if cell_el.get("r") != address:
            continue
        formula_el = cell_el.find(f"{{{_SML_NS}}}f")
        if formula_el is None:
            return False
        text, type_attr = _format_cell_value(result)
        if type_attr:
            cell_el.set("t", type_attr)
        elif "t" in cell_el.attrib:
            del cell_el.attrib["t"]
        value_el = cell_el.find(f"{{{_SML_NS}}}v")
        if value_el is None:
            value_el = etree.Element(f"{{{_SML_NS}}}v")
            formula_el.addnext(value_el)
        value_el.text = text
        return True
    return False


def evaluate_and_patch_lookup_formulas(file_path: Path) -> dict[str, object]:
    """Run this **after** `_recalc_xlsx`'s LibreOffice pass, on the same
    file -- this module's own evaluation reads other cells' *already-
    recalculated* cached values (needed when a lookup/return range itself
    contains ordinary formulas), and LibreOffice's own recalc of every
    other formula is unaffected by (and unrelated to) what this function
    does to the two it can't compute.

    Returns `{"formulas_evaluated": int, "results": {"Sheet!A1": <str(result)>}}`.
    A cell whose formula isn't a supported single-call XLOOKUP/XMATCH is
    left completely untouched -- not scanned as an error, not modified.
    """
    from openpyxl import load_workbook

    values_workbook = load_workbook(str(file_path), data_only=True)
    formula_workbook = load_workbook(str(file_path), data_only=False)

    matches: dict[str, dict[str, Any]] = {}
    for sheet_name in formula_workbook.sheetnames:
        sheet = formula_workbook[sheet_name]
        if not hasattr(sheet, "iter_rows"):
            continue
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type != "f" or not isinstance(cell.value, str):
                    continue
                result = evaluate_formula(cell.value, sheet_name, values_workbook)
                if result is _UNRESOLVABLE:
                    continue
                matches.setdefault(sheet_name, {})[cell.coordinate] = result
    values_workbook.close()
    formula_workbook.close()

    if not matches:
        return {"formulas_evaluated": 0, "results": {}}

    import zipfile

    from lxml import etree

    with zipfile.ZipFile(file_path, "r") as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}

    sheet_paths = _sheet_archive_paths(entries)
    results: dict[str, str] = {}
    for sheet_name, cells in matches.items():
        archive_path = sheet_paths.get(sheet_name)
        if archive_path is None or archive_path not in entries:
            continue
        sheet_root = etree.fromstring(entries[archive_path])
        for address, result in cells.items():
            if _patch_cell(sheet_root, address, result):
                results[f"{sheet_name}!{address}"] = (
                    result.code if isinstance(result, FormulaError) else str(result)
                )
        entries[archive_path] = etree.tostring(
            sheet_root, xml_declaration=True, encoding="UTF-8", standalone=True
        )

    with zipfile.ZipFile(file_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)

    return {"formulas_evaluated": len(results), "results": results}
