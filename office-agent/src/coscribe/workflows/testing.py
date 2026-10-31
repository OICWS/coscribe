"""A test run of a workflow: the real steps, once, reported step by step --
how the assistant checks a workflow it drafted or revised before the user
reviews it, and reproduces a saved one's failure.

It runs on a throwaway checkpointer, so nothing lands in a task's run
history, and stops before an approval step: a test has no one to ask, and
what comes after an approval is exactly what the approval guards. It also
stops before a step marked irreversible unless told to send for real
(decision 0029): checking a fix must not file a claim or pay an invoice."""

from __future__ import annotations

import contextlib
import csv
import json
import time
import uuid
from collections.abc import Awaitable, Callable
from datetime import datetime
from pathlib import Path
from typing import Any

from langgraph.checkpoint.memory import InMemorySaver

from ..tools import browser_run
from ..workflow_spec import (
    AgentStep,
    ApprovalStep,
    BranchStep,
    LoopStep,
    Step,
    ToolStep,
    Workflow,
    walk,
)
from .engine import StepContext, StepRecord, WorkflowNotRunnable, WorkflowRun, preview

_OUTPUT_CHARS = 400
_MAX_FILES = 10
# Where a tool's result names a file it wrote, moved or downloaded.
_FILE_KEYS = frozenset({"path", "file", "files", "destination"})


def uses_browser(workflow: Workflow) -> bool:
    return any(
        isinstance(placed.step, ToolStep) and placed.step.tool.startswith("browser_")
        for placed in walk(workflow.steps)
    )


def _seconds(record: StepRecord) -> float | None:
    if not record.started_at or not record.finished_at:
        return None
    try:
        start = datetime.fromisoformat(record.started_at)
        end = datetime.fromisoformat(record.finished_at)
    except ValueError:
        return None
    return round((end - start).total_seconds(), 1)


def _clip(value: Any) -> Any:
    shown = preview(value)
    text = shown if isinstance(shown, str) else json.dumps(shown, ensure_ascii=False, default=str)
    return text if len(text) <= _OUTPUT_CHARS else f"{text[:_OUTPUT_CHARS]}…"


def _named_files(value: Any) -> list[str]:
    found: list[str] = []
    if isinstance(value, dict):
        for key, item in value.items():
            if key in _FILE_KEYS and isinstance(item, str):
                found.append(item)
            elif key in _FILE_KEYS and isinstance(item, list):
                found += [x for x in item if isinstance(x, str)]
            else:
                found += _named_files(item)
    elif isinstance(value, list):
        for item in value:
            found += _named_files(item)
    return found


def _table_shape(path: Path) -> dict[str, int]:
    """Rows and columns of a spreadsheet's first sheet, or {} if it can't be read."""
    try:
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True)
            try:
                sheet = book.worksheets[0]
                return {"rows": sheet.max_row or 0, "columns": sheet.max_column or 0}
            finally:
                book.close()
        if path.suffix.lower() == ".csv":
            with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
                rows = list(csv.reader(handle))
            return {"rows": len(rows), "columns": max((len(r) for r in rows), default=0)}
    except Exception:  # noqa: BLE001 -- a file that isn't a real table just shows no shape
        return {}
    return {}


def output_files(outputs: list[Any], root: Path, since: float) -> list[dict[str, Any]]:
    """Files the steps' results (`outputs`) name that were written during
    the run -- what a reviewer checks first."""
    files: list[dict[str, Any]] = []
    seen: set[Path] = set()
    for output in outputs:
        for name in _named_files(output):
            path = (root / name).resolve() if not Path(name).is_absolute() else Path(name)
            if path in seen or not path.is_file() or path.stat().st_mtime < since:
                continue
            seen.add(path)
            try:
                shown = str(path.relative_to(root.resolve()))
            except ValueError:
                shown = str(path)
            files.append({"path": shown, "size": path.stat().st_size, **_table_shape(path)})
            if len(files) == _MAX_FILES:
                return files
    return files


_GUARD_PREFIX = "before_"


def _guarded(workflow: Workflow) -> tuple[Workflow, dict[str, Step]]:
    """The workflow with an approval step in front of every irreversible
    one -- the engine already stops a test there -- and which guard stands
    for which step."""
    taken = {placed.step.id for placed in walk(workflow.steps)}
    guards: dict[str, Step] = {}

    def guard(step: Step) -> ApprovalStep:
        gid = _GUARD_PREFIX + step.id
        while gid in taken:
            gid += "_"
        taken.add(gid)
        guards[gid] = step
        return ApprovalStep(
            id=gid,
            title=f"Before {step.title}"[:120],
            message=f"{step.title} can't be taken back.",
        )

    def rebuild(steps: list[Step]) -> list[Step]:
        out: list[Step] = []
        for step in steps:
            if isinstance(step, LoopStep):
                step = step.model_copy(update={"steps": rebuild(step.steps)})
            elif isinstance(step, BranchStep):
                step = step.model_copy(
                    update={"then": rebuild(step.then), "otherwise": rebuild(step.otherwise)}
                )
            elif isinstance(step, (ToolStep, AgentStep)) and step.irreversible:
                out.append(guard(step))
            out.append(step)
        return out

    return workflow.model_copy(update={"steps": rebuild(workflow.steps)}), guards


async def run_test(
    workflow: Workflow,
    ctx: StepContext,
    inputs: dict[str, Any] | None = None,
    on_step: Callable[[StepRecord], Awaitable[None]] | None = None,
    thread_id: str = "",
    *,
    send_for_real: bool = False,
) -> dict[str, Any]:
    """{"status": "passed" | "failed" | "stopped_at_approval" |
    "stopped_before_irreversible", "seconds", "steps": [{"id", "title",
    "status", "seconds", "output", "error"}], "failed_step", "error"} --
    every step listed, "not reached" for those the run never got to; with
    "stopped_before" ({"id", "title"}) when an irreversible step stopped it.
    `send_for_real` runs irreversible steps too."""
    records: dict[str, StepRecord] = {}
    shown = workflow
    guards: dict[str, Step] = {}
    if not send_for_real:
        workflow, guards = _guarded(workflow)

    async def record_step(record: StepRecord) -> None:
        # Inside a loop the last pass stands for the step.
        records[record.step_id] = record
        if on_step is not None:
            await on_step(record)

    run = WorkflowRun(workflow, ctx, InMemorySaver(), f"test-{uuid.uuid4().hex[:8]}", record_step)
    started = time.monotonic()
    # Whole seconds: some file systems keep coarse modification times.
    started_at = int(time.time())
    lock = browser_run(thread_id) if uses_browser(workflow) else contextlib.nullcontext()
    async with lock:
        try:
            outcome = await run.start(inputs or {})
        except WorkflowNotRunnable as exc:
            return {
                "status": "failed",
                "seconds": 0.0,
                "steps": [],
                "files": [],
                "failed_step": None,
                "error": str(exc),
            }
    status = {"completed": "passed", "waiting": "stopped_at_approval"}.get(outcome.status, "failed")
    stopped_before = guards.get(outcome.step_id or "") if status == "stopped_at_approval" else None
    if stopped_before is not None:
        status = "stopped_before_irreversible"
    steps = []
    for placed in walk(shown.steps):
        record = records.get(placed.step.id)
        step: dict[str, Any] = {
            "id": placed.step.id,
            "title": placed.step.title,
            "status": record.status if record else "not reached",
            "seconds": _seconds(record) if record else None,
            "output": _clip(record.output) if record and record.output is not None else None,
            "error": record.error if record else None,
        }
        if record and record.status == "failed" and record.checks:
            step["checks"] = record.checks
        steps.append(step)
    extra = (
        {"stopped_before": {"id": stopped_before.id, "title": stopped_before.title}}
        if stopped_before is not None
        else {}
    )
    return {
        "status": status,
        **extra,
        "seconds": round(time.monotonic() - started, 1),
        "steps": steps,
        "files": output_files([r.output for r in records.values()], ctx.workspace_root, started_at),
        "failed_step": outcome.step_id if status == "failed" else None,
        "error": outcome.error if status == "failed" else None,
    }
